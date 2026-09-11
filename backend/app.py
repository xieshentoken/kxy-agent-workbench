from __future__ import annotations

import asyncio
import csv
import fnmatch
import hashlib
import io
import json
import mimetypes
import os
import re
import shutil
import signal
import sqlite3
import subprocess
import sys
import tempfile
import threading
import time
import uuid
import zipfile
from collections import deque
from contextlib import asynccontextmanager
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable

from fastapi import FastAPI, File, HTTPException, Query, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse, Response, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, ConfigDict, Field

try:
    from PIL.Image import DecompressionBombError as _PILDecompressionBombError
except Exception:  # Pillow is an optional import for non-image backend paths.
    class _PILDecompressionBombError(Exception):
        pass

try:  # LFX is the only graph runtime used by kxy.
    from lfx.custom.custom_component.component import Component
    from lfx.graph import Graph
    from lfx.io import DataInput, MultilineInput, Output, StrInput
    from lfx.schema import Data

    LFX_IMPORT_ERROR: str | None = None
except Exception as exc:  # pragma: no cover - exercised by the setup error path.
    Component = Graph = DataInput = MultilineInput = Output = StrInput = Data = None  # type: ignore[assignment]
    LFX_IMPORT_ERROR = f"{type(exc).__name__}: {exc}"

try:  # The configuration owner publishes this module independently.
    from .agent_config import (
        build_headless_command,
        check_skill_dependencies,
        ensure_agent_config_schema,
        prepare_headless_environment,
        resolve_model_binding,
        snapshot_model_binding,
        wrap_headless_command,
        parse_final_output,
        prepare_mcp_configuration,
        mcp_runtime_environment,
        router as agent_config_router,
    )
    AGENT_CONFIG_IMPORT_ERROR: str | None = None
except ImportError:  # Keep the V1 runtime importable while the handoff is pending.
    build_headless_command = check_skill_dependencies = ensure_agent_config_schema = prepare_headless_environment = None  # type: ignore[assignment]
    resolve_model_binding = snapshot_model_binding = wrap_headless_command = parse_final_output = None  # type: ignore[assignment]
    prepare_mcp_configuration = None  # type: ignore[assignment]
    mcp_runtime_environment = None  # type: ignore[assignment]
    agent_config_router = None
    AGENT_CONFIG_IMPORT_ERROR = "backend.agent_config is not installed"

from .portable import (
    PORTABLE_MAX_ARCHIVE_BYTES,
    PortablePackageError,
    PortablePreviewExpired,
    build_package,
    commit_package,
    preview_package,
)


ROOT = Path(__file__).resolve().parents[1]
DATA_ROOT = Path(os.environ.get("KXY_DATA_ROOT", ROOT / "data")).expanduser().resolve()
DB_PATH = DATA_ROOT / "db" / "kxy.sqlite3"
FILES_ROOT = DATA_ROOT / "files"
SKILLS_ROOT = DATA_ROOT / "skills"
RUNS_ROOT = DATA_ROOT / "runs"
OUTPUTS_ROOT = DATA_ROOT / "outputs"
BRANDING_ROOT = ROOT / "branding"
MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_TEXT_BYTES = 250_000
MAX_SKILL_BYTES = 20 * 1024 * 1024
MAX_SKILL_FILES = 400
MAX_BACKGROUND_BYTES = 10 * 1024 * 1024
MAX_BACKGROUND_PIXELS = 40_000_000
MAX_EVENT_TEXT = 12_000
MAX_DISPLAY_CHARS = 512
MAX_WORKFLOW_DEPTH = 4
MAX_EXPANDED_NODES = 200
MAX_EXPANDED_EDGES = 1000
MAX_OUTPUT_SCHEMA_BYTES = 200_000
MAX_MODELS_RESPONSE_BYTES = 1_000_000
HUMAN_POLL_SECONDS = 0.25
LOOP_DEFAULT_MAX_ROUNDS = 3
LOOP_MAX_ROUNDS = 50
LOOP_DEFAULT_ACTIVE_BUDGET_SECONDS = 1800
LOOP_MAX_ACTIVE_BUDGET_SECONDS = 86_400
LOOP_DEFAULT_PREVIOUS_SUMMARY_CHARS = 4_000
LOOP_MAX_PREVIOUS_SUMMARY_CHARS = 20_000
FILTER_PATTERN_MODES = {"glob", "regex"}
FILTER_NAME_SCOPES = {"basename", "relative_path"}
FILTER_MAX_PATTERN_CHARS = 512
FILTER_MAX_EXTENSION_ITEMS = 100
FILTER_MAX_CANDIDATES = 1_000
FILTER_DEFAULT_TIMEOUT_MS = 100
FILTER_MAX_TIMEOUT_MS = 1_000
LOOP_REVIEWER_PROMPT = (
    "Review the latest executor result against the original goal. "
    "Return only one JSON object with this exact shape: "
    '{"passed":true|false,"issues":["..."],"next_action":"..."}. '
    "passed must be a boolean, issues must be an array of strings, and "
    "next_action must be a string. Do not add Markdown or other keys."
)
DEFAULT_ANALYZER_PROMPT = (
    "请基于输入资料给出可核查、带边界的分析。\n\n"
    "如需读取本次运行资料，请打开 input-context.json；输入文件副本位于 inputs/。"
    "若已选择 Skill，请读取 skill-manifest.json 中列出的快照；生成文件请写入 outputs/。"
)
MAX_DEFAULT_ANALYZER_PROMPT_CHARS = 100_000
ALLOWED_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".html",
    ".htm",
    ".yaml",
    ".yml",
    ".csv",
    ".xlsx",
    ".pdf",
    ".jpg",
    ".jpeg",
    ".png",
}
SAFE_NAME_RE = re.compile(r"[^A-Za-z0-9._()\-\u4e00-\u9fff ]+")
HEX_COLOR_RE = re.compile(r"^#[0-9a-fA-F]{6}$")
SECRET_RE = re.compile(
    r"(?i)(api[_ -]?key|access[_ -]?token|refresh[_ -]?token|token|password|secret|authorization)"
    r"(\s*[:=]\s*|\s+)[^\s,;]+"
)

DEFAULT_SETTINGS = {
    "palette": "paper",
    "accent": "#987a5d",
    "canvas": "#faf8f5",
    "font": "system",
    "custom_font": "",
    "text_font": "",
    "code_font": "",
    "font_size": 14,
    "code_font_size": 13,
    "background_image": "",
    "motion_enabled": True,
    "motion_intensity": 30,
    "undo_limit": 5,
    # These are workspace defaults only.  Existing workflows and explicit
    # node settings keep their own values.
    "agent_check_on_settings_open": False,
    "auto_check_mounts": True,
    "default_analyzer": {
        "mode": "cli",
        "agent_id": "codex",
        "prompt": DEFAULT_ANALYZER_PROMPT,
    },
    "output_defaults": {
        "export_formats": [],
        "allowed_file_extensions": [],
        "json_mode": "full",
    },
}
MOTION_SCALE_V8_MIGRATION_KEY = "motion_scale_v8_migrated"
MOTION_INTENSITY_MAX = 500
ALLOWED_PALETTES = {"paper", "sage", "mist"}
ALLOWED_FONTS = {"system", "serif", "mono"}
ALLOWED_CREDENTIAL_ENVS = {"OPENAI_API_KEY", "ANTHROPIC_API_KEY", "GEMINI_API_KEY"}
DEFAULT_API_ENDPOINTS = {
    "openai": "https://api.openai.com/v1",
    "anthropic": "https://api.anthropic.com/v1",
    "google": "https://generativelanguage.googleapis.com/v1beta",
}
API_FORMATS = {
    "openai-chat-completions",
    "openai-responses",
    "anthropic-messages",
    "google-generative-ai",
}
API_FORMAT_ENV_NAMES = {
    "openai-chat-completions": "OPENAI_API_KEY",
    "openai-responses": "OPENAI_API_KEY",
    "anthropic-messages": "ANTHROPIC_API_KEY",
    "google-generative-ai": "GEMINI_API_KEY",
}
OUTPUT_FORMATS = {"auto", "markdown", "text", "json"}
CONTAINER_EXPORT_FORMATS = {"markdown", "text", "json"}
GENERATED_FILE_EXTENSIONS = {
    ".txt",
    ".md",
    ".json",
    ".csv",
    ".pdf",
    ".docx",
    ".xlsx",
    ".pptx",
    ".png",
    ".jpg",
    ".jpeg",
    ".webp",
    ".svg",
    ".html",
    ".mp4",
    ".avi",
    ".zip",
}
JSON_MODES = {"content", "full"}
BACKGROUND_ROOT = DATA_ROOT / "appearance" / "background"
BRANDING_LOGO_FILES = (
    ("logo.png", "image/png"),
    ("logo.webp", "image/webp"),
    ("logo.svg", "image/svg+xml"),
)
FONT_CACHE_SECONDS = 30.0
_FONT_CACHE: tuple[float, list[dict[str, Any]], str, str | None] | None = None
FONT_FAMILY_RE = re.compile(r"^[^\\\";:{}\[\]<>\x00-\x1f\x7f]{1,160}$")

RUN_PROCESSES: dict[str, subprocess.Popen[str]] = {}
RUN_CANCEL_EVENTS: dict[str, threading.Event] = {}
RUN_APPROVAL_EVENTS: dict[tuple[str, str], threading.Event] = {}
RUN_LOCK = threading.RLock()


class RunCancelled(RuntimeError):
    """Raised inside an LFX component when the owning run is cancelled."""


class RunRejected(RuntimeError):
    """Raised when a human explicitly rejects the active run."""


class LoopBudgetExceeded(RuntimeError):
    """Raised when a loop reaches its active execution budget."""


class LoopStopped(RunCancelled):
    """Raised when the user stops a paused loop without accepting output."""


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def json_text(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def redact(value: str, max_length: int = MAX_EVENT_TEXT) -> str:
    return SECRET_RE.sub(lambda match: f"{match.group(1)}=***", value)[:max_length]


def redact_object(value: Any, key: str = "", redaction_values: tuple[str, ...] = (), schema_context: bool = False) -> Any:
    lower_key = str(key).lower()
    secret_key = lower_key in {"api_key", "apikey", "token", "access_token", "refresh_token", "secret", "password", "authorization", "client_secret"} or lower_key.endswith(("_api_key", "_password", "_secret", "_token"))
    if secret_key and not schema_context:
        return "***"
    if isinstance(value, dict):
        return {
            str(item_key): redact_object(item_value, str(item_key), redaction_values, schema_context or lower_key == "output_schema")
            for item_key, item_value in value.items()
        }
    if isinstance(value, list):
        return [redact_object(item, key, redaction_values, schema_context) for item in value]
    if isinstance(value, str):
        safe = redact(value, max_length=MAX_TEXT_BYTES)
        for secret_value in redaction_values:
            safe = safe.replace(secret_value, "[REDACTED]")
        return safe
    return value


def safe_display_name(name: str | None) -> str:
    raw = Path(name or "untitled").name.replace("\x00", "")
    raw = SAFE_NAME_RE.sub("_", raw).strip(" .") or "untitled"
    return raw[:MAX_DISPLAY_CHARS]


def safe_storage_name(display_name: str, digest: str) -> str:
    """Use a short, byte-safe on-disk basename while retaining the display name."""
    suffix = Path(display_name).suffix.lower()
    suffix = SAFE_NAME_RE.sub("", suffix)
    if len(suffix.encode("utf-8")) > 32:
        suffix = ""
    return f"{digest[:32]}{suffix}"


def safe_run_input_name(display_name: str, digest: str) -> str:
    """Keep run input paths short; the original display name remains in the manifest."""
    return safe_storage_name(display_name, digest)


def resolved_inside(path: Path, root: Path) -> bool:
    try:
        path.resolve().relative_to(root.resolve())
    except ValueError:
        return False
    return True


def ensure_roots() -> None:
    for path in (DATA_ROOT, DB_PATH.parent, FILES_ROOT, SKILLS_ROOT, RUNS_ROOT, OUTPUTS_ROOT, BACKGROUND_ROOT):
        path.mkdir(parents=True, exist_ok=True)


def connect_db() -> sqlite3.Connection:
    ensure_roots()
    connection = sqlite3.connect(DB_PATH, check_same_thread=False)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys=ON")
    return connection


def init_db() -> None:
    interrupted_ids: list[str] = []
    with connect_db() as db:
        db.executescript(
            """
            CREATE TABLE IF NOT EXISTS settings (
                key TEXT PRIMARY KEY,
                value TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS workflows (
                id TEXT PRIMARY KEY,
                name TEXT NOT NULL,
                version TEXT NOT NULL,
                document TEXT NOT NULL,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS files (
                id TEXT PRIMARY KEY,
                sha256 TEXT NOT NULL UNIQUE,
                display_name TEXT NOT NULL,
                stored_path TEXT NOT NULL,
                mime TEXT NOT NULL,
                size INTEGER NOT NULL,
                extension TEXT NOT NULL,
                status TEXT NOT NULL,
                preview TEXT,
                metadata TEXT NOT NULL,
                created_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS skills (
                id TEXT PRIMARY KEY,
                name TEXT NOT NULL,
                description TEXT NOT NULL,
                root_path TEXT NOT NULL,
                snapshot_hash TEXT NOT NULL,
                metadata TEXT NOT NULL,
                created_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS analyzers (
                id TEXT PRIMARY KEY,
                name TEXT NOT NULL,
                config TEXT NOT NULL,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS grants (
                id TEXT PRIMARY KEY,
                canonical_path TEXT NOT NULL,
                revoked_at TEXT,
                created_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS credentials (
                id TEXT PRIMARY KEY,
                provider TEXT NOT NULL,
                credential_ref TEXT NOT NULL,
                endpoint TEXT,
                env_name TEXT NOT NULL,
                created_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS credential_metadata (
                credential_id TEXT PRIMARY KEY,
                name TEXT NOT NULL DEFAULT '',
                api_format TEXT NOT NULL DEFAULT '',
                models TEXT NOT NULL DEFAULT '[]',
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS appearance_assets (
                id TEXT PRIMARY KEY,
                path TEXT NOT NULL,
                mime TEXT NOT NULL,
                sha256 TEXT NOT NULL,
                size INTEGER NOT NULL,
                width INTEGER NOT NULL,
                height INTEGER NOT NULL,
                created_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS runs (
                id TEXT PRIMARY KEY,
                workflow_id TEXT,
                status TEXT NOT NULL,
                snapshot TEXT NOT NULL,
                run_dir TEXT NOT NULL,
                output_manifest TEXT,
                error TEXT,
                cancel_requested INTEGER NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL,
                started_at TEXT,
                finished_at TEXT
            );
            CREATE TABLE IF NOT EXISTS run_nodes (
                run_id TEXT NOT NULL,
                node_id TEXT NOT NULL,
                status TEXT NOT NULL,
                message TEXT,
                output TEXT,
                started_at TEXT,
                finished_at TEXT,
                PRIMARY KEY (run_id, node_id),
                FOREIGN KEY (run_id) REFERENCES runs(id) ON DELETE CASCADE
            );
            CREATE TABLE IF NOT EXISTS run_events (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                run_id TEXT NOT NULL,
                event_type TEXT NOT NULL,
                payload TEXT NOT NULL,
                created_at TEXT NOT NULL,
                FOREIGN KEY (run_id) REFERENCES runs(id) ON DELETE CASCADE
            );
            CREATE INDEX IF NOT EXISTS run_events_run_id_idx ON run_events(run_id, id);
            CREATE TABLE IF NOT EXISTS run_approvals (
                run_id TEXT NOT NULL,
                node_id TEXT NOT NULL,
                node_path TEXT NOT NULL,
                status TEXT NOT NULL,
                content TEXT NOT NULL,
                confirm_label TEXT NOT NULL DEFAULT '确认继续',
                inputs TEXT NOT NULL,
                allow_return INTEGER NOT NULL DEFAULT 0,
                revision_text TEXT,
                revision_note TEXT,
                revision_at TEXT,
                decision TEXT,
                note TEXT,
                created_at TEXT NOT NULL,
                decided_at TEXT,
                PRIMARY KEY (run_id, node_id),
                FOREIGN KEY (run_id) REFERENCES runs(id) ON DELETE CASCADE
            );
            CREATE INDEX IF NOT EXISTS run_approvals_run_id_idx ON run_approvals(run_id, created_at);
            CREATE TABLE IF NOT EXISTS component_presets (
                id TEXT PRIMARY KEY,
                kind TEXT NOT NULL,
                name TEXT NOT NULL,
                config TEXT NOT NULL,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            CREATE INDEX IF NOT EXISTS component_presets_kind_idx ON component_presets(kind, updated_at);
            CREATE TABLE IF NOT EXISTS run_attempts (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                run_id TEXT NOT NULL,
                attempt_no INTEGER NOT NULL,
                status TEXT NOT NULL,
                reason TEXT,
                reused_nodes TEXT NOT NULL DEFAULT '[]',
                rerun_nodes TEXT NOT NULL DEFAULT '[]',
                created_at TEXT NOT NULL,
                started_at TEXT,
                finished_at TEXT,
                error TEXT,
                UNIQUE(run_id, attempt_no),
                FOREIGN KEY (run_id) REFERENCES runs(id) ON DELETE CASCADE
            );
            CREATE INDEX IF NOT EXISTS run_attempts_run_id_idx ON run_attempts(run_id, attempt_no);
            CREATE UNIQUE INDEX IF NOT EXISTS run_attempts_active_idx ON run_attempts(run_id) WHERE status IN ('queued', 'running');
            CREATE TABLE IF NOT EXISTS loop_rounds (
                run_id TEXT NOT NULL,
                attempt_no INTEGER NOT NULL,
                loop_path TEXT NOT NULL,
                round_no INTEGER NOT NULL,
                status TEXT NOT NULL,
                original_goal TEXT NOT NULL DEFAULT '',
                original_input TEXT NOT NULL DEFAULT '{}',
                previous_summary TEXT NOT NULL DEFAULT '',
                executor_node_id TEXT,
                executor_input TEXT,
                executor_output TEXT,
                reviewer_node_id TEXT,
                reviewer_input TEXT,
                reviewer_output TEXT,
                review_passed INTEGER,
                review_issues TEXT NOT NULL DEFAULT '[]',
                next_action TEXT,
                active_seconds REAL NOT NULL DEFAULT 0,
                effective_max_rounds INTEGER NOT NULL DEFAULT 3,
                budget_seconds REAL NOT NULL DEFAULT 1800,
                node_outputs TEXT NOT NULL DEFAULT '{}',
                node_statuses TEXT NOT NULL DEFAULT '{}',
                control TEXT,
                control_payload TEXT,
                error TEXT,
                created_at TEXT NOT NULL,
                started_at TEXT,
                finished_at TEXT,
                PRIMARY KEY (run_id, attempt_no, loop_path, round_no),
                FOREIGN KEY (run_id) REFERENCES runs(id) ON DELETE CASCADE
            );
            CREATE INDEX IF NOT EXISTS loop_rounds_run_idx ON loop_rounds(run_id, loop_path, attempt_no, round_no);
            """
        )
        approval_columns = {
            row["name"] for row in db.execute("PRAGMA table_info(run_approvals)").fetchall()
        }
        if "confirm_label" not in approval_columns:
            db.execute(
                "ALTER TABLE run_approvals ADD COLUMN confirm_label TEXT NOT NULL DEFAULT '确认继续'"
            )
        for column, definition in (
            ("allow_return", "INTEGER NOT NULL DEFAULT 0"),
            ("revision_text", "TEXT"),
            ("revision_note", "TEXT"),
            ("revision_at", "TEXT"),
        ):
            if column not in approval_columns:
                db.execute(f"ALTER TABLE run_approvals ADD COLUMN {column} {definition}")
        loop_columns = {row["name"] for row in db.execute("PRAGMA table_info(loop_rounds)").fetchall()}
        if "effective_max_rounds" not in loop_columns:
            db.execute(
                "ALTER TABLE loop_rounds ADD COLUMN effective_max_rounds INTEGER NOT NULL DEFAULT 3"
            )
        if "budget_seconds" not in loop_columns:
            db.execute(
                "ALTER TABLE loop_rounds ADD COLUMN budget_seconds REAL NOT NULL DEFAULT 1800"
            )
        existing_motion = db.execute(
            "SELECT value FROM settings WHERE key=?",
            ("motion_intensity",),
        ).fetchone()
        migration_done = db.execute(
            "SELECT value FROM settings WHERE key=?",
            (MOTION_SCALE_V8_MIGRATION_KEY,),
        ).fetchone()
        for key, value in DEFAULT_SETTINGS.items():
            db.execute(
                "INSERT OR IGNORE INTO settings(key, value) VALUES (?, ?)",
                (key, _setting_storage(value)),
            )
        if migration_done is None:
            # A pre-V8 database already has the V7 0..100 value.  A fresh
            # database did not have a row before the defaults were inserted,
            # so its new default must remain 30.  The marker makes this
            # conversion idempotent across every later startup.
            if existing_motion is not None:
                try:
                    legacy_value = int(existing_motion["value"])
                except (TypeError, ValueError):
                    # Invalid V7 values fell back to the old full-strength
                    # appearance.  Convert that legacy default once to the
                    # new scale's equivalent value (30), not 30 as a legacy
                    # input which would become 9.
                    legacy_value = 100
                if 0 <= legacy_value <= 100:
                    migrated_value = int(legacy_value * 3 / 10 + 0.5)
                    db.execute(
                        "UPDATE settings SET value=? WHERE key=?",
                        (str(migrated_value), "motion_intensity"),
                    )
            db.execute(
                "INSERT INTO settings(key, value) VALUES (?, ?)",
                (MOTION_SCALE_V8_MIGRATION_KEY, "1"),
            )
        interrupted_ids = [
            row["id"]
            for row in db.execute("SELECT id FROM runs WHERE status IN ('pending','running','waiting')").fetchall()
        ]
        db.execute(
            "UPDATE runs SET status='interrupted', finished_at=?, error=? "
            "WHERE status IN ('pending','running','waiting')",
            (utc_now(), "Server restarted before the run finished."),
        )
        db.execute(
            "UPDATE run_attempts SET status='interrupted', finished_at=COALESCE(finished_at, ?), "
            "error=COALESCE(error, ?) WHERE status IN ('queued','running') "
            "AND run_id IN (SELECT id FROM runs WHERE status='interrupted')",
            (utc_now(), "Server restarted before the attempt finished."),
        )
        db.execute("UPDATE run_nodes SET status='interrupted', finished_at=? WHERE status IN ('pending','running','waiting') AND run_id IN (SELECT id FROM runs WHERE status='interrupted')", (utc_now(),))
        db.execute("UPDATE run_approvals SET status='interrupted', decided_at=? WHERE status='pending' AND run_id IN (SELECT id FROM runs WHERE status='interrupted')", (utc_now(),))
        db.execute(
            "UPDATE loop_rounds SET status='interrupted', finished_at=COALESCE(finished_at, ?), "
            "error=COALESCE(error, ?) WHERE status IN ('queued','running','waiting','budget_waiting','round_limit_waiting') "
            "AND run_id IN (SELECT id FROM runs WHERE status='interrupted')",
            (utc_now(), "Server restarted before the loop round finished."),
        )
    # If a test or an embedding process calls init_db in the same interpreter,
    # wake any in-process workers so they cannot overwrite the durable restart
    # state with a later cancelled/failed result.
    with RUN_LOCK:
        for run_id in interrupted_ids:
            event = RUN_CANCEL_EVENTS.get(run_id)
            if event:
                event.set()
            process = RUN_PROCESSES.get(run_id)
            if process:
                kill_process_group(process)
            for (approval_run_id, _node_id), approval_handle in list(RUN_APPROVAL_EVENTS.items()):
                if approval_run_id == run_id:
                    approval_handle.set()
    # A hard process kill skips execute_cli's finally block.  Remove only
    # task-local CLI state below the exact interrupted run roots; never sweep a
    # native home or a broad data directory.
    for run_id in interrupted_ids:
        run_root = (RUNS_ROOT / str(run_id)).resolve()
        if not resolved_inside(run_root, RUNS_ROOT) or not run_root.is_dir():
            continue
        workspace_root = run_root / "workspace"
        for state_root in (workspace_root.rglob(".cli-state") if workspace_root.is_dir() else []):
            if state_root.is_symlink() or not state_root.is_dir():
                continue
            if resolved_inside(state_root, workspace_root):
                shutil.rmtree(state_root, ignore_errors=True)


def row_dict(row: sqlite3.Row | None) -> dict[str, Any] | None:
    return dict(row) if row is not None else None


SERVICE_MODEL_ID_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_.:@/+-]{0,255}$")


def default_api_format(provider: str, env_name: str = "") -> str:
    provider_value = str(provider or "").strip().lower()
    env_value = str(env_name or "").strip().upper()
    if provider_value == "google" or env_value == "GEMINI_API_KEY":
        return "google-generative-ai"
    if provider_value == "anthropic" or env_value == "ANTHROPIC_API_KEY":
        return "anthropic-messages"
    # Existing OpenAI credentials already feed Codex's verified Responses
    # mapping. Keep that behavior for rows created before V6 metadata existed.
    if provider_value == "openai" or env_value == "OPENAI_API_KEY":
        return "openai-responses"
    return "openai-chat-completions"


def normalize_api_format(value: Any, provider: str = "", env_name: str = "") -> str:
    result = default_api_format(provider, env_name) if value in (None, "") else str(value).strip().lower()
    if result not in API_FORMATS:
        raise ValueError("api_format 必须是受支持的 API 接口格式")
    return result


def api_format_env_name(api_format: str) -> str:
    try:
        return API_FORMAT_ENV_NAMES[api_format]
    except KeyError as exc:
        raise ValueError("api_format 必须是受支持的 API 接口格式") from exc


def normalize_service_env_name(value: Any, api_format: str, *, default: bool = True) -> str:
    expected = api_format_env_name(api_format)
    actual = str(value or "").strip().upper()
    if not actual and default:
        return expected
    if actual != expected:
        raise ValueError(f"{api_format} 必须使用 {expected}")
    return expected


def normalize_service_models(value: Any, *, secret: str | None = None) -> list[dict[str, str]]:
    if value in (None, ""):
        return []
    if not isinstance(value, list):
        raise ValueError("models 必须是模型对象数组")
    if len(value) > 500:
        raise ValueError("models 最多保存 500 个模型")
    result: list[dict[str, str]] = []
    seen: set[str] = set()
    for item in value:
        if isinstance(item, str):
            raw_id = item
            raw_alias = item
        elif isinstance(item, dict):
            raw_id = item.get("id") or item.get("name")
            raw_alias = item.get("alias") or item.get("display_name") or item.get("displayName") or item.get("name") or raw_id
        else:
            raise ValueError("models 中每一项都必须是对象")
        model_id = str(raw_id or "").strip()
        if not model_id or not SERVICE_MODEL_ID_RE.fullmatch(model_id):
            raise ValueError("models 中包含无效模型 id")
        alias = re.sub(r"[\x00-\x1f\x7f]", " ", str(raw_alias or model_id)).strip()[:256] or model_id
        if secret and (secret in model_id or secret in alias):
            # A provider response is untrusted input. Do not reflect a leaked
            # key into the catalog response or durable metadata.
            continue
        if model_id in seen:
            continue
        seen.add(model_id)
        result.append({"id": model_id, "alias": alias})
    return result


def credential_metadata_row(credential_id: str, db: sqlite3.Connection | None = None) -> dict[str, Any] | None:
    owns_connection = db is None
    connection = db or connect_db()
    try:
        try:
            row = connection.execute(
                "SELECT credential_id, name, api_format, models, updated_at FROM credential_metadata WHERE credential_id=?",
                (str(credential_id),),
            ).fetchone()
        except sqlite3.OperationalError:
            # A direct legacy import can be read before init_db creates the
            # additive metadata table. The six-column credentials row remains
            # fully usable in that case.
            return None
        return dict(row) if row is not None else None
    finally:
        if owns_connection:
            connection.close()


def credential_public(row: sqlite3.Row | dict[str, Any], *, configured: bool | None = None, db: sqlite3.Connection | None = None) -> dict[str, Any]:
    raw = dict(row)
    metadata = credential_metadata_row(str(raw.get("id", "")), db) or {}
    provider = str(raw.get("provider") or "custom")
    env_name = str(raw.get("env_name") or "")
    try:
        api_format = normalize_api_format(metadata.get("api_format"), provider, env_name)
    except ValueError:
        api_format = default_api_format(provider, env_name)
    try:
        models = normalize_service_models(json.loads(metadata.get("models") or "[]"))
    except (TypeError, ValueError, json.JSONDecodeError):
        models = []
    result = {
        "id": str(raw["id"]),
        "name": str(metadata.get("name") or provider)[:160],
        "provider": provider[:80],
        "api_format": api_format,
        "models": models,
        "credential_ref": str(raw.get("credential_ref") or ""),
        "endpoint": str(raw.get("endpoint") or ""),
        "env_name": env_name,
        "created_at": raw.get("created_at"),
    }
    if configured is not None:
        result["configured"] = bool(configured)
    return result


def save_credential_metadata(
    db: sqlite3.Connection,
    credential_id: str,
    *,
    name: str,
    api_format: str,
    models: list[dict[str, str]],
) -> None:
    db.execute(
        "INSERT INTO credential_metadata(credential_id, name, api_format, models, updated_at) VALUES (?, ?, ?, ?, ?) "
        "ON CONFLICT(credential_id) DO UPDATE SET name=excluded.name, api_format=excluded.api_format, "
        "models=excluded.models, updated_at=excluded.updated_at",
        (credential_id, name[:160], api_format, json_text(models), utc_now()),
    )


def file_row(row: sqlite3.Row) -> dict[str, Any]:
    result = dict(row)
    result["metadata"] = json.loads(result.pop("metadata") or "{}")
    result.pop("stored_path", None)
    result["path"] = None
    return result


def skill_row(row: sqlite3.Row) -> dict[str, Any]:
    result = dict(row)
    result["metadata"] = json.loads(result.pop("metadata") or "{}")
    result.pop("root_path", None)
    return result


def analyzer_row(row: sqlite3.Row) -> dict[str, Any]:
    result = dict(row)
    result["config"] = json.loads(result["config"] or "{}")
    result["config"] = sanitize_config(result["config"])[0]
    return result


def sanitize_config(config: dict[str, Any]) -> tuple[dict[str, Any], list[str]]:
    """Strip embedded secrets at every depth; opaque credential references are portable."""
    removed: list[str] = []
    def clean(value: Any, prefix: str = "", schema_context: bool = False) -> Any:
        if isinstance(value, dict):
            result = {}
            for key, item in value.items():
                label = f"{prefix}.{key}" if prefix else str(key)
                lower = str(key).lower()
                # JSON Schema is a data definition, not a credential-bearing
                # runtime configuration.  Preserve property names such as
                # `token_count` and `keywords` while keeping the normal
                # recursive credential stripping for every other field.
                if schema_context or lower == "output_schema":
                    result[key] = clean(item, label, True)
                    continue
                secret = lower in {"api_key", "apikey", "token", "access_token", "refresh_token", "secret", "password", "authorization", "client_secret"} or lower.endswith(("_api_key", "_password", "_secret"))
                reference = lower.endswith(("_ref", "_id"))
                if secret and not reference:
                    if item not in (None, "", False):
                        removed.append(label)
                else:
                    result[key] = clean(item, label, False)
            return result
        if isinstance(value, list):
            return [clean(item, f"{prefix}[{index}]", schema_context) for index, item in enumerate(value)]
        return value
    return clean(config), removed


PRESET_NODE_TYPES = frozenset(
    {"file", "text", "input", "filter", "analyzer", "container", "condition", "human", "blackbox"}
)


def sanitize_component_preset_config(config: dict[str, Any]) -> tuple[dict[str, Any], list[str]]:
    """Keep local references and resolved bindings while stripping key literals."""

    # ``sanitize_config`` already removes actual API keys/passwords at every
    # depth.  Opaque credential_id/credential_ref values and resolved_model
    # metadata are safe local references: dropping them would make an applied
    # blackbox silently lose its nested Agent/model association.
    return sanitize_config(config)


def component_preset_row(row: sqlite3.Row, source: str = "component") -> dict[str, Any]:
    result = dict(row)
    result["config"] = json.loads(result.get("config") or "{}")
    result["config"] = sanitize_component_preset_config(result["config"])[0]
    result["source"] = source
    result["kind"] = str(result.get("kind") or "analyzer")
    return result


NODE_TYPES = {"file", "text", "input", "filter", "analyzer", "condition", "container", "human", "blackbox"}
SUBFLOW_MARKERS = {"subflow_input", "subflow_output"}
CONDITION_HANDLES = {"true", "false"}
CONDITION_OPERATORS = {"equals", "contains", "exists", "gt", "gte", "lt", "lte"}
CONDITION_MISSING_STRATEGIES = {"error", "false", "true"}


def _normalise_filter_extension(value: Any) -> str:
    raw = str(value or "").strip().lower()
    if raw in {"", "*", ".*", "*.*"}:
        return "*"
    if raw.startswith("*."):
        raw = raw[1:]
    if not raw.startswith("."):
        raw = "." + raw
    if len(raw) > 32 or not re.fullmatch(r"\.[a-z0-9][a-z0-9._-]*", raw):
        raise ValueError("filter extensions must use safe dot-prefixed extensions")
    return raw


def normalize_filter_config(data: dict[str, Any] | None) -> dict[str, Any]:
    raw = data if isinstance(data, dict) else {}
    pattern = str(raw.get("name_pattern", raw.get("pattern", "*")) or "*")
    if len(pattern) > FILTER_MAX_PATTERN_CHARS:
        raise ValueError(f"filter name pattern exceeds {FILTER_MAX_PATTERN_CHARS} characters")
    mode = str(raw.get("pattern_mode", "glob") or "glob").strip().lower()
    if mode not in FILTER_PATTERN_MODES:
        raise ValueError("filter pattern_mode must be glob or regex")
    scope = str(raw.get("name_scope", raw.get("match_scope", "basename")) or "basename").strip().lower()
    if scope not in FILTER_NAME_SCOPES:
        raise ValueError("filter name_scope must be basename or relative_path")
    raw_extensions = raw.get("extensions", raw.get("formats", ["*"]))
    if raw_extensions in (None, ""):
        raw_extensions = ["*"]
    if isinstance(raw_extensions, str):
        raw_extensions = re.split(r"[,\n]", raw_extensions)
    if not isinstance(raw_extensions, list):
        raise ValueError("filter extensions must be an array")
    custom = raw.get("custom_extensions", "")
    if isinstance(custom, str):
        custom_values = re.split(r"[,\n]", custom) if custom.strip() else []
    elif isinstance(custom, list):
        custom_values = custom
    elif custom in (None, ""):
        custom_values = []
    else:
        raise ValueError("filter custom_extensions must be a comma-separated string or array")
    values = [_normalise_filter_extension(item) for item in [*raw_extensions, *custom_values]]
    extensions: list[str] = []
    for extension in values:
        if extension not in extensions:
            extensions.append(extension)
    if not extensions:
        raise ValueError("filter extensions must not be empty; choose * or a specific extension")
    if len(extensions) > FILTER_MAX_EXTENSION_ITEMS:
        raise ValueError(f"filter extensions exceed {FILTER_MAX_EXTENSION_ITEMS} items")
    try:
        timeout_ms = int(raw.get("timeout_ms", raw.get("timeout", FILTER_DEFAULT_TIMEOUT_MS)))
    except (TypeError, ValueError) as exc:
        raise ValueError("filter timeout_ms must be an integer") from exc
    if not 1 <= timeout_ms <= FILTER_MAX_TIMEOUT_MS:
        raise ValueError(f"filter timeout_ms must be between 1 and {FILTER_MAX_TIMEOUT_MS}")
    if mode == "regex":
        try:
            re.compile(pattern)
        except re.error as exc:
            raise ValueError(f"filter regex is invalid: {exc}") from exc
    return {
        "name_pattern": pattern,
        "pattern_mode": mode,
        "name_scope": scope,
        "extensions": extensions,
        "timeout_ms": timeout_ms,
    }


def normalize_output_format(config: dict[str, Any]) -> str:
    """Resolve output handling without changing the user's prompt."""
    if config.get("expect_json") is True:
        return "json"
    value = config.get("output_format", "auto")
    if not isinstance(value, str) or value not in OUTPUT_FORMATS:
        raise ValueError("output_format must be auto, markdown, text, or json")
    return value


def normalize_container_formats(value: Any) -> list[str]:
    """Normalize optional reply exports; an explicit [] means no reply file."""
    if value is None:
        return ["markdown", "json"]
    values = [value] if isinstance(value, str) else value
    if not isinstance(values, list):
        raise ValueError("export_formats must be an array")
    result: list[str] = []
    for item in values:
        if not isinstance(item, str) or item not in CONTAINER_EXPORT_FORMATS:
            raise ValueError("export_formats supports markdown, text, and json")
        if item not in result:
            result.append(item)
    return result


def normalize_allowed_file_extensions(value: Any) -> list[str] | None:
    """Return V6 file allowlist; None preserves pre-V6 container behavior."""
    if value is None:
        return None
    if not isinstance(value, list):
        raise ValueError("allowed_file_extensions must be an array")
    result: list[str] = []
    for item in value:
        if not isinstance(item, str):
            raise ValueError("allowed_file_extensions must contain strings")
        extension = item.strip().lower()
        if not extension.startswith(".") or len(extension) > 16 or not re.fullmatch(r"\.[a-z0-9]+", extension):
            raise ValueError("allowed_file_extensions must use dot-prefixed lowercase extensions")
        if extension not in GENERATED_FILE_EXTENSIONS:
            raise ValueError(f"unsupported generated file extension: {extension}")
        if extension not in result:
            result.append(extension)
    return result


def output_schema_errors(schema: Any) -> list[str]:
    """Validate a local-only JSON Schema without ever resolving remote refs."""
    if schema in (None, ""):
        return []
    if not isinstance(schema, dict):
        return ["output_schema must be a JSON object"]
    try:
        encoded = json.dumps(schema, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    except (TypeError, ValueError):
        return ["output_schema must be JSON serializable"]
    if len(encoded) > MAX_OUTPUT_SCHEMA_BYTES:
        return [f"output_schema exceeds {MAX_OUTPUT_SCHEMA_BYTES} bytes"]
    errors: list[str] = []

    def inspect_refs(value: Any) -> None:
        if isinstance(value, dict):
            for reference_key in ("$ref", "$dynamicRef", "$recursiveRef"):
                if reference_key not in value:
                    continue
                ref = value[reference_key]
                if not isinstance(ref, str) or not ref.startswith("#"):
                    errors.append("output_schema only allows local $ref values")
            for item in value.values():
                inspect_refs(item)
        elif isinstance(value, list):
            for item in value:
                inspect_refs(item)

    inspect_refs(schema)
    if errors:
        return errors
    try:
        import jsonschema

        jsonschema.validators.validator_for(schema).check_schema(schema)
    except Exception:
        return ["output_schema is not a valid JSON Schema"]
    return []


def validate_json_schema_value(value: Any, schema: Any) -> None:
    errors = output_schema_errors(schema)
    if errors:
        raise ValueError(errors[0])
    if schema in (None, ""):
        return
    try:
        import jsonschema
        from referencing import Registry

        def reject_retrieve(_uri: str) -> Any:
            raise RuntimeError("remote JSON Schema retrieval is disabled")

        registry = Registry(retrieve=reject_retrieve)
        validator = jsonschema.validators.validator_for(schema)(schema, registry=registry)
        validator.validate(value)
    except Exception as exc:
        # Do not include validator paths or the returned object in a run error;
        # both can contain sensitive input values.
        if exc.__class__.__module__.startswith("jsonschema"):
            raise ValueError("JSON output does not satisfy output_schema") from exc
        raise ValueError("output_schema validation failed") from exc


def validate_analyzer_output_config(data: dict[str, Any], label: str) -> None:
    try:
        output_format = normalize_output_format(data)
    except ValueError as exc:
        raise ValueError(f"Analyzer {label} {exc}") from exc
    errors = output_schema_errors(data.get("output_schema"))
    if errors:
        raise ValueError(f"Analyzer {label}: {errors[0]}")
    if data.get("output_schema") not in (None, "") and output_format != "json":
        raise ValueError(f"Analyzer {label}: output_schema requires output_format=json")


def normalize_workflow(document: dict[str, Any]) -> dict[str, Any]:
    if not isinstance(document, dict):
        raise ValueError("Workflow must be an object")
    result = json.loads(json_text(document))
    result.setdefault("version", "kxy.workflow.v1")
    result.setdefault("nodes", [])
    result.setdefault("edges", [])
    result.setdefault("settings", {})
    for node in result.get("nodes", []):
        if not isinstance(node, dict) or node.get("type") != "blackbox":
            continue
        data = node.setdefault("data", {})
        if not isinstance(data, dict):
            continue
        nested = data.get("workflow")
        if nested is None:
            nested = data.get("subflow")
        if isinstance(nested, dict):
            data["workflow"] = normalize_workflow(nested)
        if isinstance(data.get("loop"), dict):
            raw_loop = data["loop"]
            review_fields = raw_loop.get("review_fields")
            if not isinstance(review_fields, dict):
                review_fields = raw_loop.get("feedback_fields") if isinstance(raw_loop.get("feedback_fields"), dict) else {}
            feedback = raw_loop.get("feedback") if isinstance(raw_loop.get("feedback"), dict) else {}
            data["loop"] = {
                **raw_loop,
                "enabled": raw_loop.get("enabled") is True,
                "max_rounds": raw_loop.get("max_rounds", LOOP_DEFAULT_MAX_ROUNDS),
                "active_budget_seconds": raw_loop.get("active_budget_seconds", LOOP_DEFAULT_ACTIVE_BUDGET_SECONDS),
                "previous_summary_chars": raw_loop.get("previous_summary_chars", LOOP_DEFAULT_PREVIOUS_SUMMARY_CHARS),
                "input_field": raw_loop.get("input_field") or "text",
                "review_fields": {
                    "passed": review_fields.get("passed", "passed"),
                    "issues": review_fields.get("issues", "issues"),
                    "next_action": review_fields.get("next_action", "next_action"),
                },
                "feedback": {
                    "latest_result": feedback.get("latest_result", True) is True,
                    "issues": feedback.get("issues", True) is True,
                    "next_action": feedback.get("next_action", True) is True,
                    "history_summary": feedback.get("history_summary", True) is True,
                },
            }
            # The reviewer prompt lives on the selected inner analyzer.  Do
            # not retain a second executable prompt source on the loop card.
            data["loop"].pop("reviewer_prompt", None)
    return result


def _nested_workflow(node: dict[str, Any]) -> dict[str, Any] | None:
    data = node.get("data")
    if not isinstance(data, dict):
        return None
    nested = data.get("workflow")
    if nested is None:
        nested = data.get("subflow")
    return nested if isinstance(nested, dict) else None


def loop_settings(value: Any) -> dict[str, Any]:
    """Return the small, serializable loop configuration used by V9."""

    raw = value if isinstance(value, dict) else {}
    fields = raw.get("review_fields")
    if not isinstance(fields, dict):
        # V9 early drafts called this feedback_fields.  Read it for imported
        # drafts, but keep the public representation on review_fields.
        fields = raw.get("feedback_fields") if isinstance(raw.get("feedback_fields"), dict) else {}
    feedback = raw.get("feedback") if isinstance(raw.get("feedback"), dict) else {}
    return {
        "enabled": raw.get("enabled") is True,
        "executor_id": str(raw.get("executor_id") or "").strip(),
        "reviewer_id": str(raw.get("reviewer_id") or "").strip(),
        "goal": str(raw.get("goal") or ""),
        "input_field": str(raw.get("input_field") or "text").strip() or "text",
        "review_fields": {
            "passed": str(fields.get("passed") or "passed").strip() or "passed",
            "issues": str(fields.get("issues") or "issues").strip() or "issues",
            "next_action": str(fields.get("next_action") or "next_action").strip() or "next_action",
        },
        "feedback": {
            "latest_result": feedback.get("latest_result", True) is True,
            "issues": feedback.get("issues", True) is True,
            "next_action": feedback.get("next_action", True) is True,
            "history_summary": feedback.get("history_summary", True) is True,
        },
        "max_rounds": raw.get("max_rounds", LOOP_DEFAULT_MAX_ROUNDS),
        "active_budget_seconds": raw.get("active_budget_seconds", LOOP_DEFAULT_ACTIVE_BUDGET_SECONDS),
        "previous_summary_chars": raw.get("previous_summary_chars", LOOP_DEFAULT_PREVIOUS_SUMMARY_CHARS),
    }


def loop_enabled(node: dict[str, Any]) -> bool:
    data = node.get("data") if isinstance(node, dict) else None
    return isinstance(data, dict) and isinstance(data.get("loop"), dict) and data["loop"].get("enabled") is True


def loop_review_gate_node(workflow: dict[str, Any] | None, reviewer_id: str) -> str | None:
    """Return the only supported human return gate: reviewer -> human -> output."""
    if not isinstance(workflow, dict) or not isinstance(workflow.get("nodes"), list) or not isinstance(workflow.get("edges"), list):
        return None
    nodes = {str(item.get("id")): item for item in workflow["nodes"] if isinstance(item, dict) and item.get("id")}
    output_ids = {node_id for node_id, item in nodes.items() if item.get("type") == "subflow_output"}
    candidates = []
    for node_id, item in nodes.items():
        data = item.get("data") if isinstance(item.get("data"), dict) else {}
        if item.get("type") != "human" or data.get("review_gate") is not True:
            continue
        incoming = [edge for edge in workflow["edges"] if isinstance(edge, dict) and str(edge.get("target")) == node_id]
        outgoing = [edge for edge in workflow["edges"] if isinstance(edge, dict) and str(edge.get("source")) == node_id]
        if len(incoming) == 1 and len(outgoing) == 1 and str(incoming[0].get("source")) == reviewer_id and str(outgoing[0].get("target")) in output_ids:
            candidates.append(node_id)
    return candidates[0] if len(candidates) == 1 else None


def _path_label(path: tuple[str, ...], node_id: str | None = None) -> str:
    parts = [*path, node_id] if node_id else list(path)
    return "/".join(str(part) for part in parts) or "workflow"


def _validate_workflow_level(
    workflow: dict[str, Any],
    *,
    path: tuple[str, ...],
    depth: int,
    allow_markers: bool,
    counters: dict[str, int],
    errors: list[str],
    warnings: list[str],
    root_order: list[str] | None = None,
    inside_loop: bool = False,
) -> None:
    label = _path_label(path)
    if workflow.get("version") != "kxy.workflow.v1":
        errors.append(f"Unsupported workflow version at {label}: {workflow.get('version')}")
    nodes = workflow.get("nodes")
    edges = workflow.get("edges")
    if not isinstance(nodes, list) or not isinstance(edges, list):
        errors.append(f"Workflow at {label} must contain nodes and edges arrays")
        return
    counters["nodes"] += len(nodes)
    counters["edges"] += len(edges)
    if not nodes:
        errors.append(f"Workflow at {label} must contain at least one node")
        return

    allowed_types = NODE_TYPES | (SUBFLOW_MARKERS if allow_markers else set())
    ids: set[str] = set()
    node_map: dict[str, dict[str, Any]] = {}
    for node in nodes:
        if not isinstance(node, dict):
            errors.append(f"Every node in {label} must be an object")
            continue
        node_id = str(node.get("id", "")).strip()
        node_type = str(node.get("type", "")).strip()
        display_id = _path_label(path, node_id or "unnamed node")
        if not node_id:
            errors.append(f"A node is missing its id in {label}")
        elif node_id in ids:
            errors.append(f"Duplicate node id: {display_id}")
        else:
            ids.add(node_id)
            node_map[node_id] = node
        if node_id and not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_-]{0,99}", node_id):
            errors.append(f"Invalid node id: {display_id}")
        if not isinstance(node.get("data", {}), dict):
            errors.append(f"Node {display_id} data must be an object")
            continue
        if node_type not in allowed_types:
            errors.append(f"Unsupported node type for {display_id}: {node_type}")
        if node_type == "human":
            content = str(node["data"].get("content", ""))
            if len(content) > MAX_TEXT_BYTES:
                errors.append(f"Human content is too large for {display_id}")
            if node["data"].get("review_gate") is True and not inside_loop:
                warnings.append(
                    f"Human {display_id} is marked as a loop review gate but is outside an enabled loop; return for revision is disabled"
                )
        if node_type == "filter":
            try:
                normalize_filter_config(node["data"])
            except ValueError as exc:
                errors.append(f"Filter {display_id}: {exc}")
        if node_type == "analyzer":
            try:
                validate_analyzer_output_config(node["data"], display_id)
            except ValueError as exc:
                errors.append(str(exc))
        if node_type == "container":
            try:
                normalize_container_formats(node["data"].get("export_formats"))
            except ValueError as exc:
                errors.append(f"Container {display_id}: {exc}")
            if "allowed_file_extensions" in node["data"]:
                try:
                    normalize_allowed_file_extensions(node["data"].get("allowed_file_extensions"))
                except ValueError as exc:
                    errors.append(f"Container {display_id}: {exc}")
            json_mode = node["data"].get("json_mode", "full")
            if json_mode not in JSON_MODES:
                errors.append(f"Container {display_id}: json_mode must be content or full")
            if inside_loop and str(node["data"].get("grant_id") or "").strip():
                errors.append(
                    f"Container {display_id} uses an output grant inside a loop; "
                    "place the authorized output container after the loop"
                )
        if node_type == "blackbox":
            nested = _nested_workflow(node)
            enabled_loop = loop_enabled(node)
            if enabled_loop and inside_loop:
                errors.append(
                    f"Nested enabled loops are not supported at {display_id}; "
                    "use one loop and place any authorized output after it"
                )
            if enabled_loop:
                settings = loop_settings(node.get("data", {}).get("loop"))
                executor_id = settings["executor_id"]
                reviewer_id = settings["reviewer_id"]
                if executor_id == reviewer_id:
                    errors.append(f"Loop {display_id} needs distinct executor_id and reviewer_id")
                if not str(settings["goal"]).strip():
                    warnings.append(f"Loop {display_id} has no explicit goal; the original input will be used")
                try:
                    max_rounds = int(settings["max_rounds"])
                except (TypeError, ValueError):
                    max_rounds = 0
                if not 1 <= max_rounds <= LOOP_MAX_ROUNDS:
                    errors.append(f"Loop {display_id}: max_rounds must be between 1 and {LOOP_MAX_ROUNDS}")
                try:
                    budget = float(settings["active_budget_seconds"])
                except (TypeError, ValueError):
                    budget = 0
                if not 1 <= budget <= LOOP_MAX_ACTIVE_BUDGET_SECONDS:
                    errors.append(
                        f"Loop {display_id}: active_budget_seconds must be between 1 and "
                        f"{LOOP_MAX_ACTIVE_BUDGET_SECONDS}"
                    )
                try:
                    summary_chars = int(settings["previous_summary_chars"])
                except (TypeError, ValueError):
                    summary_chars = 0
                if not 0 <= summary_chars <= LOOP_MAX_PREVIOUS_SUMMARY_CHARS:
                    errors.append(
                        f"Loop {display_id}: previous_summary_chars must be between 0 and "
                        f"{LOOP_MAX_PREVIOUS_SUMMARY_CHARS}"
                    )
                for field_name in ("input_field", *settings["review_fields"].values()):
                    if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_.-]{0,127}", str(field_name)):
                        errors.append(f"Loop {display_id}: invalid feedback field {field_name}")
                if nested is not None and isinstance(nested.get("nodes"), list):
                    nested_nodes = {
                        str(item.get("id")): item
                        for item in nested["nodes"]
                        if isinstance(item, dict) and item.get("id")
                    }
                    for role, selected_id in (("executor_id", executor_id), ("reviewer_id", reviewer_id)):
                        selected = nested_nodes.get(selected_id)
                        if selected is None or selected.get("type") != "analyzer":
                            errors.append(
                                f"Loop {display_id}: {role} must name an inner analyzer node"
                            )
                    marked_gates = [
                        str(item.get("id"))
                        for item in nested.get("nodes", [])
                        if isinstance(item, dict)
                        and item.get("type") == "human"
                        and isinstance(item.get("data"), dict)
                        and item["data"].get("review_gate") is True
                    ]
                    gate_id = loop_review_gate_node(nested, reviewer_id)
                    if marked_gates and gate_id is None:
                        errors.append(
                            f"Loop {display_id}: return-for-revision human gate must be directly reviewer -> human -> subflow_output"
                        )
            if nested is None:
                errors.append(f"Blackbox {display_id} is missing data.workflow")
            elif depth >= MAX_WORKFLOW_DEPTH:
                errors.append(f"Blackbox nesting exceeds depth {MAX_WORKFLOW_DEPTH} at {display_id}")
            else:
                _validate_workflow_level(
                    normalize_workflow(nested),
                    path=(*path, node_id),
                    depth=depth + 1,
                    allow_markers=True,
                    counters=counters,
                    errors=errors,
                    warnings=warnings,
                    inside_loop=inside_loop or enabled_loop,
                )

    adjacency: dict[str, list[str]] = {node_id: [] for node_id in ids}
    reverse: dict[str, list[str]] = {node_id: [] for node_id in ids}
    indegree: dict[str, int] = {node_id: 0 for node_id in ids}
    incoming: dict[str, int] = {node_id: 0 for node_id in ids}
    edge_keys: set[tuple[str, str, str, str]] = set()
    for edge in edges:
        if not isinstance(edge, dict):
            errors.append(f"Every edge in {label} must be an object")
            continue
        source = str(edge.get("source", ""))
        target = str(edge.get("target", ""))
        source_handle = str(edge.get("sourceHandle") or "result")
        target_handle = str(edge.get("targetHandle") or "items")
        if source not in ids or target not in ids:
            errors.append(f"Dangling edge in {label}: {source} → {target}")
            continue
        if source == target:
            errors.append(f"Self-cycle is not allowed: {_path_label(path, source)}")
        key = (source, target, source_handle, target_handle)
        if key in edge_keys:
            errors.append(f"Duplicate edge in {label}: {source} → {target}")
        edge_keys.add(key)
        source_type = node_map[source].get("type")
        target_type = node_map[target].get("type")
        if source_type == "condition" and source_handle not in CONDITION_HANDLES:
            errors.append(f"Condition {_path_label(path, source)} must use true or false output handle")
        if source_type != "condition" and source_handle != "result":
            errors.append(f"Node {_path_label(path, source)} must use the result output handle")
        if target_handle != "items":
            errors.append(f"Node {_path_label(path, target)} must use the items input handle")
        if target_type in {"file", "text", "input"}:
            errors.append(f"Invalid target for {_path_label(path, target)}: source nodes cannot receive edges")
        adjacency[source].append(target)
        reverse[target].append(source)
        indegree[target] += 1
        incoming[target] += 1

    queue = deque(node_id for node_id, degree in indegree.items() if degree == 0)
    visited: list[str] = []
    while queue:
        current = queue.popleft()
        visited.append(current)
        for target in adjacency[current]:
            indegree[target] -= 1
            if indegree[target] == 0:
                queue.append(target)
    if len(visited) != len(ids):
        errors.append(f"Workflow at {label} contains a cycle")
    if root_order is not None:
        root_order.extend(visited)

    marker_ids = {node_type: [node_id for node_id, node in node_map.items() if node.get("type") == node_type] for node_type in SUBFLOW_MARKERS}
    if allow_markers:
        if len(marker_ids["subflow_input"]) != 1 or len(marker_ids["subflow_output"]) != 1:
            errors.append(f"Blackbox workflow at {label} requires exactly one subflow_input and one subflow_output")
        else:
            input_id = marker_ids["subflow_input"][0]
            output_id = marker_ids["subflow_output"][0]
            if incoming[input_id] != 0:
                errors.append(f"subflow_input {_path_label(path, input_id)} must not have an internal incoming edge")
            source_roots = [
                node_id
                for node_id, node in node_map.items()
                if node.get("type") in {"file", "text", "input", "blackbox"} and incoming[node_id] == 0
            ]
            if not adjacency[input_id] and not source_roots:
                errors.append(f"subflow_input {_path_label(path, input_id)} must feed an inner node")
            if not reverse[output_id]:
                errors.append(f"subflow_output {_path_label(path, output_id)} must receive an inner result")
            if adjacency[output_id]:
                errors.append(f"subflow_output {_path_label(path, output_id)} must not have an internal outgoing edge")

            # Every inner node must participate in the input-to-output subflow. This
            # prevents hidden components from running outside the blackbox boundary.
            reachable: set[str] = set()
            frontier = [input_id, *source_roots]
            while frontier:
                current = frontier.pop()
                if current in reachable:
                    continue
                reachable.add(current)
                frontier.extend(adjacency[current])
            can_reach_output: set[str] = set()
            frontier = [output_id]
            while frontier:
                current = frontier.pop()
                if current in can_reach_output:
                    continue
                can_reach_output.add(current)
                frontier.extend(reverse[current])
            for node_id in ids - reachable:
                errors.append(f"Inner node {_path_label(path, node_id)} is not reachable from subflow_input")
            # The input marker may only activate fixed source roots through the
            # compiled hidden gate; it need not have a visible path to output.
            for node_id in (ids - can_reach_output) - {input_id}:
                errors.append(f"Inner node {_path_label(path, node_id)} cannot reach subflow_output")

    for node_id, node in node_map.items():
        node_type = node.get("type")
        data = node.get("data") or {}
        if node_type == "analyzer" and not str(data.get("prompt", "")).strip():
            warnings.append(f"Analyzer {_path_label(path, node_id)} has an empty prompt")
        if node_type == "condition":
            operator = str(data.get("operator", "equals"))
            if operator not in CONDITION_OPERATORS:
                errors.append(f"Condition {_path_label(path, node_id)} has an unsupported operator")
            if incoming.get(node_id, 0) > 1:
                errors.append(f"Condition {_path_label(path, node_id)} needs one upstream result; collect inputs first")
            if not str(data.get("field", "")).strip():
                warnings.append(f"Condition {_path_label(path, node_id)} has no field yet")
            missing_strategy = str(data.get("missing_strategy", "error") or "error")
            if missing_strategy not in CONDITION_MISSING_STRATEGIES:
                errors.append(
                    f"Condition {_path_label(path, node_id)} has unsupported missing_strategy"
                )

    counters["levels"] += 1


def validate_workflow(document: dict[str, Any]) -> dict[str, Any]:
    workflow = normalize_workflow(document)
    errors: list[str] = []
    warnings: list[str] = []
    _, embedded_secrets = sanitize_config(workflow)
    if embedded_secrets:
        errors.append("Workflow contains embedded credentials; use a credential reference")
    counters = {"nodes": 0, "edges": 0, "levels": 0}
    order: list[str] = []
    _validate_workflow_level(
        workflow,
        path=(),
        depth=0,
        allow_markers=False,
        counters=counters,
        errors=errors,
        warnings=warnings,
        root_order=order,
    )
    if counters["nodes"] > MAX_EXPANDED_NODES or counters["edges"] > MAX_EXPANDED_EDGES:
        errors.append(
            f"Expanded workflow exceeds {MAX_EXPANDED_NODES} nodes or {MAX_EXPANDED_EDGES} edges "
            f"({counters['nodes']} nodes, {counters['edges']} edges)"
        )
    return {
        "ok": not errors,
        "errors": errors,
        "warnings": warnings,
        "order": order,
        "expanded_nodes": counters["nodes"],
        "expanded_edges": counters["edges"],
    }


def nested_value(payload: Any, field_name: str) -> tuple[bool, Any]:
    value: Any = payload
    if isinstance(value, dict) and "data" in value and isinstance(value["data"], dict):
        value = value["data"]
    for part in field_name.split("."):
        if isinstance(value, dict) and part in value:
            value = value[part]
        else:
            return False, None
    return True, value


def active_payloads(items: Any) -> list[dict[str, Any]]:
    if items is None:
        return []
    values = items if isinstance(items, list) else [items]
    result: list[dict[str, Any]] = []
    for item in values:
        raw = item.data if hasattr(item, "data") else item
        if isinstance(raw, dict) and (
            raw.get("status") == "skipped" or ("active" in raw and raw.get("active") is False)
        ):
            continue
        if isinstance(raw, dict):
            result.append(raw)
        else:
            result.append({"text": str(raw)})
    return result


def evaluate_condition(
    items: Any,
    field_name: str,
    operator: str,
    expected: str,
    missing_strategy: str = "error",
) -> bool:
    import math
    if operator not in {"equals", "contains", "gt", "gte", "lt", "lte", "exists"}:
        raise ValueError(f"Unsupported condition operator: {operator}")
    if missing_strategy not in CONDITION_MISSING_STRATEGIES:
        raise ValueError(f"Unsupported condition missing strategy: {missing_strategy}")
    if not field_name or any(not part for part in field_name.split(".")):
        raise ValueError("Condition field must be a nonempty dotted field path")
    payloads = active_payloads(items)
    if len(payloads) != 1:
        raise ValueError("Condition requires exactly one upstream result; collect multiple inputs first")
    found, actual = nested_value(payloads[0], field_name)
    if operator == "exists":
        return found
    if not found:
        if missing_strategy == "false":
            return False
        if missing_strategy == "true":
            return True
        raise ValueError(f"Condition field '{field_name}' is missing in upstream output")
    if operator == "equals":
        return str(actual).casefold() == str(expected).casefold()
    if operator == "contains":
        return str(expected).casefold() in str(actual).casefold()
    try:
        if isinstance(actual, bool):
            raise ValueError("Boolean is not a numeric measurement")
        left, right = float(actual), float(expected)
        if not math.isfinite(left) or not math.isfinite(right):
            raise ValueError("Non-finite number")
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Condition numeric comparison requires finite numbers for '{field_name}'") from exc
    return {"gt": left > right, "gte": left >= right, "lt": left < right, "lte": left <= right}[operator]


def topological_order(workflow: dict[str, Any]) -> list[str]:
    validation = validate_workflow(workflow)
    if not validation["ok"]:
        raise ValueError("; ".join(validation["errors"]))
    return validation["order"]


def parse_text_file(path: Path, extension: str) -> tuple[str, str, dict[str, Any]]:
    """Bound previews while preserving source locations and literal table values."""
    try:
        if extension == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            chunks, refs = [], []
            size, truncated = 0, False
            for index, page in enumerate(reader.pages):
                if index >= 2000:
                    truncated = True
                    break
                raw = page.extract_text() or ""
                refs.append(index + 1)
                if raw.strip():
                    chunks.append(f"[page {index + 1}]\n{raw}")
                    size += len(raw)
                if size > MAX_TEXT_BYTES:
                    truncated = index + 1 < len(reader.pages) or size > MAX_TEXT_BYTES
                    break
            text = "\n\n".join(chunks)
            metadata = {"pages": len(reader.pages), "page_refs": refs, "truncated": truncated}
            status = ("parsed-truncated" if truncated else "parsed") if text.strip() else "empty-or-scanned"
            return text, status, metadata
        if extension in {".xlsx", ".csv"}:
            chunks, refs = [], {}
            size, truncated = 0, False
            def append_rows(rows: Any, label: str) -> int:
                nonlocal size, truncated
                count = 0
                for number, row in enumerate(rows, 1):
                    if number > 10000 or size >= MAX_TEXT_BYTES:
                        truncated = True
                        break
                    stream = io.StringIO()
                    csv.writer(stream, lineterminator="").writerow(["" if value is None else str(value) for value in row])
                    line = f"[{label} row {number}] {stream.getvalue()}"
                    chunks.append(line)
                    size += len(line)
                    count = number
                return count
            if extension == ".xlsx":
                from openpyxl import load_workbook
                workbook = load_workbook(path, read_only=True, data_only=True)
                try:
                    for sheet in workbook.worksheets:
                        refs[sheet.title] = append_rows(sheet.iter_rows(values_only=True), f"sheet {sheet.title}")
                        if truncated:
                            break
                finally:
                    workbook.close()
                metadata = {"sheets": refs, "truncated": truncated, "formula_values": "cached-values-only"}
            else:
                with path.open("r", encoding="utf-8-sig", errors="strict", newline="") as handle:
                    rows = append_rows(csv.reader(handle), "csv")
                metadata = {"rows": rows, "truncated": truncated}
            return "\n".join(chunks), "parsed-truncated" if truncated else "parsed", metadata
        if extension in {".txt", ".md", ".markdown"}:
            text = path.read_text(encoding="utf-8-sig", errors="strict")
            return text, "parsed" if text.strip() else "empty", {}
        if extension in {".html", ".htm"}:
            from bs4 import BeautifulSoup

            raw = path.read_text(encoding="utf-8-sig", errors="strict")
            soup = BeautifulSoup(raw, "html.parser")
            for element in soup.find_all(["script", "style", "noscript", "template"]):
                element.decompose()
            root = soup.body or soup
            text = "\n".join(line.strip() for line in root.get_text("\n").splitlines() if line.strip())
            return text, "parsed" if text else "empty", {"source_format": "html", "body_only": soup.body is not None}
        if extension in {".yaml", ".yml"}:
            text = path.read_text(encoding="utf-8-sig", errors="strict")
            return text, "parsed" if text.strip() else "empty", {"source_format": "yaml", "structured": False}
        if extension in {".jpg", ".jpeg", ".png"}:
            return "", "retained-image", {"visual_input": True}
        return "", "unsupported-retained", {"supported": sorted(ALLOWED_EXTENSIONS)}
    except Exception as exc:
        return "", "unreadable: " + type(exc).__name__, {"error": str(exc)[:240]}


def store_file(upload_name: str, content: bytes) -> dict[str, Any]:
    if len(content) > MAX_FILE_BYTES:
        raise HTTPException(status_code=413, detail=f"File exceeds {MAX_FILE_BYTES // (1024 * 1024)} MB limit")
    display_name = safe_display_name(upload_name)
    extension = Path(display_name).suffix.lower()
    digest = hashlib.sha256(content).hexdigest()
    with connect_db() as db:
        existing = db.execute("SELECT * FROM files WHERE sha256=?", (digest,)).fetchone()
    reparse_existing = bool(
        existing is not None
        and existing["status"] == "unsupported-retained"
        and extension in {".html", ".htm", ".yaml", ".yml"}
        and str(existing["extension"] or "").lower() == extension
    )
    if existing is not None and not reparse_existing:
        return file_row(existing)
    if reparse_existing:
        destination = Path(str(existing["stored_path"]))
    else:
        destination_dir = FILES_ROOT / digest
        destination = destination_dir / safe_storage_name(display_name, digest)
        destination_dir.mkdir(parents=True, exist_ok=True)
        if not destination.exists():
            temp = destination_dir / f".{digest[:20]}.{uuid.uuid4().hex}.tmp"
            temp.write_bytes(content)
            temp.replace(destination)
    preview, status, metadata = parse_text_file(destination, extension)
    truncated = len(preview.encode("utf-8")) > MAX_TEXT_BYTES
    if truncated:
        encoded = preview.encode("utf-8")[:MAX_TEXT_BYTES]
        preview = encoded.decode("utf-8", errors="ignore")
        metadata["truncated"] = True
        status = "parsed-truncated"
    mime = mimetypes.guess_type(display_name)[0] or "application/octet-stream"
    with connect_db() as db:
        if reparse_existing:
            db.execute(
                "UPDATE files SET preview=?, status=?, metadata=? WHERE id=?",
                (preview, status, json_text(metadata), digest),
            )
        else:
            db.execute(
                "INSERT OR IGNORE INTO files(id, sha256, display_name, stored_path, mime, size, extension, status, preview, metadata, created_at) "
                "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (
                    digest,
                    digest,
                    display_name,
                    str(destination),
                    mime,
                    len(content),
                    extension,
                    status,
                    preview,
                    json_text(metadata),
                    utc_now(),
                ),
            )
        row = db.execute("SELECT * FROM files WHERE id=?", (digest,)).fetchone()
    assert row is not None
    return file_row(row)


def snapshot_skill(root: Path) -> tuple[str, int, int]:
    digest = hashlib.sha256()
    count = 0
    total = 0
    for path in sorted(p for p in root.rglob("*") if p.is_file()):
        if path.is_symlink():
            raise ValueError(f"Skill contains a symlink: {path.relative_to(root)}")
        relative = path.relative_to(root).as_posix().encode("utf-8")
        content = path.read_bytes()
        digest.update(relative + b"\0" + content)
        count += 1
        total += len(content)
        if count > MAX_SKILL_FILES or total > MAX_SKILL_BYTES:
            raise ValueError("Skill is larger than the import limit")
    return digest.hexdigest(), count, total


def skill_modified_at(root: Path) -> str:
    latest: float | None = None
    try:
        for path in root.rglob("*"):
            if path.is_symlink() or not path.is_file():
                continue
            latest = path.stat().st_mtime if latest is None else max(latest, path.stat().st_mtime)
    except OSError:
        return "unavailable"
    if latest is None:
        return "unavailable"
    return datetime.fromtimestamp(latest, timezone.utc).isoformat()


def skill_metadata(root: Path) -> tuple[str, str, dict[str, Any]]:
    import yaml
    skill_file = root / "SKILL.md"
    if not skill_file.is_file():
        raise ValueError("Select a single skill directory containing SKILL.md at its root")
    if skill_file.stat().st_size > 512_000:
        raise ValueError("SKILL.md is larger than 512 KB")
    text = skill_file.read_text(encoding="utf-8-sig")
    lines = text.splitlines()
    if not lines or lines[0].strip() != "---":
        raise ValueError("SKILL.md requires YAML frontmatter with name and description")
    end = next((index for index in range(1, len(lines)) if lines[index].strip() == "---"), None)
    if end is None:
        raise ValueError("SKILL.md frontmatter is unterminated")
    try:
        frontmatter = yaml.safe_load("\n".join(lines[1:end]))
    except yaml.YAMLError as exc:
        raise ValueError("SKILL.md frontmatter is invalid YAML") from exc
    if not isinstance(frontmatter, dict):
        raise ValueError("Skill frontmatter must be a mapping")
    name, description = frontmatter.get("name"), frontmatter.get("description")
    if not isinstance(name, str) or not re.fullmatch(r"[a-z0-9]+(?:-[a-z0-9]+)*", name) or len(name) > 64:
        raise ValueError("Skill name must contain at most 64 lowercase letters, digits or hyphens")
    if not isinstance(description, str) or not description.strip() or len(description) > 1024:
        raise ValueError("Skill description must be a nonempty string up to 1024 characters")
    compatibility = []
    if frontmatter.get("compatibility"):
        compatibility.append(str(frontmatter["compatibility"])[:1024])
    compatibility.extend(line.strip() for line in lines[end + 1:] if re.search(r"compatib|dependenc|requirement|install", line, re.IGNORECASE))
    return name, description.strip(), {"skill_file": "SKILL.md", "compatibility": compatibility[:20]}


def import_skill_root(source_root: Path) -> dict[str, Any]:
    source_root = source_root.resolve()
    if not source_root.is_dir():
        raise ValueError("Skill source must be a directory")
    if any(path.is_symlink() for path in [source_root, *source_root.rglob("*")]):
        raise ValueError("Symlinks are not accepted in skill imports")
    name, description, metadata = skill_metadata(source_root)
    snapshot_hash, _, _ = snapshot_skill(source_root)
    skill_id = uuid.uuid4().hex
    destination = SKILLS_ROOT / skill_id
    shutil.copytree(source_root, destination, symlinks=False)
    snapshot_hash, count, total = snapshot_skill(destination)
    metadata.update({"file_count": count, "byte_count": total, "modified_at": skill_modified_at(destination)})
    with connect_db() as db:
        db.execute(
            "INSERT INTO skills(id, name, description, root_path, snapshot_hash, metadata, created_at) VALUES (?, ?, ?, ?, ?, ?, ?)",
            (skill_id, name, description, str(destination), snapshot_hash, json_text(metadata), utc_now()),
        )
        row = db.execute("SELECT * FROM skills WHERE id=?", (skill_id,)).fetchone()
    assert row is not None
    return skill_row(row)


def import_skill_zip(content: bytes, display_name: str) -> dict[str, Any]:
    if len(content) > MAX_SKILL_BYTES:
        raise ValueError("Skill archive is larger than the import limit")
    temp_root = Path(tempfile.mkdtemp(prefix="kxy-skill-", dir=DATA_ROOT))
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > MAX_SKILL_FILES or sum(member.file_size for member in members) > MAX_SKILL_BYTES:
                raise ValueError("Skill archive contains too many files")
            for member in members:
                member_path = Path(member.filename)
                if member.filename.startswith("/") or ".." in member_path.parts:
                    raise ValueError("Skill archive contains a path traversal")
                if member.is_dir():
                    continue
                if member.external_attr >> 16 & 0o170000 == 0o120000:
                    raise ValueError("Skill archive contains a symlink")
                target = (temp_root / member_path).resolve()
                if not resolved_inside(target, temp_root):
                    raise ValueError("Skill archive escapes its extraction root")
                target.parent.mkdir(parents=True, exist_ok=True)
                if member.file_size > MAX_SKILL_BYTES:
                    raise ValueError("Skill member is too large")
                target.write_bytes(archive.read(member))
        roots = [temp_root] if (temp_root / "SKILL.md").exists() else [p for p in temp_root.iterdir() if p.is_dir()]
        if len(roots) != 1:
            raise ValueError("Archive must contain one skill directory or a root SKILL.md")
        return import_skill_root(roots[0])
    finally:
        shutil.rmtree(temp_root, ignore_errors=True)


def cli_binary(name: str) -> str | None:
    configured = os.environ.get(f"KXY_{name.upper()}_BIN", name)
    if "/" in configured:
        candidate = Path(configured).expanduser()
        return str(candidate.resolve()) if candidate.is_file() and os.access(candidate, os.X_OK) else None
    return shutil.which(configured)


def probe_cli(name: str) -> dict[str, Any]:
    binary = cli_binary(name)
    if not binary:
        return {"name": name, "available": False, "status": "NOT_INSTALLED", "version": None}
    try:
        process = subprocess.run([binary, "--version"], capture_output=True, text=True, timeout=8, check=False)
        version = redact((process.stdout or process.stderr).strip().splitlines()[0] if (process.stdout or process.stderr) else "")
        return {
            "name": name,
            "available": process.returncode == 0,
            "status": "READY" if process.returncode == 0 else "UNHEALTHY",
            "version": version,
            "binary": Path(binary).name,
        }
    except (OSError, subprocess.TimeoutExpired) as exc:
        return {"name": name, "available": False, "status": type(exc).__name__, "version": None}


KEYCHAIN_PERMISSION_STATUS = 100001


def keychain_error_message(status: int, action: str) -> str:
    if int(status) == KEYCHAIN_PERMISSION_STATUS:
        return (
            "kxy 服务启动环境受限，macOS 拒绝了 Keychain 访问（100001 / Operation not permitted）。"
            "请从本机 Terminal 启动 ./start.sh 或 Start.command。不要反复解锁钥匙串。"
        )
    return f"Keychain {action} 失败（状态 {status}），请解锁登录钥匙串并检查访问权限"


def is_keychain_permission_error(error: BaseException) -> bool:
    message = str(error)
    return str(KEYCHAIN_PERMISSION_STATUS) in message or "Operation not permitted" in message


def keychain_has(ref: str) -> bool:
    if sys.platform != "darwin" or not shutil.which("security"):
        return False
    try:
        process = subprocess.run(
            ["security", "find-generic-password", "-a", "kxy", "-s", f"kxy/{ref}"],
            capture_output=True,
            text=True,
            timeout=8,
            check=False,
        )
        if process.returncode != 0:
            output = f"{process.stdout or ''}\n{process.stderr or ''}"
            if process.returncode == KEYCHAIN_PERMISSION_STATUS or "100001" in output or "operation not permitted" in output.lower():
                raise RuntimeError(keychain_error_message(KEYCHAIN_PERMISSION_STATUS, "读取"))
        return process.returncode == 0
    except (OSError, subprocess.TimeoutExpired):
        return False


def keychain_secret(ref: str) -> str:
    if sys.platform != 'darwin':
        raise RuntimeError('macOS Keychain 不可用，请使用 CLI 原生登录')
    import ctypes
    security = ctypes.CDLL('/System/Library/Frameworks/Security.framework/Security')
    find = security.SecKeychainFindGenericPassword
    find.argtypes = [ctypes.c_void_p,ctypes.c_uint32,ctypes.c_char_p,ctypes.c_uint32,ctypes.c_char_p,ctypes.POINTER(ctypes.c_uint32),ctypes.POINTER(ctypes.c_void_p),ctypes.c_void_p]
    find.restype = ctypes.c_int32
    free = security.SecKeychainItemFreeContent
    free.argtypes = [ctypes.c_void_p,ctypes.c_void_p]
    free.restype = ctypes.c_int32
    service=('kxy/'+ref).encode(); account=b'kxy'
    size=ctypes.c_uint32(); data=ctypes.c_void_p()
    status=find(None,len(service),service,len(account),account,ctypes.byref(size),ctypes.byref(data),None)
    if status != 0: raise RuntimeError(keychain_error_message(status, '读取'))
    try:
        if not data.value or not size.value: raise RuntimeError('Keychain 配置为空')
        return ctypes.string_at(data,size.value).decode('utf-8')
    finally:
        free(None,data)


def credential_endpoint_env(env_name: str) -> str | None:
    return {
        "OPENAI_API_KEY": "OPENAI_BASE_URL",
        "ANTHROPIC_API_KEY": "ANTHROPIC_BASE_URL",
        "GEMINI_API_KEY": "GEMINI_BASE_URL",
    }.get(env_name)


def build_cli_argv(config, prompt, workspace, last_message, attachments, output_grant=None):
    cli = str(config.get('cli', 'codex')).lower()
    if cli not in {'codex', 'opencode'}:
        raise ValueError('MVP 支持 Codex 和 OpenCode 执行；Claude/pi 仅检测')
    binary = cli_binary(cli)
    if not binary:
        raise FileNotFoundError(f'{cli} CLI 未安装或不可执行')
    model = str(config.get('model') or '').strip()
    if cli == 'codex':
        argv = [binary, 'exec', '--ephemeral', '--ignore-user-config', '--json', '--output-last-message', str(last_message), '--sandbox', 'workspace-write', '--skip-git-repo-check', '-C', str(workspace)]
        if model:
            argv += ['--model', model]
        effort = str(config.get('effort') or config.get('variant') or '').strip()
        if effort:
            if effort not in {'minimal', 'low', 'medium', 'high', 'xhigh', 'max', 'ultra'}:
                raise ValueError('不支持的 Codex effort')
            argv += ['--config', f'model_reasoning_effort={json.dumps(effort)}']
        argv += ['--config', 'sandbox_workspace_write.network_access=true']
        if config.get('_endpoint'):
            if config.get('_env_name') != 'OPENAI_API_KEY':
                raise ValueError('Codex API 配置需要 OpenAI Responses 兼容密钥')
            argv += ['--config', 'model_provider="kxy"', '--config', 'model_providers.kxy=' + '{name="kxy",base_url=' + json.dumps(config['_endpoint']) + ',env_key="OPENAI_API_KEY",wire_api="responses"}']
        elif config.get('_env_name'):
            argv += ['--config', 'model_provider="kxy"', '--config', 'model_providers.kxy={name="kxy",base_url="https://api.openai.com/v1",env_key="OPENAI_API_KEY",wire_api="responses"}']
        for attachment in attachments:
            if attachment.suffix.lower() in {'.png', '.jpg', '.jpeg'}:
                argv += ['--image', str(attachment)]
    else:
        argv = [binary, 'run', '--pure', '--format', 'json', '--dir', str(workspace)]
        if model:
            argv += ['--model', model]
        variant = str(config.get('variant') or config.get('effort') or '').strip()
        if variant:
            argv += ['--variant', variant]
    return [*argv, prompt]


def sandbox_profile(workspace, grant, network):
    # Only the node directory is writable. Export grants are never CLI permissions.
    lines = ['(version 1)', '(allow default)', '(deny file-write*)', '(allow file-write* (literal "/dev/null") (literal "/dev/tty") (literal "/dev/urandom"))']
    lines.append('(allow file-write* (subpath ' + json.dumps(str(workspace.resolve())) + '))')
    if not network:
        lines.append('(deny network*)')
    return '\n'.join(lines)


def sandbox_argv(argv, workspace, grant, network, cli):
    if cli == 'codex':
        if not network:
            raise RuntimeError('当前 Codex 云模型执行需要联网；请开启网络。本 MVP 不提供 Codex 离线模型执行。')
        # Codex applies its own workspace sandbox; nesting seatbelt breaks macOS launch.
        return argv, None
    sandbox = shutil.which('sandbox-exec')
    if not sandbox or sys.platform != 'darwin':
        raise RuntimeError('当前 OpenCode 执行需要 macOS sandbox-exec 文件写入隔离')
    profile_file = workspace / '.kxy-seatbelt.sb'
    profile_file.write_text(sandbox_profile(workspace, None, network), encoding='utf-8')
    return [sandbox, '-f', str(profile_file), '--', *argv], profile_file


def extract_final_output(lines: list[str], last_message: Path) -> str:
    if last_message.exists():
        text = last_message.read_text(encoding="utf-8", errors="replace").strip()
        if text:
            if len(text) > MAX_TEXT_BYTES:
                raise RuntimeError("最终回答超过250000字符，请将长报告写入outputs文件并返回摘要")
            return redact(text, max_length=MAX_TEXT_BYTES)
    candidates: list[str] = []
    for line in lines:
        raw = line.strip()
        if not raw:
            continue
        try:
            value = json.loads(raw)
        except json.JSONDecodeError:
            continue
        if not isinstance(value, dict):
            continue
        # OpenCode --format json emits {type: "text", part: {text: ...}}.
        if value.get("type") == "text" and isinstance(value.get("part"), dict):
            text = value["part"].get("text")
            if isinstance(text, str) and text.strip():
                candidates.append(text.strip())
        # Codex --json emits an item.completed event containing an agent_message.
        if value.get("type") in {"item.completed", "message.completed", "agent_message"}:
            item = value.get("item", value)
            if isinstance(item, dict) and item.get("type") in {None, "agent_message", "message"}:
                text = item.get("text") or item.get("content")
                if isinstance(text, str) and text.strip():
                    candidates.append(text.strip())
    if candidates:
        if len(candidates[-1]) > MAX_TEXT_BYTES:
            raise RuntimeError("最终回答超过250000字符，请将长报告写入outputs文件并返回摘要")
        return redact(candidates[-1], max_length=MAX_TEXT_BYTES)
    return ""


@dataclass
class RunState:
    run_id: str
    workspace: Path
    cancel_event: threading.Event
    config: dict[str, Any]
    attempt_no: int = 1
    resume_outputs: dict[str, Any] = field(default_factory=dict)
    node_started: set[str] = field(default_factory=set)
    node_finished: set[str] = field(default_factory=set)
    reused_nodes: set[str] = field(default_factory=set)
    node_paths: dict[str, str] = field(default_factory=dict)
    blackbox_paths: dict[str, dict[str, Any]] = field(default_factory=dict)
    loop_scope: str | None = None
    loop_path: str | None = None
    loop_round_no: int | None = None
    loop_context: dict[str, Any] | None = None
    loop_input_payload: dict[str, Any] | None = None
    loop_budget_seconds: float | None = None
    loop_active_seconds: float = 0.0

    @property
    def node_workspace_root(self) -> Path:
        if self.attempt_no <= 1:
            return self.workspace / "nodes"
        return self.workspace / "attempts" / str(self.attempt_no) / "nodes"

    @property
    def active_node_workspace_root(self) -> Path:
        root = self.node_workspace_root
        if not self.loop_scope:
            return root
        safe_scope = re.sub(r"[^A-Za-z0-9_-]+", "_", str(self.loop_scope)).strip("._") or "loop"
        return root / "loops" / safe_scope[:96]

    def begin_loop_call(self) -> float | None:
        if self.loop_budget_seconds is None:
            return None
        remaining = self.loop_budget_seconds - self.loop_active_seconds
        if remaining <= 0:
            raise LoopBudgetExceeded("循环活动执行预算已用尽，请增加有界预算或停止循环")
        return time.monotonic()

    def finish_loop_call(self, started: float | None) -> None:
        if started is not None:
            self.loop_active_seconds += max(0.0, time.monotonic() - started)

    def restore_node(self, node_id: str) -> tuple[bool, Any]:
        """Restore a durable result once per node without changing its route."""

        if node_id not in self.resume_outputs:
            return False, None
        if node_id not in self.reused_nodes:
            self.reused_nodes.add(node_id)
            self.node_started.add(node_id)
            self.node_finished.add(node_id)
            append_event(
                self.run_id,
                "node_reused",
                {"node_id": node_id, "attempt": self.attempt_no},
            )
        return True, self.resume_outputs[node_id]

    def start_node(self, node_id: str) -> None:
        if node_id in self.node_started:
            return
        self.node_started.add(node_id)
        with connect_db() as db:
            current = db.execute("SELECT status FROM runs WHERE id=?", (self.run_id,)).fetchone()
        if current is not None and current["status"] == "interrupted":
            return
        update_run_node(self.run_id, node_id, "running", started_at=utc_now())
        append_event(self.run_id, "node_started", {"node_id": node_id})

    def finish_node(self, node_id: str, status: str, output: Any = None, message: str | None = None) -> None:
        if node_id in self.node_finished and status != "failed":
            return
        with connect_db() as db:
            current = db.execute("SELECT status FROM runs WHERE id=?", (self.run_id,)).fetchone()
        if current is not None and current["status"] == "interrupted" and status != "interrupted":
            return
        self.node_finished.add(node_id)
        update_run_node(self.run_id, node_id, status, output=output, message=message, finished_at=utc_now())
        append_event(self.run_id, "node_finished", {"node_id": node_id, "status": status, "message": message})

    def set_waiting(self, node_id: str) -> None:
        with connect_db() as db:
            current = db.execute("SELECT status FROM runs WHERE id=?", (self.run_id,)).fetchone()
        if current is not None and current["status"] == "interrupted":
            return
        update_run_node(self.run_id, node_id, "waiting")
        with connect_db() as db:
            db.execute("UPDATE runs SET status='waiting' WHERE id=? AND status IN ('pending','running')", (self.run_id,))
        for path, info in self.blackbox_paths.items():
            if node_id in info.get("inner", []):
                update_run_node(self.run_id, path, "waiting")
        append_event(self.run_id, "human_waiting", {"node_id": node_id, "node_path": self.node_paths.get(node_id, node_id)})

    def resume_running(self, node_id: str) -> None:
        with connect_db() as db:
            current = db.execute("SELECT status FROM runs WHERE id=?", (self.run_id,)).fetchone()
        if current is not None and current["status"] == "interrupted":
            return
        update_run_node(self.run_id, node_id, "running")
        with connect_db() as db:
            db.execute("UPDATE runs SET status='running' WHERE id=? AND status='waiting'", (self.run_id,))
        for path, info in self.blackbox_paths.items():
            if node_id in info.get("inner", []):
                update_run_node(self.run_id, path, "running")
        append_event(self.run_id, "human_resumed", {"node_id": node_id, "node_path": self.node_paths.get(node_id, node_id)})


def append_event(run_id: str, event_type: str, payload: dict[str, Any]) -> None:
    redacted_payload = redact(json_text(payload))
    try:
        safe_payload = json.loads(redacted_payload)
    except json.JSONDecodeError:
        safe_payload = {"text": redacted_payload}
    with connect_db() as db:
        db.execute(
            "INSERT INTO run_events(run_id, event_type, payload, created_at) VALUES (?, ?, ?, ?)",
            (run_id, event_type, json_text(safe_payload), utc_now()),
        )


def update_run_node(
    run_id: str,
    node_id: str,
    status: str,
    *,
    output: Any = None,
    message: str | None = None,
    started_at: str | None = None,
    finished_at: str | None = None,
) -> None:
    with connect_db() as db:
        # Loop rounds create scoped internal checkpoints lazily.  The insert is
        # additive and keeps the ordinary V8 pre-seeded node rows unchanged.
        db.execute(
            "INSERT OR IGNORE INTO run_nodes(run_id, node_id, status) VALUES (?, ?, 'pending')",
            (run_id, node_id),
        )
        db.execute(
            "UPDATE run_nodes SET status=?, message=COALESCE(?, message), output=COALESCE(?, output), "
            "started_at=COALESCE(?, started_at), finished_at=COALESCE(?, finished_at) WHERE run_id=? AND node_id=?",
            (
                status,
                redact(message) if message else None,
                json_text(redact_object(output)) if output is not None else None,
                started_at,
                finished_at,
                run_id,
                node_id,
            ),
        )


def _approval_row(run_id: str, node_id: str) -> sqlite3.Row | None:
    with connect_db() as db:
        return db.execute("SELECT * FROM run_approvals WHERE run_id=? AND node_id=?", (run_id, node_id)).fetchone()


def create_pending_approval(
    run_id: str,
    node_id: str,
    node_path: str,
    content: str,
    inputs: list[dict[str, Any]],
    confirm_label: str = "确认继续",
    allow_return: bool = False,
) -> sqlite3.Row:
    bounded_content = redact(str(content or ""), max_length=MAX_TEXT_BYTES)
    bounded_label = redact(str(confirm_label or "确认继续"), max_length=120) or "确认继续"
    input_json = json_text(redact_object(inputs))
    created = utc_now()
    with connect_db() as db:
        db.execute(
            "INSERT OR IGNORE INTO run_approvals(run_id,node_id,node_path,status,content,confirm_label,inputs,allow_return,created_at) "
            "VALUES (?,?,?,?,?,?,?,?,?)",
            (run_id, node_id, node_path, "pending", bounded_content, bounded_label, input_json, 1 if allow_return else 0, created),
        )
        row = db.execute("SELECT * FROM run_approvals WHERE run_id=? AND node_id=?", (run_id, node_id)).fetchone()
    if row is None:
        raise RuntimeError("无法创建人工确认断点")
    if row["status"] == "pending" and row["created_at"] == created:
        append_event(
            run_id,
            "approval_created",
            {"node_id": node_id, "node_path": node_path, "content": bounded_content},
        )
    return row


def approval_event(run_id: str, node_id: str) -> threading.Event:
    key = (run_id, node_id)
    with RUN_LOCK:
        return RUN_APPROVAL_EVENTS.setdefault(key, threading.Event())


def approval_payload(row: sqlite3.Row) -> dict[str, Any]:
    revision_text = row["revision_text"] if "revision_text" in row.keys() else None
    return {
        "node_id": row["node_path"] or row["node_id"],
        "flat_node_id": row["node_id"],
        "status": row["status"],
        "content": row["content"],
        "confirm_label": row["confirm_label"] or "确认继续",
        "inputs": json.loads(row["inputs"] or "[]"),
        "original_inputs": json.loads(row["inputs"] or "[]"),
        "allow_return": bool(row["allow_return"]) if "allow_return" in row.keys() else False,
        "revision_text": revision_text,
        "revision_note": row["revision_note"] if "revision_note" in row.keys() else None,
        "revision_at": row["revision_at"] if "revision_at" in row.keys() else None,
        "decision": row["decision"],
        "note": row["note"],
        "created_at": row["created_at"],
        "decided_at": row["decided_at"],
    }


def source_payload(file_id, workspace):
    manifest_path = workspace / 'manifest.json'
    if manifest_path.is_file():
        snapshot = json.loads(manifest_path.read_text())
        if file_id in snapshot.get('sources', {}):
            source = snapshot['sources'][file_id]
            original = (workspace / source['file_path']).resolve()
            if not resolved_inside(original, workspace) or hashlib.sha256(original.read_bytes()).hexdigest() != source['sha256']:
                raise ValueError('运行资料快照校验失败')
            return source
    raise ValueError(f'运行快照中缺少资料：{file_id}；请重新绑定资料')


def input_attachment_refs(data: dict[str, Any]) -> list[dict[str, str]]:
    """Read the V13 attachment list while retaining legacy input.file_id."""
    refs: list[dict[str, str]] = []
    raw_attachments = data.get("attachments")
    if isinstance(raw_attachments, list):
        for item in raw_attachments:
            if not isinstance(item, dict):
                continue
            file_id = str(item.get("file_id") or "").strip()
            if not file_id:
                continue
            name = str(item.get("name") or item.get("display_name") or "").strip()
            relative = str(item.get("relative_path") or name or "").strip()
            display = str(item.get("display_name") or name or relative or "").strip()
            refs.append({"file_id": file_id, "name": name, "relative_path": relative, "display_name": display})
    raw_ids = data.get("file_ids")
    if isinstance(raw_ids, str):
        raw_ids = [raw_ids]
    if isinstance(raw_ids, list):
        for raw_id in raw_ids:
            file_id = str(raw_id or "").strip()
            if file_id and not any(item["file_id"] == file_id for item in refs):
                refs.append({"file_id": file_id, "name": "", "relative_path": "", "display_name": ""})
    legacy_id = str(data.get("file_id") or "").strip()
    if legacy_id and not any(item["file_id"] == legacy_id for item in refs):
        refs.insert(0, {"file_id": legacy_id, "name": "", "relative_path": "", "display_name": ""})
    return refs


def safe_relative_display(value: str, fallback: str) -> str:
    raw = str(value or "").replace("\\", "/").replace("\x00", "").strip()
    parts = [safe_display_name(part) for part in raw.split("/") if part not in {"", ".", ".."}]
    return "/".join(parts)[:MAX_DISPLAY_CHARS] or safe_display_name(fallback)


def input_source_payload(data: dict[str, Any], workspace: Path) -> dict[str, Any]:
    text = str(data.get("text") or "")[:MAX_TEXT_BYTES]
    label = str(data.get("label") or "Text input")[:MAX_DISPLAY_CHARS]
    attachments: list[dict[str, Any]] = []
    for ref in input_attachment_refs(data):
        source = json.loads(json_text(source_payload(ref["file_id"], workspace)))
        base_name = str(source.get("name") or ref["file_id"])
        requested_name = ref.get("name") or ref.get("display_name") or base_name
        relative = safe_relative_display(ref.get("relative_path") or requested_name or base_name, base_name)
        display = safe_relative_display(requested_name or relative, base_name)
        source["name"] = display
        source["display_name"] = display
        source["relative_path"] = relative
        source["active"] = True
        attachments.append(source)
    return {
        "status": "source",
        "active": True,
        "name": label,
        "text": text,
        "input_text": text,
        "items": attachments,
        "attachments": [
            {
                "file_id": item.get("file_id"),
                "name": item.get("name") or item.get("display_name"),
                "display_name": item.get("display_name") or item.get("name"),
                "relative_path": item.get("relative_path") or item.get("name"),
                "sha256": item.get("sha256"),
            }
            for item in attachments
        ],
        "file_ids": [str(item.get("file_id")) for item in attachments if item.get("file_id")],
    }


def _skill_source_agent(item: dict[str, Any], target_agent: str) -> str:
    metadata = item.get("metadata") if isinstance(item.get("metadata"), dict) else {}
    values: list[str] = []
    for key in ("source_agents", "native_agents"):
        raw = metadata.get(key, [])
        raw = [raw] if isinstance(raw, str) else raw
        if isinstance(raw, list):
            values.extend(str(value).strip().lower() for value in raw if str(value).strip())
    for key in ("source_agent", "native_agent"):
        if metadata.get(key):
            values.append(str(metadata[key]).strip().lower())
    values = list(dict.fromkeys(values))
    if target_agent in values:
        return target_agent
    return values[0] if values else "local"


def _update_run_snapshot(workspace: Path, update: Callable[[dict[str, Any]], None]) -> dict[str, Any]:
    """Atomically persist runtime mapping metadata in the run manifest and DB."""

    manifest_path = workspace / "manifest.json"
    run_root = workspace.parent
    if not resolved_inside(workspace, RUNS_ROOT) or not resolved_inside(run_root, RUNS_ROOT):
        raise ValueError("运行快照路径无效")
    run_id = run_root.name
    with RUN_LOCK:
        snapshot = json.loads(manifest_path.read_text(encoding="utf-8"))
        if not isinstance(snapshot, dict):
            raise ValueError("运行快照格式无效")
        update(snapshot)
        temporary = manifest_path.with_name(f".manifest.{uuid.uuid4().hex}.tmp")
        temporary.write_text(json_text(snapshot), encoding="utf-8")
        temporary.replace(manifest_path)
        with connect_db() as db:
            db.execute("UPDATE runs SET snapshot=? WHERE id=?", (json_text(snapshot), run_id))
    return snapshot


def mount_skills(skill_ids, node_workspace, target_agent="codex", run_workspace: Path | None = None):
    if not skill_ids:
        return []
    run_workspace = run_workspace or node_workspace.parents[1]
    snapshot = json.loads((run_workspace / 'manifest.json').read_text())
    available = {item['id']: item for item in snapshot.get('skills', [])}
    mounted = []
    for skill_id in skill_ids:
        item = available.get(skill_id)
        if item is None:
            raise ValueError('运行快照中缺少所选 Skill')
        root = run_workspace / item['path']
        if not resolved_inside(root, run_workspace) or snapshot_skill(root)[0] != item['snapshot_hash']:
            raise ValueError('Skill 运行快照校验失败')
        source_agent = _skill_source_agent(item, str(target_agent).strip().lower())
        mapping: dict[str, Any] | None = None
        copy_root = root
        if source_agent != "local":
            try:
                from . import agent_config

                validated_target = agent_config._valid_agent_id(str(target_agent).strip().lower())
                if source_agent not in agent_config.SKILLHUB_SUPPORTED_AGENTS:
                    raise ValueError(f"来源 Agent {source_agent} 没有可验证的 SkillHub projection 适配器")
                mapping = agent_config._skillhub_project_snapshot(
                    root,
                    kxy_skill_id=str(item["id"]),
                    name=str(item["name"]),
                    category=(item.get("metadata") or {}).get("native_category", "") if isinstance(item.get("metadata"), dict) else "",
                    snapshot_hash=str(item["snapshot_hash"]),
                    source_agent=source_agent,
                    target_agent=validated_target,
                    source_path=str((item.get("metadata") or {}).get("native_source") or "") if isinstance(item.get("metadata"), dict) else None,
                    projection_root=node_workspace / "skillhub-projections" / validated_target,
                )
                projection = mapping.get("projection")
                if mapping.get("cross_agent"):
                    if not isinstance(projection, str) or not projection:
                        raise ValueError("SkillHub projection 未返回安全路径")
                    candidate = (DATA_ROOT / projection).absolute()
                    allowed_projection_root = (node_workspace / "skillhub-projections").absolute()
                    try:
                        candidate.relative_to(allowed_projection_root)
                    except ValueError as exc:
                        raise ValueError("SkillHub runtime projection 越过 node workspace") from exc
                    if not candidate.exists() and not candidate.is_symlink():
                        raise ValueError("SkillHub runtime projection 不存在")
                    copy_root = candidate
            except (AttributeError, OSError, ValueError) as exc:
                raise ValueError(f"SkillHub 运行时挂载失败：{exc}") from exc
        destination = node_workspace / 'skills' / f"{item['name']}-{item['snapshot_hash'][:10]}"
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copytree(copy_root, destination, symlinks=False)
        mounted_item = {**item, 'path': str(destination.relative_to(node_workspace)), 'source_agent': source_agent, 'target_agent': str(target_agent).strip().lower()}
        if mapping is not None:
            mounted_item["skillhub_mapping"] = {
                key: mapping.get(key)
                for key in ("source_agent", "target_agent", "cross_agent", "status", "message", "upstream_skill_id", "snapshot_hash", "upstream_commit", "projection")
                if mapping.get(key) is not None
            }
            mapping_key = f"{node_workspace.name}:{item['id']}"
            _update_run_snapshot(
                run_workspace,
                lambda current, key=mapping_key, value=mounted_item["skillhub_mapping"]: current.setdefault("skill_mappings", {}).update({key: value}),
            )
        mounted.append(mounted_item)
    (node_workspace / 'skill-manifest.json').write_text(json_text({'skills': mounted}), encoding='utf-8')
    return mounted


def payload_from_data(value: Any) -> dict[str, Any]:
    if hasattr(value, "data") and isinstance(value.data, dict):
        return value.data
    if isinstance(value, dict):
        return value
    return {"text": str(value)}


def payload_children(value: Any) -> list[dict[str, Any]]:
    """Return only the transport wrappers we own; do not walk arbitrary model JSON."""
    payload = payload_from_data(value)
    children: list[dict[str, Any]] = []
    for key in ("items", "upstream", "inputs"):
        child = payload.get(key)
        if isinstance(child, list):
            children.extend(payload_from_data(item) for item in child)
        elif isinstance(child, dict):
            children.append(payload_from_data(child))
    return children


def _filter_candidate_name(value: dict[str, Any], scope: str) -> str:
    relative = str(value.get("relative_path") or value.get("display_name") or value.get("name") or value.get("path") or "")
    relative = relative.replace("\\", "/").strip().lstrip("/")
    if scope == "basename":
        logical_name = str(value.get("name") or value.get("display_name") or relative)
        return Path(logical_name.replace("\\", "/")).name
    return relative


def _filter_candidate_key(value: dict[str, Any], kind: str) -> str:
    identity = value.get("file_id") or value.get("path") or value.get("relative_path") or value.get("name")
    digest = value.get("sha256") or value.get("size") or ""
    return f"{kind}:{identity}:{digest}:{value.get('relative_path') or value.get('name') or ''}"


def _filter_summary(value: dict[str, Any], *, matched: bool, scope: str) -> dict[str, Any]:
    name = _filter_candidate_name(value, scope)
    summary: dict[str, Any] = {
        "name": name or "untitled",
        "relative_path": str(value.get("relative_path") or value.get("name") or name or ""),
        "matched": matched,
    }
    for key in ("file_id", "display_name", "mime", "size", "sha256", "source_status"):
        if value.get(key) not in (None, ""):
            summary[key] = value[key]
    return summary


def _safe_regex_search_batch(pattern: str, candidates: list[str], timeout_ms: int) -> list[bool]:
    """Run one bounded regex batch in a killable child."""
    if any(len(candidate) > FILTER_MAX_PATTERN_CHARS * 4 for candidate in candidates):
        raise ValueError("filter candidate path is too long")
    code = (
        "import json,re,sys; p=json.load(sys.stdin); "
        "print(json.dumps([bool(re.search(p['pattern'], c)) for c in p['candidates']]))"
    )
    process = subprocess.Popen(
        [sys.executable, "-c", code],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        start_new_session=True,
    )
    try:
        stdout, stderr = process.communicate(
            json.dumps({"pattern": pattern, "candidates": candidates}, ensure_ascii=False),
            timeout=max(0.25, timeout_ms / 1000),
        )
    except subprocess.TimeoutExpired as exc:
        try:
            os.killpg(process.pid, signal.SIGKILL)
        except OSError:
            process.kill()
        process.communicate()
        raise ValueError(f"filter regex timed out after {timeout_ms} ms") from exc
    if process.returncode != 0:
        detail = stderr.strip()[:240] if stderr else "regex worker failed"
        raise ValueError(f"filter regex execution failed: {detail}")
    try:
        result = json.loads(stdout.strip() or "[]")
        if not isinstance(result, list) or len(result) != len(candidates) or any(type(item) is not bool for item in result):
            raise ValueError("filter regex worker returned invalid output")
        return result
    except (TypeError, ValueError, json.JSONDecodeError) as exc:
        raise ValueError("filter regex worker returned invalid output") from exc


def _filter_name_matches(config: dict[str, Any], candidate: str) -> bool:
    pattern = str(config["name_pattern"])
    if config["pattern_mode"] == "glob":
        return fnmatch.fnmatchcase(candidate.casefold(), pattern.casefold())
    return _safe_regex_search_batch(pattern, [candidate], int(config["timeout_ms"]))[0]


def filter_payload_items(items: Any, config: dict[str, Any] | None = None) -> dict[str, Any]:
    """Filter owned source/artifact transport envelopes without leaking excluded bodies."""
    normalized = normalize_filter_config(config)
    raw_values = items if isinstance(items, list) else [items]
    values = [json.loads(json_text(payload_from_data(value))) for value in raw_values if value is not None]
    candidates: dict[str, dict[str, Any]] = {}
    candidate_order: list[str] = []

    def is_source(value: dict[str, Any]) -> bool:
        return bool(value.get("file_id") or value.get("file_path")) and (
            value.get("status") in {"source", "filtered"}
            or value.get("source_status") is not None
            or value.get("file_path") is not None
        )

    def register(value: dict[str, Any], kind: str) -> None:
        if kind == "source" and not is_source(value):
            return
        key = _filter_candidate_key(value, kind)
        if key not in candidates:
            candidates[key] = {"key": key, "kind": kind, "value": value}
            candidate_order.append(key)

    def collect(value: Any) -> None:
        if isinstance(value, list):
            for child in value:
                collect(child)
            return
        if not isinstance(value, dict):
            return
        if is_source(value):
            register(value, "source")
        artifacts = value.get("artifacts")
        if isinstance(artifacts, list):
            for artifact in artifacts:
                if isinstance(artifact, dict) and (artifact.get("path") or artifact.get("name")):
                    register(artifact, "artifact")
        for key in ("items", "upstream", "inputs", "attachments"):
            child = value.get(key)
            if isinstance(child, (list, dict)):
                collect(child)

    for value in values:
        collect(value)
    if len(candidates) > FILTER_MAX_CANDIDATES:
        raise ValueError(f"filter input contains more than {FILTER_MAX_CANDIDATES} attachments")
    if not candidates:
        raise ValueError("filter found no attachments to filter")

    matched_keys: set[str] = set()
    matched: list[dict[str, Any]] = []
    excluded: list[dict[str, Any]] = []
    extensions = set(normalized["extensions"])
    candidate_names = [
        _filter_candidate_name(candidates[key]["value"], normalized["name_scope"])
        for key in candidate_order
    ]
    if any(len(name) > FILTER_MAX_PATTERN_CHARS * 4 for name in candidate_names):
        raise ValueError("filter candidate path is too long")
    regex_results = (
        _safe_regex_search_batch(normalized["name_pattern"], candidate_names, int(normalized["timeout_ms"]))
        if normalized["pattern_mode"] == "regex"
        else []
    )
    for index, key in enumerate(candidate_order):
        candidate = candidates[key]
        value = candidate["value"]
        name = candidate_names[index]
        extension = Path(str(value.get("relative_path") or value.get("name") or value.get("path") or "")).suffix.lower()
        extension_ok = "*" in extensions or extension in extensions
        name_ok = regex_results[index] if regex_results else _filter_name_matches(normalized, name)
        candidate["matched"] = bool(extension_ok and name_ok)
        if candidate["matched"]:
            matched_keys.add(key)
            matched.append(_filter_summary(value, matched=True, scope=normalized["name_scope"]))
        else:
            excluded.append(_filter_summary(value, matched=False, scope=normalized["name_scope"]))
    if not matched_keys:
        raise ValueError(f"filter matched no attachments (excluded {len(excluded)})")

    def historical_reference(value: Any) -> Any:
        if isinstance(value, list):
            return [historical_reference(child) for child in value]
        if not isinstance(value, dict):
            return value
        result: dict[str, Any] = {}
        refs: list[Any] = []
        for key, child in value.items():
            if key in {"text", "content", "structured", "file_path", "path", "items", "upstream", "inputs", "artifacts"}:
                if key in {"items", "upstream", "inputs"} and isinstance(child, (list, dict)):
                    nested = child if isinstance(child, list) else [child]
                    refs.extend(historical_reference(item) for item in nested)
                continue
            result[key] = historical_reference(child)
        if refs:
            result["source_refs"] = refs
        result["execution_only"] = False
        return result

    def visible_content(value: Any) -> Any:
        if not isinstance(value, dict):
            return value
        if value.get("status") == "filtered" and value.get("reason") == "excluded_by_filter":
            return None
        for key in ("upstream", "items", "inputs"):
            child = value.get(key)
            if isinstance(child, list) and child:
                values = [visible_content(item) for item in child]
                values = [item for item in values if item is not None]
                return values[0] if len(values) == 1 else values
            if isinstance(child, dict):
                return visible_content(child)
        if "content" in value:
            return value.get("content")
        if "structured" in value:
            return value.get("structured")
        return value.get("text")

    def transform(value: Any) -> Any:
        if isinstance(value, list):
            return [transform(child) for child in value]
        if not isinstance(value, dict):
            return value
        if is_source(value):
            key = _filter_candidate_key(value, "source")
            if key not in matched_keys:
                return {
                    **_filter_summary(value, matched=False, scope=normalized["name_scope"]),
                    "status": "filtered",
                    "active": True,
                    "reason": "excluded_by_filter",
                }
        result: dict[str, Any] = {}
        for key, child in value.items():
            if key == "file_path" and is_source(value) and _filter_candidate_key(value, "source") not in matched_keys:
                continue
            if key == "artifacts" and isinstance(child, list):
                kept: list[Any] = []
                for artifact in child:
                    if not isinstance(artifact, dict):
                        continue
                    artifact_key = _filter_candidate_key(artifact, "artifact")
                    if artifact_key in matched_keys:
                        kept.append(transform(artifact))
                    elif artifact_key in candidates:
                        continue
                    else:
                        kept.append(transform(artifact))
                result[key] = kept
                continue
            if key == "excluded_artifacts" and isinstance(child, list):
                result[key] = [
                    _filter_summary(artifact, matched=False, scope=normalized["name_scope"])
                    for artifact in child
                    if isinstance(artifact, dict)
                ]
                continue
            if key == "original_result":
                result[key] = historical_reference(child)
                continue
            if key in {"items", "upstream", "inputs", "attachments"} and isinstance(child, (list, dict)):
                transformed = transform(child)
                result[key] = transformed
                continue
            result[key] = child
        owned_children = [result.get(key) for key in ("items", "upstream", "inputs") if key in result]
        if owned_children:
            child_values: list[Any] = []
            for child in owned_children:
                child_values.extend(child if isinstance(child, list) else [child])
            visible_text = [str(child.get("text")) for child in child_values if isinstance(child, dict) and child.get("text")]
            input_text = result.get("input_text")
            if isinstance(input_text, str) and input_text:
                visible_text.insert(0, input_text)
            result["text"] = "\n\n".join(visible_text)
            if "content" in result:
                content_values = [visible_content(child) for child in child_values]
                content_values = [item for item in content_values if item is not None]
                result["content"] = content_values[0] if len(content_values) == 1 else content_values
            if "structured" in result:
                structured_values = [visible_content(child) for child in child_values]
                structured_values = [item for item in structured_values if item is not None]
                result["structured"] = structured_values[0] if len(structured_values) == 1 else structured_values
        return result

    filtered_values = [transform(value) for value in values]
    text_parts = [str(value.get("text")) for value in filtered_values if isinstance(value, dict) and value.get("text")]
    return {
        "status": "filter",
        "active": True,
        "items": filtered_values,
        "upstream": filtered_values,
        "text": "\n\n".join(text_parts),
        "matched": matched,
        "excluded": excluded,
        "matched_count": len(matched),
        "excluded_count": len(excluded),
        "filter": normalized,
    }


def loop_scope_id(loop_path: str, round_no: int) -> str:
    exact = f"{loop_path}\x00{round_no}"
    raw = re.sub(r"[^A-Za-z0-9_-]+", "_", f"{loop_path}__round-{round_no}").strip("._") or "loop"
    digest = hashlib.sha256(exact.encode("utf-8")).hexdigest()[:16]
    return f"{raw[:54]}__{digest}"


def scoped_component_id(scope: str, flat_id: str) -> str:
    if not scope:
        return flat_id
    candidate = f"{scope}__{flat_id}"
    if len(candidate) <= 96:
        return candidate
    digest = hashlib.sha256(candidate.encode("utf-8")).hexdigest()[:20]
    return f"{candidate[:75]}__{digest}"


def _loop_json(value: Any) -> str:
    return json_text(redact_object(value))


def ensure_loop_round(
    run_id: str,
    attempt_no: int,
    loop_path: str,
    round_no: int,
    *,
    goal: str,
    original_input: Any,
    previous_summary: str,
    executor_node_id: str,
    reviewer_node_id: str,
    max_rounds: int = LOOP_DEFAULT_MAX_ROUNDS,
    budget_seconds: float = LOOP_DEFAULT_ACTIVE_BUDGET_SECONDS,
) -> None:
    now = utc_now()
    with connect_db() as db:
        db.execute(
            "INSERT OR IGNORE INTO loop_rounds(" 
            "run_id,attempt_no,loop_path,round_no,status,original_goal,original_input,previous_summary," 
            "executor_node_id,reviewer_node_id,effective_max_rounds,budget_seconds,created_at,started_at) "
            "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (
                run_id,
                attempt_no,
                loop_path,
                round_no,
                "running",
                redact(goal, max_length=MAX_TEXT_BYTES),
                _loop_json(original_input),
                redact(previous_summary, max_length=LOOP_MAX_PREVIOUS_SUMMARY_CHARS),
                executor_node_id,
                reviewer_node_id,
                max_rounds,
                budget_seconds,
                now,
                now,
            ),
        )
        db.execute(
            "UPDATE loop_rounds SET status='running', error=NULL, started_at=COALESCE(started_at, ?), "
            "finished_at=NULL WHERE run_id=? AND attempt_no=? AND loop_path=? AND round_no=? "
            "AND status IN ('interrupted','budget_waiting','round_limit_waiting')",
            (now, run_id, attempt_no, loop_path, round_no),
        )


def update_loop_round(run_id: str, attempt_no: int, loop_path: str, round_no: int, **updates: Any) -> None:
    allowed = {
        "status", "original_goal", "original_input", "previous_summary", "executor_node_id",
        "executor_input", "executor_output", "reviewer_node_id", "reviewer_input", "reviewer_output",
        "review_passed", "review_issues", "next_action", "active_seconds", "effective_max_rounds",
        "budget_seconds", "node_outputs", "node_statuses", "control", "control_payload", "error",
        "started_at", "finished_at",
    }
    fields = [key for key in updates if key in allowed]
    if not fields:
        return
    values: list[Any] = []
    json_fields = {
        "original_input", "executor_input", "executor_output", "reviewer_input", "reviewer_output",
        "review_issues", "node_outputs", "node_statuses", "control_payload",
    }
    for key in fields:
        value = updates[key]
        if key in json_fields:
            value = _loop_json(value)
        elif key in {"original_goal", "previous_summary", "next_action", "error"} and value is not None:
            value = redact(str(value), max_length=MAX_TEXT_BYTES if key == "original_goal" else LOOP_MAX_PREVIOUS_SUMMARY_CHARS)
        values.append(value)
    values.extend([run_id, attempt_no, loop_path, round_no])
    with connect_db() as db:
        db.execute(
            f"UPDATE loop_rounds SET {', '.join(f'{key}=?' for key in fields)} "
            "WHERE run_id=? AND attempt_no=? AND loop_path=? AND round_no=?",
            values,
        )


def loop_round_rows(run_id: str, loop_path: str | None = None) -> list[sqlite3.Row]:
    with connect_db() as db:
        if loop_path is None:
            return db.execute(
                "SELECT * FROM loop_rounds WHERE run_id=? ORDER BY attempt_no, loop_path, round_no",
                (run_id,),
            ).fetchall()
        return db.execute(
            "SELECT * FROM loop_rounds WHERE run_id=? AND loop_path=? ORDER BY attempt_no, round_no",
            (run_id, loop_path),
        ).fetchall()


def loop_round_active_seconds(run_id: str, attempt_no: int, loop_path: str, round_no: int) -> float:
    with connect_db() as db:
        row = db.execute(
            "SELECT active_seconds FROM loop_rounds WHERE run_id=? AND attempt_no=? AND loop_path=? AND round_no=?",
            (run_id, attempt_no, loop_path, round_no),
        ).fetchone()
    return float(row["active_seconds"] or 0) if row is not None else 0.0


def loop_round_json(row: sqlite3.Row, key: str, fallback: Any) -> Any:
    try:
        value = json.loads(row[key] or "")
    except (TypeError, ValueError, json.JSONDecodeError):
        return fallback
    return value


def loop_round_payload(row: sqlite3.Row) -> dict[str, Any]:
    return {
        "run_id": row["run_id"],
        "attempt_no": row["attempt_no"],
        "loop_path": row["loop_path"],
        "round_no": row["round_no"],
        "status": row["status"],
        "original_goal": row["original_goal"],
        "original_input": loop_round_json(row, "original_input", {}),
        "previous_summary": row["previous_summary"],
        "executor_node_id": row["executor_node_id"],
        "executor_input": loop_round_json(row, "executor_input", {}),
        "executor_output": loop_round_json(row, "executor_output", None),
        "reviewer_node_id": row["reviewer_node_id"],
        "reviewer_input": loop_round_json(row, "reviewer_input", {}),
        "reviewer_output": loop_round_json(row, "reviewer_output", None),
        "review_passed": None if row["review_passed"] is None else bool(row["review_passed"]),
        "review_issues": loop_round_json(row, "review_issues", []),
        "next_action": row["next_action"],
        "active_seconds": float(row["active_seconds"] or 0),
        "effective_max_rounds": int(row["effective_max_rounds"] or LOOP_DEFAULT_MAX_ROUNDS),
        "budget_seconds": float(row["budget_seconds"] or LOOP_DEFAULT_ACTIVE_BUDGET_SECONDS),
        "node_outputs": loop_round_json(row, "node_outputs", {}),
        "node_statuses": loop_round_json(row, "node_statuses", {}),
        "control": row["control"],
        "control_payload": loop_round_json(row, "control_payload", {}),
        "error": row["error"],
        "created_at": row["created_at"],
        "started_at": row["started_at"],
        "finished_at": row["finished_at"],
    }


def loop_text(value: Any, limit: int = LOOP_MAX_PREVIOUS_SUMMARY_CHARS) -> str:
    payload = payload_from_data(value)
    text = payload.get("text")
    if not isinstance(text, str) or not text.strip():
        try:
            text = json_text(payload)
        except (TypeError, ValueError):
            text = str(payload)
    return redact(text, max_length=max(0, limit))


def parse_loop_review(payload: dict[str, Any], settings: dict[str, Any], loop_path: str) -> dict[str, Any]:
    fields = settings.get("review_fields") if isinstance(settings.get("review_fields"), dict) else {}
    selected: Any = payload.get("structured")
    if selected is None and isinstance(payload.get("content"), dict):
        selected = payload.get("content")
    if selected is None and isinstance(payload.get("text"), str):
        try:
            selected = json.loads(payload["text"])
        except (TypeError, ValueError, json.JSONDecodeError):
            selected = None
    if not isinstance(selected, dict):
        raise ValueError(
            f"Loop {loop_path} reviewer returned no JSON object; "
            "fix the reviewer output or edit the strict JSON template"
        )
    def selected_field(name: str, fallback: str) -> Any:
        exists, value = nested_value(selected, str(fields.get(name) or fallback))
        return value if exists else None
    passed = selected_field("passed", "passed")
    issues = selected_field("issues", "issues")
    next_action = selected_field("next_action", "next_action")
    if type(passed) is not bool:
        raise ValueError(f"Loop {loop_path} reviewer JSON must contain boolean passed")
    if not isinstance(issues, list) or any(not isinstance(item, str) for item in issues):
        raise ValueError(f"Loop {loop_path} reviewer JSON must contain string array issues")
    if not isinstance(next_action, str):
        raise ValueError(f"Loop {loop_path} reviewer JSON must contain string next_action")
    return {
        "passed": passed,
        "issues": [redact(item, max_length=MAX_EVENT_TEXT) for item in issues[:100]],
        "next_action": redact(next_action, max_length=MAX_EVENT_TEXT),
        "source": "selected_reviewer_output",
    }


def loop_resume_outputs(run_id: str, scope: str, component_ids: list[str]) -> dict[str, Any]:
    if not component_ids:
        return {}
    placeholders = ",".join("?" for _ in component_ids)
    with connect_db() as db:
        rows = db.execute(
            f"SELECT node_id,status,output FROM run_nodes WHERE run_id=? AND node_id IN ({placeholders})",
            [run_id, *component_ids],
        ).fetchall()
    result: dict[str, Any] = {}
    for row in rows:
        if row["status"] != "succeeded" or not row["output"]:
            continue
        try:
            output = json.loads(row["output"])
        except (TypeError, ValueError, json.JSONDecodeError):
            raise ValueError(f"循环节点 {row['node_id']} 的成功 checkpoint 不是合法 JSON")
        try:
            validate_checkpoint_output((RUNS_ROOT / str(run_id) / "workspace").resolve(), output, row["node_id"])
        except ValueError as exc:
            raise ValueError(f"循环节点 {row['node_id']} 的 checkpoint 校验失败：{exc}") from exc
        result[str(row["node_id"])] = output
    return result


def loop_node_checkpoint_map(run_id: str, component_ids: list[str]) -> tuple[dict[str, Any], dict[str, str]]:
    if not component_ids:
        return {}, {}
    placeholders = ",".join("?" for _ in component_ids)
    with connect_db() as db:
        rows = db.execute(
            f"SELECT node_id,status,output FROM run_nodes WHERE run_id=? AND node_id IN ({placeholders})",
            [run_id, *component_ids],
        ).fetchall()
    outputs: dict[str, Any] = {}
    statuses: dict[str, str] = {}
    for row in rows:
        node_id = str(row["node_id"])
        statuses[node_id] = str(row["status"])
        if row["output"]:
            try:
                outputs[node_id] = json.loads(row["output"])
            except (TypeError, ValueError, json.JSONDecodeError):
                continue
    return outputs, statuses


async def wait_for_loop_control(state: RunState, loop_path: str, round_no: int, reason: str) -> dict[str, Any]:
    update_loop_round(
        state.run_id,
        state.attempt_no,
        loop_path,
        round_no,
        status="budget_waiting" if reason == "active_budget" else "round_limit_waiting",
        error=reason,
        finished_at=None,
    )
    with connect_db() as db:
        db.execute(
            "UPDATE runs SET status='waiting', error=? WHERE id=? AND status IN ('pending','running','waiting')",
            ("Loop paused: " + reason, state.run_id),
        )
    append_event(
        state.run_id,
        "loop_waiting",
        {"loop_path": loop_path, "round": round_no, "reason": reason, "active_seconds": state.loop_active_seconds},
    )
    while True:
        if state.cancel_event.is_set():
            raise RunCancelled("Run cancelled by user")
        with connect_db() as db:
            row = db.execute(
                "SELECT control,control_payload FROM loop_rounds WHERE run_id=? AND attempt_no=? AND loop_path=? AND round_no=?",
                (state.run_id, state.attempt_no, loop_path, round_no),
            ).fetchone()
        if row is not None and row["control"]:
            control = str(row["control"])
            payload = loop_round_json(row, "control_payload", {})
            update_loop_round(
                state.run_id,
                state.attempt_no,
                loop_path,
                round_no,
                control=None,
                status="running",
                error=None,
            )
            with connect_db() as db:
                db.execute(
                    "UPDATE runs SET status='running',error=NULL WHERE id=? AND status='waiting'",
                    (state.run_id,),
                )
            if control == "stop":
                raise LoopStopped("用户停止了未通过的循环；没有输出被自动接受")
            if control == "continue":
                append_event(
                    state.run_id,
                    "loop_continued",
                    {"loop_path": loop_path, "round": round_no, "control": payload},
                )
                return payload if isinstance(payload, dict) else {}
        await asyncio.to_thread(time.sleep, HUMAN_POLL_SECONDS)


def _attach_task_mcp_environment(agent_id: str, env: dict[str, str], node_workspace: Path, mcp_config: dict[str, Any]) -> None:
    """Expose only the selected, workspace-local MCP config to the child."""

    candidate_path = Path(str(mcp_config.get("path") or ""))
    if candidate_path.is_symlink():
        raise RuntimeError("MCP task-local configuration must not be a symlink")
    path = candidate_path.resolve()
    if not resolved_inside(path, node_workspace) or not path.is_file():
        raise RuntimeError("MCP task-local configuration path is invalid")
    env["KXY_MCP_CONFIG"] = str(path)
    if agent_id == "opencode":
        try:
            local_config = json.loads(env.get("OPENCODE_CONFIG_CONTENT", "{}"))
            selected_config = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, ValueError, json.JSONDecodeError) as exc:
            raise RuntimeError("OpenCode task-local MCP configuration is invalid") from exc
        if not isinstance(local_config, dict) or not isinstance(selected_config, dict):
            raise RuntimeError("OpenCode task-local MCP configuration must be an object")
        local_config["mcp"] = selected_config.get("mcp", {})
        env["OPENCODE_CONFIG_CONTENT"] = json_text(local_config)
    elif agent_id == "workbuddy":
        state_root = Path(env.get("WORKBUDDY_CONFIG_DIR") or node_workspace / ".cli-state" / "workbuddy")
        if not resolved_inside(state_root, node_workspace) or state_root.is_symlink():
            raise RuntimeError("WorkBuddy task-local config directory is invalid")
        state_root.mkdir(parents=True, exist_ok=True)
        target = state_root / "mcp.json"
        shutil.copyfile(path, target)
        try:
            target.chmod(0o600)
        except OSError:
            pass
    elif agent_id == "claude":
        # Claude receives the selected file through its documented --mcp-config
        # argv flags below; keep the generic reference for diagnostics only.
        state_root = Path(env.get("CLAUDE_CONFIG_DIR") or node_workspace / ".cli-state" / "claude")
        if not resolved_inside(state_root, node_workspace) or state_root.is_symlink():
            raise RuntimeError("Claude task-local config directory is invalid")
    elif agent_id == "codex":
        # Codex is invoked with --ignore-user-config, so its selected entries
        # are supplied as explicit -c overrides below.  The audited TOML file
        # remains the source artifact for parsing and replay.
        pass


def _codex_mcp_overrides(path: Path) -> list[str]:
    def toml_value(value: Any) -> str:
        if isinstance(value, bool):
            return "true" if value else "false"
        if isinstance(value, str):
            return json.dumps(value, ensure_ascii=False)
        if isinstance(value, list):
            return "[" + ",".join(toml_value(item) for item in value) + "]"
        if isinstance(value, dict):
            return "{" + ",".join(
                f"{json.dumps(str(key), ensure_ascii=False)}={toml_value(item)}"
                for key, item in value.items()
            ) + "}"
        raise RuntimeError("Codex task-local MCP value has an unsupported TOML type")

    try:
        import tomllib

        document = tomllib.loads(path.read_text(encoding="utf-8"))
    except (ImportError, OSError, UnicodeError, ValueError, tomllib.TOMLDecodeError) as exc:
        raise RuntimeError("Codex task-local MCP TOML is invalid") from exc
    servers = document.get("mcp_servers")
    if not isinstance(servers, dict):
        raise RuntimeError("Codex task-local MCP TOML has no mcp_servers table")
    overrides: list[str] = []
    for server_id, item in servers.items():
        if not isinstance(item, dict):
            raise RuntimeError("Codex task-local MCP entry is invalid")
        fields: list[str] = []
        for key in ("command", "url"):
            if key in item:
                fields.append(f"{key}={json.dumps(str(item[key]), ensure_ascii=False)}")
        if "args" in item:
            fields.append("args=" + toml_value(item["args"]))
        if "enabled" in item:
            fields.append(f"enabled={'true' if item['enabled'] else 'false'}")
        if "env_vars" in item:
            fields.append("env_vars=" + toml_value(item["env_vars"]))
        if "default_tools_approval_mode" in item:
            fields.append(
                "default_tools_approval_mode="
                + toml_value(item["default_tools_approval_mode"])
            )
        for nested_key in ("env", "http_headers", "env_http_headers"):
            nested = item.get(nested_key)
            if isinstance(nested, dict) and nested:
                fields.append(f"{nested_key}=" + toml_value(nested))
        if not fields:
            raise RuntimeError("Codex task-local MCP entry has no runnable fields")
        key = str(server_id) if re.fullmatch(r"[A-Za-z0-9_-]+", str(server_id)) else json.dumps(str(server_id), ensure_ascii=False)
        overrides.extend(["--config", f"mcp_servers.{key}={{" + ",".join(fields) + "}"])
    return overrides


def passthrough_payload(items: Any, marker: str) -> dict[str, Any]:
    """Preserve a single structured upstream payload without forcing items.0 paths."""
    values = active_payloads(items)
    if not values:
        return {"status": "skipped", "active": False, "marker": marker, "items": [], "text": ""}
    if len(values) == 1:
        result = json.loads(json_text(values[0]))
        result.setdefault("provenance", {})
        if isinstance(result["provenance"], dict):
            result["provenance"].setdefault("transport", marker)
        return result
    text = "\n\n".join(str(item.get("text", "")) for item in values if item.get("text"))
    return {"status": marker, "active": True, "items": values, "upstream": values, "text": text}


def human_revision_payload(value: Any, revision_text: str) -> dict[str, Any]:
    """Keep executable transport refs while replacing every visible primary body."""
    values = value if isinstance(value, list) else [value]
    source_refs: dict[str, dict[str, Any]] = {}
    artifact_refs: dict[str, dict[str, Any]] = {}

    def visit(candidate: Any) -> None:
        if isinstance(candidate, list):
            for child in candidate:
                visit(child)
            return
        if not isinstance(candidate, dict):
            return
        file_path = candidate.get("file_path")
        if isinstance(file_path, str) and file_path:
            source_key = str(candidate.get("file_id") or file_path or candidate.get("relative_path") or candidate.get("name"))
            source_refs.setdefault(
                source_key,
                {
                    key: candidate[key]
                    for key in (
                        "file_id", "file_path", "display_name", "name", "relative_path", "sha256",
                        "mime", "size", "status", "source_status", "source_type", "active",
                    )
                    if candidate.get(key) not in (None, "")
                },
            )
        artifacts = candidate.get("artifacts")
        if isinstance(artifacts, list):
            for artifact in artifacts:
                if not isinstance(artifact, dict) or not isinstance(artifact.get("path"), str) or not artifact["path"]:
                    continue
                artifact_key = str(artifact.get("sha256") or artifact["path"])
                artifact_refs.setdefault(
                    artifact_key,
                    {
                        key: artifact[key]
                        for key in ("path", "name", "mime", "size", "sha256", "status")
                        if artifact.get(key) not in (None, "")
                    },
                )
        for key in ("items", "upstream", "inputs", "attachments"):
            visit(candidate.get(key))

    visit(values)
    revised: dict[str, Any] = {
        "status": "human_revision",
        "active": True,
        "text": revision_text,
        "content": revision_text,
    }
    if source_refs:
        revised["attachments"] = list(source_refs.values())
    if artifact_refs:
        revised["artifacts"] = list(artifact_refs.values())
    return revised


if Component is not None:

    class KxySourceComponent(Component):
        display_name = "KXY Source"
        name = "KxySource"
        # `gate` is only used by compiled blackbox-internal source roots. It is
        # optional so the public V1 source nodes retain their original shape.
        inputs = [
            DataInput(name="payload", display_name="Payload", required=False),
            DataInput(name="gate", display_name="Blackbox gate", required=False),
        ]
        outputs = [Output(name="result", display_name="Source", method="build_result")]

        def build_result(self) -> Data:
            state: RunState = self.run_state
            restored, payload = state.restore_node(self.get_id())
            if restored:
                self.last_result = Data(data=payload)
                return self.last_result
            state.start_node(self.get_id())
            if self.gate is not None and not active_payloads(self.gate):
                output = {"status": "skipped", "active": False, "reason": "Blackbox input is inactive", "text": ""}
                self.last_result = Data(data=output)
                state.finish_node(self.get_id(), "skipped", output=output)
                return self.last_result
            self.last_result = self.payload if isinstance(self.payload, Data) else Data(data=payload_from_data(self.payload))
            state.finish_node(self.get_id(), "succeeded", output=payload_from_data(self.last_result))
            return self.last_result


    class KxySubflowInputComponent(Component):
        display_name = "KXY Subflow Input"
        name = "KxySubflowInput"
        inputs = [DataInput(name="items", display_name="Inputs", is_list=True, required=False)]
        outputs = [Output(name="result", display_name="Subflow Input", method="build_result")]

        def build_result(self) -> Data:
            state: RunState = self.run_state
            restored, payload = state.restore_node(self.get_id())
            if restored:
                self.last_result = Data(data=payload)
                return self.last_result
            state.start_node(self.get_id())
            # A box with no external edge is a valid source workflow. Seed it
            # once with an active empty envelope; an explicitly skipped upstream
            # remains inactive and is never turned into implicit approval/data.
            if self.items is None and state.loop_input_payload is not None:
                output = passthrough_payload([state.loop_input_payload], "subflow_input")
            elif self.items is None:
                output = {"status": "subflow_input", "active": True, "items": [], "text": ""}
            else:
                output = passthrough_payload(self.items, "subflow_input")
            self.last_result = Data(data=output)
            state.finish_node(self.get_id(), "skipped" if output.get("status") == "skipped" else "succeeded", output=output)
            return self.last_result


    class KxySubflowOutputComponent(Component):
        display_name = "KXY Subflow Output"
        name = "KxySubflowOutput"
        inputs = [DataInput(name="items", display_name="Inputs", is_list=True, required=False)]
        outputs = [Output(name="result", display_name="Subflow Output", method="build_result")]

        def build_result(self) -> Data:
            state: RunState = self.run_state
            restored, payload = state.restore_node(self.get_id())
            if restored:
                self.last_result = Data(data=payload)
                return self.last_result
            state.start_node(self.get_id())
            output = passthrough_payload(self.items, "subflow_output")
            self.last_result = Data(data=output)
            state.finish_node(self.get_id(), "skipped" if output.get("status") == "skipped" else "succeeded", output=output)
            return self.last_result


    class KxyFilterComponent(Component):
        display_name = "KXY Input Filter"
        name = "KxyFilter"
        inputs = [DataInput(name="items", display_name="Inputs", is_list=True, required=False)]
        outputs = [Output(name="result", display_name="Filtered inputs", method="build_result")]

        def build_result(self) -> Data:
            state: RunState = self.run_state
            node_id = str(self.get_id())
            restored, payload = state.restore_node(node_id)
            if restored:
                self.last_result = Data(data=payload)
                return self.last_result
            state.start_node(node_id)
            try:
                output = filter_payload_items(self.items, getattr(self, "node_config", {}) or {})
                self.last_result = Data(data=output)
                state.finish_node(node_id, "succeeded", output=output)
                return self.last_result
            except Exception as exc:
                state.finish_node(node_id, "failed", message=str(exc))
                raise


    class KxyHumanComponent(Component):
        display_name = "KXY Human Approval"
        name = "KxyHuman"
        inputs = [
            DataInput(name="items", display_name="Inputs", is_list=True, required=False),
            MultilineInput(name="content", display_name="Confirmation content", value=""),
            StrInput(name="confirm_label", display_name="Confirm label", value="确认继续"),
        ]
        outputs = [Output(name="result", display_name="Approved result", method="build_result")]

        async def build_result(self) -> Data:
            state: RunState = self.run_state
            node_id = str(self.get_id())
            restored, payload = state.restore_node(node_id)
            if restored:
                self.last_result = Data(data=payload)
                return self.last_result
            state.start_node(node_id)
            values = active_payloads(self.items)
            if not values:
                output = {"status": "skipped", "active": False, "reason": "No active upstream result", "text": ""}
                self.last_result = Data(data=output)
                state.finish_node(node_id, "skipped", output=output)
                return self.last_result
            if bool(getattr(self, "allow_return", False)) and isinstance(state.loop_context, dict):
                executor_component_id = str(state.loop_context.get("executor_component_id") or "")
                if not executor_component_id:
                    raise RuntimeError("人工闸门缺少当前循环 executor checkpoint")
                with connect_db() as db:
                    executor_row = db.execute(
                        "SELECT status, output FROM run_nodes WHERE run_id=? AND node_id=?",
                        (state.run_id, executor_component_id),
                    ).fetchone()
                if executor_row is None or executor_row["status"] != "succeeded" or not executor_row["output"]:
                    raise RuntimeError("人工闸门只能审阅当前轮已完成的 executor 结果")
                try:
                    executor_payload = json.loads(executor_row["output"])
                except (TypeError, ValueError, json.JSONDecodeError) as exc:
                    raise RuntimeError("当前循环 executor checkpoint 不是合法 JSON") from exc
                if not isinstance(executor_payload, dict):
                    raise RuntimeError("当前循环 executor checkpoint 不是对象")
                current_round = state.loop_context.get("round")
                values = [
                    {**executor_payload, "role": "executor", "round": current_round},
                    *[
                        {**item, "role": "reviewer", "round": current_round}
                        for item in values
                        if isinstance(item, dict)
                    ],
                ]
            content = str(self.content or "")[:MAX_TEXT_BYTES]
            node_path = state.node_paths.get(node_id, node_id)
            row = create_pending_approval(
                state.run_id,
                node_id,
                node_path,
                content,
                values,
                str(self.confirm_label or "确认继续")[:120],
                bool(getattr(self, "allow_return", False)),
            )
            if row["status"] != "pending":
                # A second invocation must not turn an old decision into implicit
                # approval. A resumed run keeps this same component suspended.
                raise RuntimeError(f"人工确认断点 {node_path} 已有终态 {row['status']}，拒绝自动重放")
            state.set_waiting(node_id)
            event = approval_event(state.run_id, node_id)
            try:
                while True:
                    if state.cancel_event.is_set():
                        state.finish_node(node_id, "cancelled", message="Run cancelled by user")
                        raise RunCancelled("Run cancelled by user")
                    current = _approval_row(state.run_id, node_id)
                    if current is None:
                        state.finish_node(node_id, "failed", message="Approval record disappeared")
                        raise RuntimeError("人工确认记录不存在")
                    status = current["status"]
                    if status in {"approved", "returned"}:
                        state.resume_running(node_id)
                        output = passthrough_payload(values, "human_approved")
                        revision_present = "revision_text" in current.keys() and current["revision_text"] is not None
                        revision_text = str(current["revision_text"]) if revision_present else ""
                        approval_info = {
                            "decision": current["decision"] or "approve",
                            "note": current["note"],
                            "decided_at": current["decided_at"],
                            "revised": revision_present,
                        }
                        if revision_present:
                            revised_value = human_revision_payload(values, revision_text)
                            source_refs = [
                                {
                                    key: item.get(key)
                                    for key in ("file_id", "display_name", "name", "relative_path", "sha256", "mime", "size", "status")
                                    if item.get(key) not in (None, "")
                                }
                                for item in revised_value.get("attachments", [])
                                if isinstance(item, dict)
                            ]
                            artifact_refs = [
                                {
                                    key: item.get(key)
                                    for key in ("name", "mime", "size", "sha256", "status")
                                    if item.get(key) not in (None, "")
                                }
                                for item in revised_value.get("artifacts", [])
                                if isinstance(item, dict)
                            ]
                            output = {
                                "status": "human_approved",
                                "active": True,
                                "text": revision_text,
                                "content": revision_text,
                                "items": [json.loads(json_text(revised_value))],
                                "upstream": [json.loads(json_text(revised_value))],
                                "original_result": {
                                    "status": "historical_original",
                                    "approval": approval_info,
                                    "source_refs": source_refs,
                                    "artifact_refs": artifact_refs,
                                },
                                "revision": {
                                    "text": revision_text,
                                    "note": current["revision_note"],
                                    "at": current["revision_at"],
                                },
                            }
                        output["approval"] = approval_info
                        output["status"] = "returned" if status == "returned" else "approved"
                        self.last_result = Data(data=output)
                        state.finish_node(node_id, "succeeded", output=output)
                        return self.last_result
                    if status == "rejected":
                        state.finish_node(node_id, "rejected", message=current["note"] or "Rejected by user")
                        raise RunRejected(current["note"] or "人工确认已拒绝")
                    if status in {"cancelled", "interrupted"}:
                        state.finish_node(node_id, status, message="Approval no longer active")
                        raise RunCancelled("人工确认已取消或因服务重启中断")
                    # Waiting in a worker thread keeps the event loop available to
                    # other LFX work and to independent kxy runs.
                    await asyncio.to_thread(event.wait, HUMAN_POLL_SECONDS)
                    event.clear()
            finally:
                with RUN_LOCK:
                    RUN_APPROVAL_EVENTS.pop((state.run_id, node_id), None)


    class KxyAnalyzerComponent(Component):
        display_name = "KXY Analyzer"
        name = "KxyAnalyzer"
        inputs = [
            DataInput(name="items", display_name="Inputs", is_list=True, required=False),
            MultilineInput(name="prompt", display_name="Prompt", value=""),
        ]
        outputs = [Output(name="result", display_name="Analysis", method="build_result")]

        async def build_result(self) -> Data:
            state: RunState = self.run_state
            restored, payload = state.restore_node(self.get_id())
            if restored:
                self.last_result = Data(data=payload)
                return self.last_result
            state.start_node(self.get_id())
            values = active_payloads(self.items)
            if not values:
                output = {"status": "skipped", "reason": "No active upstream result", "text": "", "output_format": normalize_output_format(getattr(self, "node_config", {}) or {})}
                self.last_result = Data(data=output)
                state.finish_node(self.get_id(), "skipped", output=output)
                return self.last_result
            if (
                isinstance(state.loop_context, dict)
                and state.loop_context.get("reviewer_component_id") == str(self.get_id())
                and state.loop_input_payload is not None
            ):
                # The reviewer receives the selected upstream result plus a
                # separate visible round-context payload.  This keeps the
                # original goal and feedback available even when the graph
                # only connects executor -> reviewer.
                values = [
                    *values,
                    {"status": "loop_round_context", "active": True, "round_context": state.loop_input_payload},
                ]
            self.last_input_payloads = values
            if state.cancel_event.is_set():
                output = {"status": "cancelled", "text": "", "output_format": normalize_output_format(getattr(self, "node_config", {}) or {})}
                self.last_result = Data(data=output)
                state.finish_node(self.get_id(), "cancelled", output=output)
                return self.last_result
            loop_started = None
            try:
                loop_started = state.begin_loop_call()
                result = self.runner(str(self.prompt or ""), values, self)
                output = {"status": "succeeded", **result}
                self.last_result = Data(data=output)
                state.finish_node(self.get_id(), "succeeded", output=output)
                return self.last_result
            except Exception as exc:
                state.finish_node(self.get_id(), "failed", message=str(exc))
                raise
            finally:
                state.finish_loop_call(loop_started)


    class KxyLoopComponent(Component):
        display_name = "Loop"
        name = "KxyLoop"
        inputs = [DataInput(name="items", display_name="Inputs", is_list=True, required=False)]
        outputs = [Output(name="result", display_name="Loop result", method="build_result")]

        async def build_result(self) -> Data:
            state: RunState = self.run_state
            node_id = str(self.get_id())
            restored, payload = state.restore_node(node_id)
            if restored:
                self.last_result = Data(data=payload)
                return self.last_result
            state.start_node(node_id)
            values = active_payloads(self.items)
            if not values:
                output = {"status": "skipped", "active": False, "reason": "No active upstream result", "text": ""}
                self.last_result = Data(data=output)
                state.finish_node(node_id, "skipped", output=output)
                return self.last_result

            config = dict(getattr(self, "node_config", {}) or {})
            nested = config.get("workflow")
            if not isinstance(nested, dict):
                raise ValueError(f"Loop {getattr(self, 'original_path', node_id)} is missing its inner workflow")
            settings = loop_settings(config.get("loop"))
            loop_path = str(getattr(self, "original_path", node_id))
            executor_id = settings["executor_id"]
            reviewer_id = settings["reviewer_id"]
            inner_compiled = flatten_workflow(normalize_workflow(nested))
            executor_flat = next(
                (item["flat_id"] for item in inner_compiled["records"] if item["path"] == executor_id),
                None,
            )
            reviewer_flat = next(
                (item["flat_id"] for item in inner_compiled["records"] if item["path"] == reviewer_id),
                None,
            )
            gate_id = loop_review_gate_node(nested, reviewer_id)
            gate_flat = next(
                (item["flat_id"] for item in inner_compiled["records"] if item["path"] == gate_id),
                None,
            ) if gate_id else None
            if not executor_flat or not reviewer_flat:
                raise ValueError(f"Loop {loop_path} could not resolve executor/reviewer checkpoints")
            inner_flats = [str(item["flat_id"]) for item in inner_compiled["records"]]
            original_input = values
            goal = str(settings["goal"] or "").strip() or loop_text(values, 1200)
            input_field = str(settings["input_field"] or "text")
            selected_input: Any = values
            if len(values) == 1:
                exists, candidate = nested_value(values[0], input_field)
                if exists:
                    selected_input = candidate
            try:
                max_rounds = max(1, min(LOOP_MAX_ROUNDS, int(settings["max_rounds"])))
                budget = max(1.0, min(LOOP_MAX_ACTIVE_BUDGET_SECONDS, float(settings["active_budget_seconds"])))
            except (TypeError, ValueError) as exc:
                raise ValueError(f"Loop {loop_path} has invalid bounds") from exc
            summary_limit = max(0, min(LOOP_MAX_PREVIOUS_SUMMARY_CHARS, int(settings["previous_summary_chars"])))
            feedback = settings.get("feedback") if isinstance(settings.get("feedback"), dict) else {}
            prior_rows = loop_round_rows(state.run_id, loop_path)
            # The loop budget belongs to the run/path, not to one retry
            # attempt.  Reused rows carry zero active time, while immutable
            # original rows retain the elapsed total for fail-closed resume.
            loop_active_seconds = sum(float(row["active_seconds"] or 0) for row in prior_rows)
            # A loop owns its budget independently of the root RunState.  This
            # matters when two sibling loops appear in one ordinary DAG: a
            # human pause in one must not consume or rewrite the other's clock.
            loop_control_state = RunState(
                state.run_id,
                state.workspace,
                state.cancel_event,
                state.config,
                attempt_no=state.attempt_no,
                loop_active_seconds=loop_active_seconds,
            )
            loop_control_state.loop_path = loop_path
            persisted_config = next(
                (
                    row for row in reversed(prior_rows)
                    if row["effective_max_rounds"] is not None and row["budget_seconds"] is not None
                ),
                None,
            )
            if persisted_config is not None:
                max_rounds = max(max_rounds, min(LOOP_MAX_ROUNDS, int(persisted_config["effective_max_rounds"] or max_rounds)))
                budget = max(budget, min(LOOP_MAX_ACTIVE_BUDGET_SECONDS, float(persisted_config["budget_seconds"] or budget)))
            completed_by_round: dict[int, sqlite3.Row] = {}
            for row in prior_rows:
                if row["status"] not in {"reviewed", "reused", "succeeded"}:
                    continue
                if row["review_passed"] is None or not row["executor_output"] or not row["reviewer_output"]:
                    continue
                round_no = int(row["round_no"])
                existing = completed_by_round.get(round_no)
                if existing is None or int(row["attempt_no"]) >= int(existing["attempt_no"]):
                    completed_by_round[round_no] = row

            latest_executor: dict[str, Any] | None = None
            latest_review: dict[str, Any] | None = None
            previous_summary = ""
            original_round_input = {
                "goal": goal,
                "input_field": input_field,
                "selected_input": selected_input,
                "original_input": original_input,
            }

            def apply_loop_control(control: dict[str, Any], control_round: int) -> None:
                nonlocal max_rounds, budget
                try:
                    budget = min(
                        LOOP_MAX_ACTIVE_BUDGET_SECONDS,
                        budget + max(0.0, float(control.get("additional_seconds") or 0)),
                    )
                    max_rounds = min(
                        LOOP_MAX_ROUNDS,
                        max_rounds + max(0, int(control.get("additional_rounds") or 0)),
                    )
                except (TypeError, ValueError) as exc:
                    raise ValueError(f"Loop {loop_path} continuation bounds are invalid") from exc
                update_loop_round(
                    state.run_id,
                    state.attempt_no,
                    loop_path,
                    control_round,
                    effective_max_rounds=max_rounds,
                    budget_seconds=budget,
                )
                loop_control_state.loop_active_seconds = loop_active_seconds
                loop_control_state.loop_budget_seconds = budget

            def final_output(executor_output: dict[str, Any], review: dict[str, Any], round_no: int) -> dict[str, Any]:
                result = json.loads(json_text(executor_output))
                human_gate = review.get("human_gate") if isinstance(review.get("human_gate"), dict) else None
                revision_present = bool(human_gate and "revision_text" in human_gate and human_gate.get("revision_text") is not None)
                revised_text = str((human_gate or {}).get("revision_text")) if revision_present else ""
                if revision_present:
                    result = human_revision_payload(result, revised_text)
                    result["original_executor_ref"] = {
                        "status": "historical_original",
                        "round": round_no,
                        "reviewer_node_id": reviewer_id,
                        "source": "run_node_checkpoint_and_approval_record",
                    }
                    result["human_review"] = human_gate
                result["loop"] = {
                    "path": loop_path,
                    "round": round_no,
                    "max_rounds": max_rounds,
                    "active_budget_seconds": budget,
                    "review": review,
                    "review_provenance": {
                        "reviewer_node_id": reviewer_id,
                        "reviewer_path": f"{loop_path}/round-{round_no}/{reviewer_id}",
                        "source": "selected_reviewer_output",
                    },
                }
                return result

            try:
                round_no = 1
                while True:
                    if state.cancel_event.is_set():
                        raise RunCancelled("Run cancelled by user")
                    if round_no > max_rounds:
                        ensure_loop_round(
                            state.run_id,
                            state.attempt_no,
                            loop_path,
                            round_no,
                            goal=goal,
                            original_input=original_input,
                            previous_summary=previous_summary,
                            executor_node_id=executor_id,
                            reviewer_node_id=reviewer_id,
                            max_rounds=max_rounds,
                            budget_seconds=budget,
                        )
                        loop_control_state.loop_active_seconds = loop_active_seconds
                        control = await wait_for_loop_control(loop_control_state, loop_path, round_no, "round_limit")
                        apply_loop_control(control, round_no)
                        continue
                    prior = completed_by_round.get(round_no)
                    if prior is not None:
                        executor_output = loop_round_json(prior, "executor_output", None)
                        review = loop_round_json(prior, "reviewer_output", None)
                        if not isinstance(executor_output, dict) or not isinstance(review, dict):
                            raise ValueError(f"Loop {loop_path} round {round_no} completed checkpoint is invalid")
                        validate_checkpoint_output(
                            (RUNS_ROOT / str(state.run_id) / "workspace").resolve(),
                            executor_output,
                            f"{loop_path}/round-{round_no}/{executor_id}",
                        )
                        ensure_loop_round(
                            state.run_id,
                            state.attempt_no,
                            loop_path,
                            round_no,
                            goal=goal,
                            original_input=original_input,
                            previous_summary=previous_summary,
                            executor_node_id=executor_id,
                            reviewer_node_id=reviewer_id,
                            max_rounds=max_rounds,
                            budget_seconds=budget,
                        )
                        if int(prior["attempt_no"]) != state.attempt_no:
                            update_loop_round(
                                state.run_id,
                                state.attempt_no,
                                loop_path,
                                round_no,
                                status="reused",
                                executor_input=loop_round_json(prior, "executor_input", {}),
                                executor_output=executor_output,
                                reviewer_input=loop_round_json(prior, "reviewer_input", {}),
                                reviewer_output=review,
                                review_passed=bool(prior["review_passed"]),
                                review_issues=loop_round_json(prior, "review_issues", []),
                                next_action=prior["next_action"],
                                node_outputs=loop_round_json(prior, "node_outputs", {}),
                                node_statuses=loop_round_json(prior, "node_statuses", {}),
                                active_seconds=0,
                                finished_at=utc_now(),
                            )
                            append_event(
                                state.run_id,
                                "loop_round_reused",
                                {"loop_path": loop_path, "round": round_no, "source_attempt": prior["attempt_no"]},
                            )
                        latest_executor = executor_output
                        latest_review = review
                        next_input_revision = review.get("next_input_revision") if isinstance(review.get("next_input_revision"), dict) else None
                        if next_input_revision is not None and "text" in next_input_revision:
                            latest_executor = human_revision_payload(executor_output, str(next_input_revision.get("text") or ""))
                            latest_executor["human_revision"] = {
                                "present": True,
                                "text": str(next_input_revision.get("text") or ""),
                                "source": "loop_review_gate",
                            }
                        previous_summary = loop_text(latest_executor, summary_limit)
                        if feedback.get("issues", True):
                            previous_summary = (previous_summary + "\nissues: " + "; ".join(str(item) for item in review.get("issues", []))).strip()
                        if feedback.get("next_action", True):
                            previous_summary = (previous_summary + "\nnext_action: " + str(review.get("next_action") or "")).strip()
                        previous_summary = previous_summary[:summary_limit]
                        if bool(review.get("passed")):
                            output = final_output(executor_output, review, round_no)
                            self.last_result = Data(data=output)
                            state.finish_node(node_id, "succeeded", output=output)
                            return self.last_result
                        if round_no >= max_rounds:
                            round_no += 1
                            ensure_loop_round(
                                state.run_id,
                                state.attempt_no,
                                loop_path,
                                round_no,
                                goal=goal,
                                original_input=original_input,
                                previous_summary=previous_summary,
                                executor_node_id=executor_id,
                                reviewer_node_id=reviewer_id,
                                max_rounds=max_rounds,
                                budget_seconds=budget,
                            )
                            loop_control_state.loop_active_seconds = loop_active_seconds
                            control = await wait_for_loop_control(loop_control_state, loop_path, round_no, "round_limit")
                            apply_loop_control(control, round_no)
                        else:
                            round_no += 1
                        continue

                    ensure_loop_round(
                        state.run_id,
                        state.attempt_no,
                        loop_path,
                        round_no,
                        goal=goal,
                        original_input=original_input,
                        previous_summary=previous_summary,
                        executor_node_id=executor_id,
                        reviewer_node_id=reviewer_id,
                        max_rounds=max_rounds,
                        budget_seconds=budget,
                    )
                    if loop_active_seconds >= budget:
                        loop_control_state.loop_active_seconds = loop_active_seconds
                        control = await wait_for_loop_control(loop_control_state, loop_path, round_no, "active_budget")
                        apply_loop_control(control, round_no)
                        continue

                    context: dict[str, Any] = {
                        "status": "loop_round_input",
                        "loop_path": loop_path,
                        "round": round_no,
                        "original_goal": goal,
                        "original_input": original_input,
                        "input_field": input_field,
                        "selected_input": selected_input,
                    }
                    if feedback.get("latest_result", True):
                        context["latest_result"] = latest_executor
                    if feedback.get("issues", True):
                        context["issues"] = (latest_review or {}).get("issues", [])
                    if feedback.get("next_action", True):
                        context["next_action"] = (latest_review or {}).get("next_action", "")
                    if feedback.get("history_summary", True):
                        context["history_summary"] = previous_summary[:summary_limit]
                    reviewer_input = {"round_context": context, "executor_result": latest_executor}
                    scope = loop_scope_id(loop_path, round_no)
                    scoped_ids = [scoped_component_id(scope, item) for item in inner_flats]
                    inner_state = RunState(
                        state.run_id,
                        state.workspace,
                        state.cancel_event,
                        state.config,
                        attempt_no=state.attempt_no,
                        resume_outputs=loop_resume_outputs(state.run_id, scope, scoped_ids),
                    )
                    round_active_start = loop_active_seconds
                    inner_state.loop_scope = scope
                    inner_state.loop_path = loop_path
                    inner_state.loop_round_no = round_no
                    inner_state.loop_budget_seconds = budget
                    inner_state.loop_context = {
                        "loop_path": loop_path,
                        "round": round_no,
                        "executor_id": executor_id,
                        "reviewer_id": reviewer_id,
                        "executor_component_id": scoped_component_id(scope, executor_flat),
                        "reviewer_component_id": scoped_component_id(scope, reviewer_flat),
                    }
                    inner_state.loop_input_payload = context
                    inner_state.loop_active_seconds = loop_active_seconds
                    try:
                        graph, inner_components, _terminals = build_lfx_graph(
                            normalize_workflow(nested),
                            inner_state,
                            component_prefix=scope,
                            display_path_prefix=f"{loop_path}/round-{round_no}",
                            loop_context=inner_state.loop_context,
                        )

                        async for result in graph.async_start(inputs=[{}], open_flow_span=False):
                            if hasattr(result, "vertex"):
                                append_event(
                                    state.run_id,
                                    "lfx_vertex",
                                    {"node_id": result.vertex.id, "loop_path": loop_path, "round": round_no, "valid": getattr(result, "valid", True)},
                                )
                        loop_active_seconds = inner_state.loop_active_seconds
                        state.node_paths.update(inner_state.node_paths)
                        state.blackbox_paths.update(inner_state.blackbox_paths)
                        executor_component = inner_components.get(scoped_component_id(scope, executor_flat))
                        reviewer_component = inner_components.get(scoped_component_id(scope, reviewer_flat))
                        gate_component = inner_components.get(scoped_component_id(scope, gate_flat)) if gate_flat else None
                        executor_output = payload_from_data(getattr(executor_component, "last_result", None))
                        reviewer_payload = payload_from_data(getattr(reviewer_component, "last_result", None))
                        gate_payload = payload_from_data(getattr(gate_component, "last_result", None)) if gate_component else None
                        if not executor_component or not getattr(executor_component, "last_result", None):
                            raise ValueError(f"Loop {loop_path} round {round_no} executor produced no result")
                        if not reviewer_component or not getattr(reviewer_component, "last_result", None):
                            raise ValueError(f"Loop {loop_path} round {round_no} reviewer produced no result")
                        review = parse_loop_review(reviewer_payload, settings, loop_path)
                        next_executor_input = executor_output
                        if gate_payload and isinstance(gate_payload.get("approval"), dict):
                            approval = gate_payload["approval"]
                            revision = gate_payload.get("revision") if isinstance(gate_payload.get("revision"), dict) else None
                            human_gate = {
                                "node_id": gate_id,
                                "decision": approval.get("decision"),
                                "note": approval.get("note"),
                                "revised": bool(approval.get("revised")),
                                "revision_text": revision.get("text") if revision is not None and "text" in revision else None,
                            }
                            review["human_gate"] = human_gate
                            if approval.get("decision") == "return":
                                review["passed"] = False
                                revision_present = revision is not None and "text" in revision
                                if revision_present:
                                    revision_text = str(revision.get("text") or "")
                                    review["next_input_revision"] = {
                                        "present": True,
                                        "text": revision_text,
                                        "source": "loop_review_gate",
                                    }
                                    next_executor_input = human_revision_payload(executor_output, revision_text)
                                    next_executor_input["human_revision"] = {
                                        "present": True,
                                        "text": revision_text,
                                        "source": "loop_review_gate",
                                    }
                                feedback_note = str(approval.get("note") or ("人工闸门已提交修订正文" if revision_present else "人工闸门退回修改"))
                                review["issues"] = [*review.get("issues", []), feedback_note][:100]
                                review["next_action"] = feedback_note
                        node_outputs, node_statuses = loop_node_checkpoint_map(state.run_id, scoped_ids)
                        round_active_seconds = loop_round_active_seconds(
                            state.run_id,
                            state.attempt_no,
                            loop_path,
                            round_no,
                        ) + max(0.0, loop_active_seconds - round_active_start)
                        update_loop_round(
                            state.run_id,
                            state.attempt_no,
                            loop_path,
                            round_no,
                            status="succeeded" if review["passed"] else "reviewed",
                            executor_input=context,
                            executor_output=executor_output,
                            reviewer_input=getattr(
                                reviewer_component,
                                "last_input_payloads",
                                active_payloads(getattr(reviewer_component, "items", None)),
                            ),
                            reviewer_output=review,
                            review_passed=review["passed"],
                            review_issues=review["issues"],
                            next_action=review["next_action"],
                            active_seconds=round_active_seconds,
                            node_outputs=node_outputs,
                            node_statuses=node_statuses,
                            error=None,
                            finished_at=utc_now(),
                        )
                        latest_executor = next_executor_input
                        latest_review = review
                        previous_summary = loop_text(latest_executor, summary_limit)
                        if feedback.get("issues", True):
                            previous_summary = (previous_summary + "\nissues: " + "; ".join(review["issues"])).strip()
                        if feedback.get("next_action", True):
                            previous_summary = (previous_summary + "\nnext_action: " + review["next_action"]).strip()
                        previous_summary = previous_summary[:summary_limit]
                        if review["passed"]:
                            output = final_output(executor_output, review, round_no)
                            self.last_result = Data(data=output)
                            state.finish_node(node_id, "succeeded", output=output)
                            return self.last_result
                    except Exception as exc:
                        loop_active_seconds = inner_state.loop_active_seconds
                        state.node_paths.update(inner_state.node_paths)
                        state.blackbox_paths.update(inner_state.blackbox_paths)
                        node_outputs, node_statuses = loop_node_checkpoint_map(state.run_id, scoped_ids)
                        round_active_seconds = loop_round_active_seconds(
                            state.run_id,
                            state.attempt_no,
                            loop_path,
                            round_no,
                        ) + max(0.0, loop_active_seconds - round_active_start)
                        if isinstance(exc, LoopBudgetExceeded) or "循环活动执行预算已用尽" in str(exc):
                            update_loop_round(
                                state.run_id,
                                state.attempt_no,
                                loop_path,
                                round_no,
                                status="budget_waiting",
                                active_seconds=round_active_seconds,
                                node_outputs=node_outputs,
                                node_statuses=node_statuses,
                                error="active_budget",
                            )
                            loop_control_state.loop_active_seconds = loop_active_seconds
                            control = await wait_for_loop_control(loop_control_state, loop_path, round_no, "active_budget")
                            apply_loop_control(control, round_no)
                            continue
                        update_loop_round(
                            state.run_id,
                            state.attempt_no,
                            loop_path,
                            round_no,
                            status="failed",
                            active_seconds=round_active_seconds,
                            node_outputs=node_outputs,
                            node_statuses=node_statuses,
                            error=str(exc),
                            finished_at=utc_now(),
                        )
                        raise
                    if round_no >= max_rounds:
                        round_no += 1
                        ensure_loop_round(
                            state.run_id,
                            state.attempt_no,
                            loop_path,
                            round_no,
                            goal=goal,
                            original_input=original_input,
                            previous_summary=previous_summary,
                            executor_node_id=executor_id,
                            reviewer_node_id=reviewer_id,
                            max_rounds=max_rounds,
                            budget_seconds=budget,
                        )
                        loop_control_state.loop_active_seconds = loop_active_seconds
                        control = await wait_for_loop_control(loop_control_state, loop_path, round_no, "round_limit")
                        apply_loop_control(control, round_no)
                    else:
                        round_no += 1
                raise ValueError(f"Loop {loop_path} ended without an accepted reviewer result")
            finally:
                state.loop_scope = None
                state.loop_path = None
                state.loop_round_no = None
                state.loop_context = None
                state.loop_input_payload = None
                state.loop_budget_seconds = None


    class KxyConditionComponent(Component):
        display_name = "KXY Condition"
        name = "KxyCondition"
        inputs = [
            DataInput(name="items", display_name="Inputs", is_list=True, required=False),
            StrInput(name="field", display_name="Field", value="status"),
            StrInput(name="operator", display_name="Operator", value="equals"),
            StrInput(name="expected", display_name="Expected", value="succeeded"),
            StrInput(name="missing_strategy", display_name="Missing field strategy", value="error"),
        ]
        outputs = [
            Output(name="true", display_name="True", method="true_branch", group_outputs=True),
            Output(name="false", display_name="False", method="false_branch", group_outputs=True),
        ]

        def _decision_value(self) -> bool:
            if not hasattr(self, "_decision"):
                self._decision = evaluate_condition(self.items, self.field, self.operator, self.expected, self.missing_strategy)
                self.last_decision = self._decision
            return bool(self._decision)

        def _route(self, route: str) -> Data:
            state: RunState = self.run_state
            restored, stored = state.restore_node(self.get_id())
            if restored and isinstance(stored, dict):
                branch = str(stored.get("branch") or "")
                if branch in CONDITION_HANDLES:
                    self._decision = branch == "true"
            state.start_node(self.get_id())
            try:
                decision = self._decision_value()
                selected = "true" if decision else "false"
                stopped = "false" if selected == "true" else "true"
                self.stop(stopped)
                self.graph.exclude_branch_conditionally(self._id, output_name=stopped)
                if route != selected:
                    return Data(data={"status": "skipped", "branch": route, "active": False})
                text = "\n\n".join(str(item.get("text", "")) for item in active_payloads(self.items) if item.get("text"))
                output = {
                    "status": "condition",
                    "active": True,
                    "branch": selected,
                    "field": self.field,
                    "operator": self.operator,
                    "expected": self.expected,
                    "missing_strategy": self.missing_strategy,
                    "upstream": active_payloads(self.items),
                    "text": text,
                }
                state.finish_node(
                    self.get_id(),
                    "succeeded",
                    output={**output, "decision": decision},
                )
                return Data(data=output)
            except Exception as exc:
                state.finish_node(self.get_id(), "failed", message=str(exc))
                raise

        def true_branch(self) -> Data:
            return self._route("true")

        def false_branch(self) -> Data:
            return self._route("false")


    class KxyContainerComponent(Component):
        display_name = "KXY Output"
        name = "KxyContainer"
        inputs = [DataInput(name="items", display_name="Inputs", is_list=True, required=False)]
        outputs = [Output(name="result", display_name="Collected Output", method="build_result")]

        def build_result(self) -> Data:
            state: RunState = self.run_state
            restored, payload = state.restore_node(self.get_id())
            if restored:
                self.last_result = Data(data=payload)
                return self.last_result
            state.start_node(self.get_id())
            values = active_payloads(self.items)
            text = "\n\n".join(str(item.get("text", "")) for item in values if item.get("text"))
            config = dict(getattr(self, "node_config", {}) or {})
            export_formats = normalize_container_formats(config.get("export_formats"))
            allowed_file_extensions = normalize_allowed_file_extensions(config.get("allowed_file_extensions"))
            json_mode = config.get("json_mode", "full")
            def content_value(item: dict[str, Any]) -> Any:
                # Branch, human, and blackbox components are transport
                # wrappers. Walk their owned children before falling back to
                # the wrapper's presentation text so JSON objects do not turn
                # into JSON strings at the final container.
                for key in ("upstream", "items", "inputs"):
                    child = item.get(key)
                    if isinstance(child, list) and child:
                        nested = [content_value(value) if isinstance(value, dict) else value for value in child]
                        return nested[0] if len(nested) == 1 else nested
                    if isinstance(child, dict):
                        return content_value(child)
                if item.get("content") is not None:
                    return item["content"]
                if item.get("structured") is not None:
                    return item["structured"]
                if "text" in item:
                    return item.get("text", "")
                return item

            content_values = [content_value(item) for item in values]
            content: Any = content_values[0] if len(content_values) == 1 else content_values
            all_artifacts: list[dict[str, Any]] = []
            seen_artifacts: set[str] = set()

            def collect_artifacts(item: Any) -> None:
                if not isinstance(item, dict):
                    return
                for artifact in item.get("artifacts", []):
                    if not isinstance(artifact, dict) or not artifact.get("path"):
                        continue
                    path_value = str(artifact["path"])
                    if path_value in seen_artifacts:
                        continue
                    seen_artifacts.add(path_value)
                    all_artifacts.append(dict(artifact))
                for child in payload_children(item):
                    collect_artifacts(child)

            for value in values:
                collect_artifacts(value)
            accepted_artifacts: list[dict[str, Any]] = []
            excluded_artifacts: list[dict[str, Any]] = []
            for artifact in all_artifacts:
                extension = Path(str(artifact.get("name") or artifact.get("path") or "")).suffix.lower()
                if allowed_file_extensions is None or extension in allowed_file_extensions:
                    accepted_artifacts.append(artifact)
                else:
                    excluded_artifacts.append(artifact)
            output = {
                "status": "succeeded",
                "items": values,
                "content": content,
                "text": text,
                "artifact_count": len(accepted_artifacts),
                "input_count": len(values),
                "artifacts": accepted_artifacts,
                "excluded_artifacts": excluded_artifacts,
                "export_formats": export_formats,
                "allowed_file_extensions": allowed_file_extensions,
                "json_mode": json_mode,
            }
            self.last_result = Data(data=output)
            state.finish_node(self.get_id(), "succeeded", output=output)
            return self.last_result


def execute_cli(state, prompt, payloads, component):
    config = dict(getattr(component, 'node_config', {}))
    binding = config.get("resolved_model") if isinstance(config.get("resolved_model"), dict) else None
    agent_id = str((binding or {}).get("agent_id") or config.get("agent_id") or config.get("cli", "codex")).strip().lower()
    cli = agent_id
    node_id = str(component.get_id())
    validate_analyzer_output_config(config, node_id)
    output_format = normalize_output_format(config)
    output_schema = config.get("output_schema")
    node_workspace = state.active_node_workspace_root / node_id
    node_workspace.mkdir(parents=True, exist_ok=False)
    input_dir = node_workspace / 'inputs'; input_dir.mkdir()
    output_dir = node_workspace / 'outputs'; output_dir.mkdir()
    attachments = []
    copied_inputs: dict[str, Path] = {}

    def copy_inputs(payload):
        value = json.loads(json_text(payload))

        def copy_file(relative: Any) -> str:
            if not isinstance(relative, str) or not relative:
                raise ValueError('上游附件路径无效')
            candidate = (state.workspace / relative).resolve()
            if not resolved_inside(candidate, state.workspace) or not candidate.is_file():
                raise ValueError('上游附件路径无效')
            cache_key = str(candidate)
            target = copied_inputs.get(cache_key)
            if target is None:
                source_digest = hashlib.sha256(candidate.read_bytes()).hexdigest()
                target = input_dir / safe_storage_name(candidate.name, source_digest)
                if not target.exists():
                    shutil.copyfile(candidate, target)
                copied_inputs[cache_key] = target
            if str(target) not in {str(item) for item in attachments}:
                attachments.append(target)
            return str(target.relative_to(node_workspace))

        def walk(item: Any) -> Any:
            if isinstance(item, list):
                return [walk(child) for child in item]
            if not isinstance(item, dict):
                return item
            result = {str(key): walk(child) for key, child in item.items()}
            if result.get('file_path') is not None:
                result['file_path'] = copy_file(result['file_path'])
            artifacts = result.get('artifacts')
            if isinstance(artifacts, list):
                for artifact in artifacts:
                    if isinstance(artifact, dict) and artifact.get('path') is not None:
                        artifact['path'] = copy_file(artifact['path'])
            return result

        return walk(value)
    node_payloads = [copy_inputs(payload) for payload in payloads]
    (node_workspace / 'input-context.json').write_text(json_text({'inputs': node_payloads}), encoding='utf-8')
    selected = config.get('skill_ids') or []
    if isinstance(selected, str): selected = [selected]
    selected_mcps = config.get("mcp_ids") or []
    if isinstance(selected_mcps, str):
        selected_mcps = [selected_mcps]
    if selected and check_skill_dependencies is not None:
        dependency_binding = binding or {
            "agent_id": agent_id,
            "model_ref": config.get("model_ref"),
            "credential_id": config.get("credential_id"),
            "model": config.get("model"),
            "source": config.get("source"),
            "runtime_interpreter": config.get("runtime_interpreter"),
        }
        try:
            dependency_result = check_skill_dependencies(
                agent_id,
                selected,
                binding=dependency_binding,
                mcp_ids=selected_mcps if isinstance(selected_mcps, list) else [],
            )
        except ValueError as exc:
            # A run owns an immutable Skill snapshot.  If the source DB row
            # was removed after submission, keep that snapshot runnable; a
            # fresh run still gets the normal dependency gate above.
            snapshot_ids = set()
            try:
                snapshot_value = json.loads((state.workspace / "manifest.json").read_text(encoding="utf-8"))
                snapshot_ids = {str(item.get("id")) for item in snapshot_value.get("skills", []) if isinstance(item, dict)}
            except (OSError, TypeError, ValueError, json.JSONDecodeError):
                snapshot_ids = set()
            if "Skill 不存在" not in str(exc) or not set(str(item) for item in selected).issubset(snapshot_ids):
                raise
            dependency_result = {"skills": []}
        blocked = [item for item in dependency_result.get("skills", []) if isinstance(item, dict) and item.get("blocking")]
        if blocked:
            details = []
            for item in blocked:
                missing = item.get("missing") if isinstance(item.get("missing"), list) else []
                names = [str(dep.get("name") or "依赖") for dep in missing if isinstance(dep, dict)]
                detail = ", ".join(names) or str(item.get("message") or "runtime 未通过")
                details.append(f"{item.get('skill_id') or 'Skill'}：{detail}")
            raise RuntimeError(f"Agent {agent_id} 的 Skill 依赖检查阻止运行；请先修复：{'；'.join(details)}")
    mounted = mount_skills(selected, node_workspace, target_agent=agent_id, run_workspace=state.workspace)
    run_snapshot = json.loads((state.workspace / 'manifest.json').read_text(encoding='utf-8'))
    if selected_mcps and mcp_runtime_environment is None:
        raise RuntimeError("MCP 配置模块不可用，无法安全运行选定 MCP")
    mcp_config = None
    mcp_env: dict[str, str] = {}
    if selected_mcps:
        if prepare_mcp_configuration is None or mcp_runtime_environment is None:
            raise RuntimeError("MCP 配置模块不可用，无法安全运行选定 MCP")
        mcp_config = prepare_mcp_configuration(
            agent_id,
            selected_mcps,
            node_workspace,
            snapshot={"mcps": run_snapshot.get("mcps", [])},
        )
        mcp_env = mcp_runtime_environment(
            agent_id,
            selected_mcps,
            snapshot={"mcps": run_snapshot.get("mcps", [])},
        )
        required_env = set(mcp_config.get("required_env", [])) if isinstance(mcp_config, dict) else set()
        for name in required_env:
            if name not in mcp_env and os.environ.get(name):
                mcp_env[name] = os.environ[name]
        missing_env = sorted(name for name in required_env if name not in mcp_env and not os.environ.get(name))
        if missing_env:
            raise RuntimeError("选定 MCP 缺少运行时环境变量引用：" + ", ".join(missing_env))
        _update_run_snapshot(
            state.workspace,
            lambda current, node=node_id, info=mcp_config: current.setdefault("mcp_mappings", {}).update(
                {
                    node: {
                        key: info.get(key)
                        for key in ("agent_id", "server_ids", "relative_path", "format", "required_env", "runnable", "upstream_commit", "compatibility")
                        if info.get(key) is not None
                    }
                }
            ),
        )
    # The analyzer textarea is the complete user prompt. Input manifests,
    # selected Skill snapshots and the outputs directory remain available in
    # the isolated workspace, but are never silently appended to this value.
    full_prompt = prompt
    network = bool(config.get("network", True))
    last_message = node_workspace / 'last-message.txt'
    secret = None
    profile_file = None
    if build_headless_command is not None and prepare_headless_environment is not None and wrap_headless_command is not None:
        credential_id = (binding or {}).get("credential_id") or config.get("credential_id")
        base_env = {
            key: value
            for key, value in os.environ.items()
            if key in {'PATH','HOME','USER','LOGNAME','SHELL','LANG','LC_ALL','TERM','SSL_CERT_FILE','SSL_CERT_DIR','HTTP_PROXY','HTTPS_PROXY','ALL_PROXY','NO_PROXY','http_proxy','https_proxy','all_proxy','no_proxy'}
        }
        env = prepare_headless_environment(
            agent_id,
            node_workspace,
            credential_id=str(credential_id) if credential_id else None,
            network=network,
            base_env=base_env,
            api_format=(binding or {}).get("credential_api_format"),
            model=(binding or {}).get("model") or config.get("model"),
        )
        credential_env_name = (binding or {}).get("credential_env_name")
        if credential_env_name:
            secret = env.get(str(credential_env_name))
        if mcp_config:
            env.update(mcp_env)
            _attach_task_mcp_environment(agent_id, env, node_workspace, mcp_config)
        argv = build_headless_command(
            agent_id,
            full_prompt,
            node_workspace,
            # KXY_*_BIN is an explicit test/fixture override. It is resolved by
            # the configuration helper without changing the persisted binding.
            executable=None if os.environ.get(f"KXY_{agent_id.upper()}_BIN") else (binding or {}).get("executable"),
            runtime_interpreter=(binding or {}).get("runtime_interpreter"),
            model=(binding or {}).get("model") or config.get("model"),
            effort=(binding or {}).get("selected_effort") or config.get("effort") or config.get("variant"),
            attachments=attachments,
            skills=[node_workspace / item["path"] for item in mounted],
            expect_json=output_format == "json",
            network=network,
            last_message=last_message,
            endpoint=(binding or {}).get("credential_endpoint") or None,
            credential_env_name=credential_env_name,
            credential_provider=(binding or {}).get("credential_provider"),
        )
        if mcp_config and agent_id == "claude":
            argv = [*argv[:-1], "--mcp-config", str(mcp_config["path"]), "--strict-mcp-config", argv[-1]]
        elif mcp_config and agent_id == "codex":
            argv = [*argv[:-1], *_codex_mcp_overrides(Path(str(mcp_config["path"]))), argv[-1]]
        wrapped, profile_file = wrap_headless_command(argv, agent_id, node_workspace, network=network)
    else:
        # Development fallback for a checkout before agent_config.py exists.
        env = {key: value for key, value in os.environ.items() if key in {'PATH','HOME','USER','LOGNAME','SHELL','LANG','LC_ALL','TERM','SSL_CERT_FILE','SSL_CERT_DIR','HTTP_PROXY','HTTPS_PROXY','ALL_PROXY','NO_PROXY','http_proxy','https_proxy','all_proxy','no_proxy'}}
        credential_id = config.get('credential_id')
        if credential_id:
            with connect_db() as db:
                credential = db.execute('SELECT * FROM credentials WHERE id=?', (credential_id,)).fetchone()
            if credential is None:
                raise RuntimeError('所选凭据已删除，请重新配置')
            secret = keychain_secret(credential['credential_ref'])
            env[credential['env_name']] = secret
            config['_env_name'] = credential['env_name']
            config['_endpoint'] = credential['endpoint']
            endpoint_env = credential_endpoint_env(credential['env_name'])
            if endpoint_env and credential['endpoint']: env[endpoint_env] = credential['endpoint']
        if cli == 'opencode':
            state_dir = node_workspace / '.cli-state'
            for kind in ('data','state','cache','config','tmp'):
                (state_dir / kind).mkdir(parents=True, exist_ok=True)
            native_data = Path(os.environ.get('XDG_DATA_HOME', Path.home() / '.local/share'))
            auth = native_data / 'opencode/auth.json'
            auth_link = state_dir / 'data/opencode/auth.json'
            auth_link.parent.mkdir(parents=True, exist_ok=True)
            if auth.is_file() and not credential_id: auth_link.symlink_to(auth.resolve())
            for name, kind in [('XDG_DATA_HOME','data'),('XDG_STATE_HOME','state'),('XDG_CACHE_HOME','cache'),('XDG_CONFIG_HOME','config')]: env[name] = str(state_dir / kind)
            env.update({'TMPDIR':str(state_dir/'tmp'), 'OPENCODE_DISABLE_CLAUDE_CODE':'true', 'OPENCODE_DISABLE_CLAUDE_CODE_SKILLS':'true', 'OPENCODE_DISABLE_DEFAULT_PLUGINS':'true', 'OPENCODE_DISABLE_PROJECT_CONFIG':'true', 'OPENCODE_DISABLE_AUTOUPDATE':'true', 'OPENCODE_DISABLE_LSP_DOWNLOAD':'true'})
            local_config = {'autoupdate':False, 'share':'disabled', 'snapshot':False, 'permission':{'*':'allow','external_directory':'deny'}, 'instructions':[]}
            if credential_id and config.get('_endpoint'):
                provider = 'openai' if config['_env_name'] == 'OPENAI_API_KEY' else 'anthropic'
                local_config['provider'] = {provider:{'options':{'baseURL':config['_endpoint']}}}
            env['OPENCODE_CONFIG_CONTENT'] = json_text(local_config)
        argv = build_cli_argv(config, full_prompt, node_workspace, last_message, attachments)
        wrapped, profile_file = sandbox_argv(argv, node_workspace, None, network, cli)
    append_event(state.run_id, 'cli_started', {'node_id':node_id,'node_path':getattr(component, 'original_path', node_id),'cli':cli,'model':config.get('model'),'effort':config.get('effort') or config.get('variant'), 'skills':mounted, 'mcp_ids': [str(item) for item in selected_mcps] if isinstance(selected_mcps, list) else []})
    timeout = max(5, min(int(config.get('timeout', 900)), 3600))
    loop_deadline = None
    if state.loop_budget_seconds is not None:
        remaining = state.loop_budget_seconds - state.loop_active_seconds
        if remaining <= 0:
            raise LoopBudgetExceeded("循环活动执行预算已用尽，请增加有界预算或停止循环")
        loop_deadline = time.monotonic() + remaining
    lines = []; process = None
    secret_values = {secret} if secret else set()
    secret_values.update(value for value in mcp_env.values() if value)
    secret_values.update(
        value
        for key, value in env.items()
        if isinstance(value, str)
        and value
        and str(key).upper().endswith(("API_KEY", "AUTH_TOKEN", "ACCESS_TOKEN", "_TOKEN"))
    )
    # Replace longer values first so a short token cannot leave a suffix of a
    # longer credential visible in a structured event or diagnostic line.
    redaction_values = tuple(sorted(secret_values, key=len, reverse=True))

    def safe_text(value):
        for value_to_redact in redaction_values:
            value = value.replace(value_to_redact, '[REDACTED]')
        return redact(value, max_length=2_000_000)
    try:
        process = subprocess.Popen(wrapped, cwd=node_workspace, env=env, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, bufsize=1, start_new_session=True)
        with RUN_LOCK: RUN_PROCESSES[state.run_id] = process
        def read_output():
            total = 0
            for raw in process.stdout:
                line = safe_text(raw.rstrip())
                # Preserve full structured final events separately from bounded UI logs.
                if total < 4_000_000:
                    lines.append(line); total += len(line)
                append_event(state.run_id, 'cli_log', {'node_id':node_id,'text':line[:MAX_EVENT_TEXT]}) if total < 200_000 else None
        reader = threading.Thread(target=read_output, daemon=True); reader.start()
        started = time.monotonic()
        while process.poll() is None:
            if state.cancel_event.is_set(): raise RuntimeError('Run cancelled by user')
            elapsed = time.monotonic() - started
            if elapsed > timeout: raise TimeoutError(f'CLI 超时：{timeout} 秒')
            if loop_deadline is not None and time.monotonic() >= loop_deadline:
                raise LoopBudgetExceeded("循环活动执行预算已用尽，请增加有界预算或停止循环")
            time.sleep(.1)
        reader.join(timeout=3)
        if state.cancel_event.is_set(): raise RuntimeError('Run cancelled by user')
        if process.returncode != 0: raise RuntimeError(f'{cli} 退出码 {process.returncode}: ' + (lines[-1][:2000] if lines else '没有诊断信息'))
        has_generated_files = any(
            path.is_file() or path.is_symlink()
            for path in output_dir.rglob('*')
        )
        if parse_final_output is not None:
            try:
                parsed = parse_final_output(agent_id, lines, last_message=last_message, expect_json=output_format == "json")
            except RuntimeError:
                if output_format == "auto" and has_generated_files:
                    parsed = {"text": "", "structured": None, "format": "files-only"}
                else:
                    raise
            if parsed.get("format") == "plain" and agent_id != "hermes":
                if output_format == "auto" and has_generated_files:
                    parsed = {"text": "", "structured": None, "format": "files-only"}
                else:
                    raise RuntimeError(f"{agent_id} 未产生可验证的最终事件；仅有启动或日志文本")
            final = safe_text(str(parsed.get('text') or ''))
            structured = (
                redact_object(parsed.get('structured'), redaction_values=redaction_values)
                if parsed.get('structured') is not None
                else None
            )
        else:
            final = safe_text(extract_final_output(lines, last_message))
            structured = None
            if output_format == "json":
                try: structured = json.loads(final)
                except ValueError as exc: raise RuntimeError('分析器要求 JSON，但最终回答不是合法 JSON') from exc
                if not isinstance(structured, dict): raise RuntimeError('分析器要求单个 JSON 对象')
                structured = redact_object(structured, redaction_values=redaction_values)
        if output_format == "auto" and structured is None and final.strip():
            try:
                candidate = json.loads(final)
            except (TypeError, ValueError, json.JSONDecodeError):
                candidate = None
            if isinstance(candidate, dict):
                # Auto mode keeps the returned text byte-for-byte while
                # exposing a structured view for condition routing.
                structured = redact_object(candidate, redaction_values=redaction_values)
        if output_format == "json":
            if not isinstance(structured, dict):
                raise RuntimeError("分析器要求单个 JSON 对象")
            try:
                validate_json_schema_value(structured, output_schema)
            except ValueError as exc:
                raise RuntimeError(str(exc)) from exc
            # Store a canonical object representation so the displayed result,
            # result.json, and parser all agree on the same valid JSON object.
            final = json.dumps(structured, ensure_ascii=False, indent=2)
        if not final.strip() and not has_generated_files:
            raise RuntimeError('CLI 未产生最终回答或生成文件；运行失败')
        if final.strip() and output_format != "auto":
            answer_suffix = {"markdown": ".md", "text": ".txt", "json": ".json"}[output_format]
            answer_path = output_dir / f"answer{answer_suffix}"
            if answer_path.exists(): answer_path = output_dir / (f"kxy-answer-{uuid.uuid4().hex[:8]}{answer_suffix}")
            answer_path.write_text(final, encoding='utf-8')
        artifacts = []
        total = 0
        for path in sorted(output_dir.rglob('*')):
            if path.is_symlink(): raise RuntimeError('分析器产物含符号链接，拒绝接收')
            if not path.is_file(): continue
            if path.stat().st_size > MAX_FILE_BYTES - total: raise RuntimeError('分析器产物超过50 MB')
            content = path.read_bytes(); total += len(content)
            if len(artifacts)>=100 or total>MAX_FILE_BYTES: raise RuntimeError('分析器产物超过 100 个文件或 50 MB')
            if any(value.encode("utf-8") in content for value in redaction_values):
                # Do not leave a rejected credential-bearing artifact in the
                # run tree; the run error and redacted logs remain available.
                path.unlink(missing_ok=True)
                raise RuntimeError('产物包含凭据内容，拒绝发布')
            artifacts.append({'name':str(path.relative_to(output_dir)), 'path':str(path.relative_to(state.workspace)), 'size':len(content),'sha256':hashlib.sha256(content).hexdigest()})
        return {'text':final,'content':structured if structured is not None else final,'structured':structured,'output_format':output_format,'output_schema':output_schema if output_format == 'json' else None,'cli':cli,'cli_version':(binding or {}).get('version') or probe_cli(cli).get('version'),'model':(binding or {}).get('model') or config.get('model') or 'CLI default','model_alias':(binding or {}).get('alias'),'model_source':(binding or {}).get('source'),'effort':(binding or {}).get('selected_effort') or config.get('effort') or config.get('variant'),'skills':mounted,'artifacts':artifacts,'attachments':[item.get('file_id') for item in node_payloads if item.get('file_id')], 'node_path':getattr(component, 'original_path', node_id)}
    finally:
        if process is not None:
            kill_process_group(process)
            if process.stdout is not None: process.stdout.close()
        secret = None
        with RUN_LOCK: RUN_PROCESSES.pop(state.run_id, None)
        if profile_file: profile_file.unlink(missing_ok=True)
        if last_message is not None:
            last_message.unlink(missing_ok=True)
        # Login links and CLI state are not research artifacts and are not retained.
        shutil.rmtree(node_workspace / '.cli-state', ignore_errors=True)


def kill_process_group(process):
    # Also stop descendants when the parent exited before them.
    try:
        os.killpg(process.pid, signal.SIGTERM)
        time.sleep(.1)
        os.killpg(process.pid, signal.SIGKILL)
    except (ProcessLookupError, OSError):
        pass
    try: process.wait(timeout=2)
    except (subprocess.TimeoutExpired, OSError): pass


def write_run_manifest(run_id: str, workspace: Path, payload: dict[str, Any]) -> Path:
    manifest = workspace / "output-manifest.json"
    temp = workspace / f".output-manifest.{uuid.uuid4().hex}.tmp"
    temp.write_text(json_text(redact_object(payload)), encoding="utf-8")
    temp.replace(manifest)
    return manifest


def write_atomic_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        raise FileExistsError(f"Refusing to overwrite existing output: {path}")
    temp = path.parent / f".{path.name}.{uuid.uuid4().hex}.tmp"
    temp.write_text(json_text(redact_object(payload)), encoding="utf-8")
    temp.replace(path)


CHECKPOINT_VERSION = "kxy.run.checkpoint.v8"


def _checkpoint_file(workspace: Path, raw_path: Any, label: str) -> Path:
    if not isinstance(raw_path, str) or not raw_path.strip():
        raise ValueError(f"{label} 缺少相对路径")
    relative = Path(raw_path)
    if relative.is_absolute() or ".." in relative.parts:
        raise ValueError(f"{label} 路径越过运行 workspace")
    lexical = workspace / relative
    candidate = lexical.resolve()
    if lexical.is_symlink() or not resolved_inside(candidate, workspace) or not candidate.is_file() or candidate.is_symlink():
        raise ValueError(f"{label} 路径无效")
    return candidate


def validate_checkpoint_output(workspace: Path, output: Any, node_id: str) -> None:
    """Validate only the file-bearing envelopes owned by kxy transport."""

    if not isinstance(output, dict):
        raise ValueError(f"节点 {node_id} 没有可恢复的结构化 checkpoint")

    def inspect(value: Any, path_label: str) -> None:
        if not isinstance(value, dict):
            return
        if value.get("file_path") is not None:
            source = _checkpoint_file(workspace, value.get("file_path"), f"节点 {node_id} 的输入副本")
            expected_hash = str(value.get("sha256") or "")
            if not re.fullmatch(r"[a-f0-9]{64}", expected_hash) or hashlib.sha256(source.read_bytes()).hexdigest() != expected_hash:
                raise ValueError(f"节点 {node_id} 的输入副本哈希不匹配")
        artifacts = value.get("artifacts")
        if artifacts is not None:
            if not isinstance(artifacts, list):
                raise ValueError(f"节点 {node_id} 的产物清单无效")
            for index, artifact in enumerate(artifacts):
                if not isinstance(artifact, dict):
                    raise ValueError(f"节点 {node_id} 的产物清单无效")
                artifact_path = _checkpoint_file(
                    workspace,
                    artifact.get("path"),
                    f"节点 {node_id} 的产物 {index + 1}",
                )
                expected_hash = str(artifact.get("sha256") or "")
                if not re.fullmatch(r"[a-f0-9]{64}", expected_hash):
                    raise ValueError(f"节点 {node_id} 的产物缺少 SHA-256")
                if hashlib.sha256(artifact_path.read_bytes()).hexdigest() != expected_hash:
                    raise ValueError(f"节点 {node_id} 的产物哈希不匹配")
                if artifact.get("size") is not None and (
                    isinstance(artifact.get("size"), bool)
                    or not isinstance(artifact.get("size"), int)
                    or artifact_path.stat().st_size != artifact["size"]
                ):
                    raise ValueError(f"节点 {node_id} 的产物大小不匹配")
        # Do not inspect arbitrary model JSON in content/structured.  These
        # are the transport envelopes whose paths kxy itself created.
        for key in ("items", "upstream", "inputs"):
            child = value.get(key)
            if isinstance(child, list):
                for index, item in enumerate(child):
                    inspect(item, f"{path_label}.{key}[{index}]")
            elif isinstance(child, dict):
                inspect(child, f"{path_label}.{key}")

    inspect(output, node_id)


def validate_run_snapshot_for_resume(run_id: str, snapshot: dict[str, Any]) -> None:
    run_root = (RUNS_ROOT / str(run_id)).resolve()
    workspace = run_root / "workspace"
    if not resolved_inside(run_root, RUNS_ROOT) or not workspace.is_dir():
        raise ValueError("运行 workspace 不存在，无法恢复")
    sources = snapshot.get("sources", {})
    if not isinstance(sources, dict):
        raise ValueError("运行快照中的资料清单无效")
    for source_id, source in sources.items():
        if not isinstance(source, dict):
            raise ValueError(f"运行快照资料 {source_id} 无效")
        path = _checkpoint_file(workspace, source.get("file_path"), f"资料 {source_id}")
        expected_hash = str(source.get("sha256") or "")
        if not re.fullmatch(r"[a-f0-9]{64}", expected_hash) or hashlib.sha256(path.read_bytes()).hexdigest() != expected_hash:
            raise ValueError(f"运行快照资料 {source_id} 哈希不匹配")
    skills = snapshot.get("skills", [])
    if not isinstance(skills, list):
        raise ValueError("运行快照中的 Skill 清单无效")
    for item in skills:
        if not isinstance(item, dict):
            raise ValueError("运行快照中的 Skill 记录无效")
        raw_path = item.get("path")
        if not isinstance(raw_path, str) or not raw_path.strip() or Path(raw_path).is_absolute() or ".." in Path(raw_path).parts:
            raise ValueError("运行快照 Skill 路径无效")
        skill_root = (workspace / raw_path).resolve()
        if not resolved_inside(skill_root, workspace) or skill_root.is_symlink() or not skill_root.is_dir():
            raise ValueError(f"Skill {item.get('id', '')} 运行快照目录无效")
        try:
            actual_hash = snapshot_skill(skill_root)[0]
        except (OSError, ValueError) as exc:
            raise ValueError(f"Skill {item.get('id', '')} 运行快照无法校验") from exc
        if actual_hash != str(item.get("snapshot_hash") or ""):
            raise ValueError(f"Skill {item.get('id', '')} 运行快照哈希不匹配")


def validate_resume_grants(snapshot: dict[str, Any]) -> None:
    grants = snapshot.get("grants", [])
    if not isinstance(grants, list):
        raise ValueError("运行快照中的输出授权清单无效")
    for item in grants:
        if not isinstance(item, dict) or not item.get("id"):
            continue
        with connect_db() as db:
            current = db.execute(
                "SELECT id, canonical_path, revoked_at FROM grants WHERE id=?",
                (str(item["id"]),),
            ).fetchone()
        expected_path = str(item.get("canonical_path") or "")
        if current is None or current["revoked_at"] is not None or str(current["canonical_path"]) != expected_path:
            raise ValueError(f"输出授权 {item['id']} 已撤销、删除或路径已变化")
        root = Path(str(current["canonical_path"]))
        if root.resolve() != root or not root.is_dir():
            raise ValueError(f"输出授权 {item['id']} 目录已不可用")


def resume_checkpoint_plan(run_id: str, run_row: sqlite3.Row, db: sqlite3.Connection) -> tuple[dict[str, Any], dict[str, Any], list[str]]:
    try:
        snapshot = json.loads(run_row["snapshot"] or "{}")
    except (TypeError, ValueError, json.JSONDecodeError) as exc:
        raise ValueError("运行快照格式无效，无法恢复") from exc
    if not isinstance(snapshot, dict) or snapshot.get("checkpoint_version") != CHECKPOINT_VERSION:
        raise ValueError("该历史运行没有 V8 durable checkpoint，不能安全恢复")
    validate_run_snapshot_for_resume(run_id, snapshot)
    validate_resume_grants(snapshot)
    compiled = snapshot.get("compiled") or {}
    records = compiled.get("nodes") if isinstance(compiled, dict) else None
    if not isinstance(records, list) or not records:
        raise ValueError("运行快照缺少可恢复的编译节点")
    expected_ids = [str(item.get("flat_id")) for item in records if isinstance(item, dict) and item.get("flat_id")]
    if len(expected_ids) != len(set(expected_ids)):
        raise ValueError("运行快照包含重复的编译节点")
    rows = db.execute("SELECT * FROM run_nodes WHERE run_id=?", (run_id,)).fetchall()
    by_id = {str(row["node_id"]): row for row in rows}
    missing = [node_id for node_id in expected_ids if node_id not in by_id]
    if missing:
        raise ValueError("运行记录缺少节点 checkpoint：" + ", ".join(missing[:8]))
    approvals = {
        str(row["node_id"]): row
        for row in db.execute("SELECT * FROM run_approvals WHERE run_id=?", (run_id,)).fetchall()
    }
    if any(row["status"] == "rejected" for row in approvals.values()):
        raise ValueError("运行包含已拒绝的人工决定，不能通过恢复绕过")
    reused: dict[str, Any] = {}
    rerun: list[str] = []
    for item in records:
        if not isinstance(item, dict):
            continue
        node_id = str(item.get("flat_id") or "")
        node_row = by_id[node_id]
        if node_row["status"] != "succeeded":
            rerun.append(node_id)
            continue
        if node_row["output"] in (None, ""):
            raise ValueError(f"成功节点 {node_id} 没有 durable checkpoint")
        try:
            output = json.loads(node_row["output"])
        except (TypeError, ValueError, json.JSONDecodeError) as exc:
            raise ValueError(f"成功节点 {node_id} 的 checkpoint 格式无效") from exc
        validate_checkpoint_output((RUNS_ROOT / str(run_id) / "workspace").resolve(), output, node_id)
        node_type = str(item.get("type") or "")
        if node_type == "human":
            approval = approvals.get(node_id)
            if approval is None or approval["status"] != "approved":
                raise ValueError(f"人工节点 {node_id} 没有已确认的审批 checkpoint")
        reused[node_id] = output
    return snapshot, reused, rerun


def load_success_checkpoints(run_id: str, attempt_no: int, snapshot: dict[str, Any]) -> dict[str, Any]:
    if attempt_no <= 1:
        return {}
    with connect_db() as db:
        run_row = db.execute("SELECT * FROM runs WHERE id=?", (run_id,)).fetchone()
        if run_row is None:
            raise ValueError("运行记录不存在，无法加载 checkpoint")
        _snapshot, reused, _rerun = resume_checkpoint_plan(
            run_id,
            run_row,
            db,
        )
    return reused


def persist_granted_outputs(run_id, workflow, outputs):
    workspace = RUNS_ROOT / run_id / 'workspace'
    snapshot = json.loads((workspace / 'manifest.json').read_text())
    saved = []
    compiled = flatten_workflow(workflow)

    def receipt_files_match(root: Path, expected: dict[str, Any], boundary: Path | None = None) -> bool:
        if root.is_symlink() or not root.is_dir():
            return False
        resolved_root = root.resolve()
        if boundary is not None and not resolved_inside(resolved_root, boundary):
            return False
        receipt_path = root / "export-receipt.json"
        if receipt_path.is_symlink() or not receipt_path.is_file():
            return False
        try:
            receipt = json.loads(receipt_path.read_text(encoding="utf-8"))
        except (OSError, ValueError, json.JSONDecodeError):
            return False
        if receipt != expected:
            return False
        for item in expected.get("files", []):
            if not isinstance(item, dict):
                return False
            raw_path = str(item.get("path") or "")
            relative = Path(raw_path)
            if not raw_path or relative.is_absolute() or ".." in relative.parts:
                return False
            candidate = root / relative
            resolved = candidate.resolve()
            if candidate.is_symlink() or not resolved_inside(resolved, root) or not candidate.is_file() or resolved.is_symlink():
                return False
            if hashlib.sha256(candidate.read_bytes()).hexdigest() != str(item.get("sha256") or ""):
                return False
            try:
                expected_size = int(item.get("size", -1))
            except (TypeError, ValueError):
                return False
            if candidate.stat().st_size != expected_size:
                return False
        return True

    def safe_target_name(value: Any, fallback: str) -> str:
        candidate = Path(str(value or fallback)).name
        return candidate if candidate not in {"", ".", ".."} else fallback

    for record in compiled['records']:
        node = record['node']
        flat_id = record['flat_id']
        if node['type'] != 'container' or flat_id not in outputs:
            continue
        output = outputs[flat_id]
        if output.get('status') == 'skipped':
            continue
        data = node.get('data') or {}
        export_formats = normalize_container_formats(output.get('export_formats', data.get('export_formats')))
        allowed_file_extensions = normalize_allowed_file_extensions(
            output.get('allowed_file_extensions', data.get('allowed_file_extensions'))
        )
        json_mode = output.get('json_mode', data.get('json_mode', 'full'))
        if json_mode not in JSON_MODES:
            raise RuntimeError('容器 JSON 模式无效')

        file_payloads: list[tuple[str, bytes]] = []
        if 'json' in export_formats:
            json_payload = output if json_mode == 'full' else output.get('content', output)
            file_payloads.append(("result.json", json_text(redact_object(json_payload)).encode("utf-8")))
        if 'markdown' in export_formats:
            file_payloads.append(("result.md", str(output.get('text', '')).encode('utf-8')))
        if 'text' in export_formats:
            file_payloads.append(("result.txt", str(output.get('text', '')).encode('utf-8')))
        file_payloads.append(("provenance.json", json_text(redact_object(snapshot)).encode('utf-8')))

        artifacts = []
        for artifact in output.get('artifacts', []):
            if not isinstance(artifact, dict) or not artifact.get('path'):
                continue
            extension = Path(str(artifact.get('name') or artifact.get('path'))).suffix.lower()
            if allowed_file_extensions is not None and extension not in allowed_file_extensions:
                continue
            artifacts.append(artifact)
        copied_artifacts: list[tuple[str, Path]] = []
        used_names = {name for name, _content in file_payloads}
        for artifact in artifacts:
            raw_source = Path(str(artifact['path']))
            if raw_source.is_absolute() or '..' in raw_source.parts:
                raise RuntimeError('产物路径不安全')
            source_candidate = workspace / raw_source
            if source_candidate.is_symlink():
                raise RuntimeError('产物路径不安全')
            source = source_candidate.resolve()
            if not resolved_inside(source, workspace) or not source.is_file() or source.is_symlink():
                raise RuntimeError('产物路径不安全')
            filename = safe_target_name(artifact.get('name'), source.name)
            if filename in used_names:
                filename = f"{source.stem}-{str(artifact.get('sha256') or '')[:8]}{source.suffix}"
            if filename in used_names:
                raise RuntimeError('拒绝覆盖已存在的文件')
            used_names.add(filename)
            copied_artifacts.append((f"files/{filename}", source))

        receipt_files = [
            {"path": name, "size": len(content), "sha256": hashlib.sha256(content).hexdigest()}
            for name, content in file_payloads
        ]
        receipt_files.extend(
            {
                "path": name,
                "size": source.stat().st_size,
                "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
            }
            for name, source in copied_artifacts
        )
        grant_id = (node.get('data') or {}).get('grant_id')
        receipt = {
            "version": "kxy.export-receipt.v8",
            "run_id": str(run_id),
            "node_id": record['path'],
            "flat_node_id": flat_id,
            "grant_id": str(grant_id or ""),
            "export_formats": export_formats,
            "allowed_file_extensions": allowed_file_extensions,
            "json_mode": json_mode,
            "files": receipt_files,
        }

        local_root = workspace / 'exports'
        if local_root.is_symlink() or (local_root.exists() and not local_root.is_dir()):
            raise RuntimeError('运行 exports 目录路径无效')
        if local_root.exists() and not resolved_inside(local_root.resolve(), workspace):
            raise RuntimeError('运行 exports 目录越过 workspace 边界')
        local_root.mkdir(parents=True, exist_ok=True)
        local = local_root / flat_id
        if local.is_dir() and receipt_files_match(local, receipt, workspace):
            local_reused = True
        else:
            local_reused = False
            if local.exists() or local.is_symlink():
                local = local_root / f"{flat_id}--attempt-{uuid.uuid4().hex[:10]}"
                while local.exists() or local.is_symlink():
                    local = local_root / f"{flat_id}--attempt-{uuid.uuid4().hex[:10]}"
            local.mkdir(parents=True, exist_ok=False)

        def write_export_tree(root: Path) -> None:
            for name, content in file_payloads:
                target = root / name
                target.parent.mkdir(parents=True, exist_ok=True)
                if target.exists() or target.is_symlink():
                    raise RuntimeError('拒绝覆盖已存在的导出文件')
                target.write_bytes(content)
            for name, source in copied_artifacts:
                target = root / name
                target.parent.mkdir(parents=True, exist_ok=True)
                if target.exists() or target.is_symlink():
                    raise RuntimeError('拒绝覆盖已存在的导出文件')
                shutil.copyfile(source, target)
            receipt_path = root / "export-receipt.json"
            if receipt_path.exists() or receipt_path.is_symlink():
                raise RuntimeError('拒绝覆盖导出 receipt')
            receipt_path.write_text(json_text(receipt), encoding='utf-8')

        if not local_reused:
            write_export_tree(local)

        destination = None
        if grant_id:
            with connect_db() as db:
                grant = db.execute(
                    'SELECT * FROM grants WHERE id=? AND revoked_at IS NULL',
                    (grant_id,),
                ).fetchone()
            if grant is None:
                raise RuntimeError(f"容器 {record['path']} 的保存授权已撤销或不存在")
            root = Path(grant['canonical_path'])
            if str(root.resolve()) != str(root) or not root.is_dir():
                raise RuntimeError('授权目录已移动或被符号链接替换')
            parent = root / 'kxy' / run_id
            if parent.resolve() != parent or not resolved_inside(parent, root):
                raise RuntimeError('输出目录越过授权边界')
            parent.mkdir(parents=True, exist_ok=True)
            destination = parent / flat_id
            if destination.exists() or destination.is_symlink():
                if not destination.is_dir() or not receipt_files_match(destination, receipt, root):
                    raise RuntimeError('授权输出已存在但没有一致的完成 receipt，拒绝覆盖')
            else:
                stage = parent / ('.' + flat_id + '-' + uuid.uuid4().hex)
                try:
                    shutil.copytree(local, stage)
                    stage.rename(destination)
                finally:
                    shutil.rmtree(stage, ignore_errors=True)
        saved.append({
            'node_id': record['path'],
            'flat_node_id': flat_id,
            'path': str(destination) if destination else str(local),
            'granted': bool(grant_id),
        })
    return saved


def workflow_skill_snapshot(workflow: dict[str, Any]) -> list[dict[str, Any]]:
    selected: set[str] = set()

    def visit(document: dict[str, Any]) -> None:
        for node in document.get("nodes", []):
            if node.get("type") == "analyzer":
                values = (node.get("data") or {}).get("skill_ids", [])
                if isinstance(values, str):
                    values = [values]
                if isinstance(values, list):
                    selected.update(str(skill_id) for skill_id in values)
            elif node.get("type") == "blackbox":
                nested = _nested_workflow(node)
                if nested:
                    visit(nested)

    visit(workflow)
    if not selected:
        return []
    with connect_db() as db:
        rows = [db.execute("SELECT id, name, snapshot_hash, metadata FROM skills WHERE id=?", (skill_id,)).fetchone() for skill_id in sorted(selected)]
    missing = [skill_id for skill_id, row in zip(sorted(selected), rows, strict=True) if row is None]
    if missing:
        raise ValueError(f"Workflow references missing local skills: {', '.join(missing)}")
    result: list[dict[str, Any]] = []
    for row in rows:
        if row is None:
            continue
        try:
            metadata = json.loads(row["metadata"] or "{}")
        except (TypeError, ValueError, json.JSONDecodeError):
            metadata = {}
        result.append({"id": row["id"], "name": row["name"], "snapshot_hash": row["snapshot_hash"], "metadata": metadata if isinstance(metadata, dict) else {}})
    return result


def workflow_file_nodes(workflow: dict[str, Any], prefix: tuple[str, ...] = ()) -> list[tuple[str, dict[str, Any]]]:
    """Walk every child workflow for immutable input snapshots."""

    result: list[tuple[str, dict[str, Any]]] = []
    for node in workflow.get("nodes", []):
        if not isinstance(node, dict):
            continue
        node_id = str(node.get("id") or "")
        path = "/".join((*prefix, node_id))
        if node.get("type") == "file":
            result.append((path, node))
        elif node.get("type") == "input":
            data = node.get("data") if isinstance(node.get("data"), dict) else {}
            refs = input_attachment_refs(data)
            for index, ref in enumerate(refs):
                result.append(
                    (
                        f"{path}#attachment-{index + 1}",
                        {"id": node_id, "type": "file", "data": {"file_id": ref["file_id"]}},
                    )
                )
        elif node.get("type") == "blackbox":
            nested = _nested_workflow(node)
            if nested:
                result.extend(workflow_file_nodes(nested, (*prefix, node_id)))
    return result


def workflow_mcp_snapshot(workflow: dict[str, Any]) -> list[dict[str, Any]]:
    """Freeze selected MCP definitions and intended target agents for a run."""

    selections: dict[str, set[str]] = {}

    def visit(document: dict[str, Any]) -> None:
        for node in document.get("nodes", []):
            if node.get("type") == "analyzer":
                data = node.get("data") or {}
                values = data.get("mcp_ids", []) if isinstance(data, dict) else []
                if isinstance(values, str):
                    raise ValueError("Analyzer mcp_ids 必须是数组")
                if not isinstance(values, list):
                    raise ValueError("Analyzer mcp_ids 必须是数组")
                binding = data.get("resolved_model") if isinstance(data, dict) else None
                target = str((binding or {}).get("agent_id") or data.get("agent_id") or data.get("cli") or "codex").strip().lower()
                for raw_id in values:
                    skill_id = str(raw_id or "").strip()
                    if not skill_id:
                        raise ValueError("Analyzer mcp_ids 不能包含空 id")
                    selections.setdefault(skill_id, set()).add(target)
            elif node.get("type") == "blackbox":
                nested = _nested_workflow(node)
                if nested:
                    visit(nested)

    visit(workflow)
    if not selections:
        return []
    try:
        from . import agent_config

        index = agent_config._mcp_load_index()
    except (AttributeError, ValueError) as exc:
        raise ValueError(f"MCP central index unavailable: {exc}") from exc
    result: list[dict[str, Any]] = []
    for server_id in sorted(selections):
        definition = index.get(server_id)
        if definition is None:
            raise ValueError(f"Workflow references missing imported MCP: {server_id}")
        if definition.get("reason"):
            raise ValueError(f"MCP {server_id} is unsupported: {definition['reason']}")
        targets = sorted(selections[server_id])
        for target in targets:
            try:
                supported, reason = agent_config._mcp_support(target)
            except ValueError as exc:
                raise ValueError(str(exc)) from exc
            if not supported:
                raise ValueError(f"MCP {server_id} cannot run on {target}: {reason}")
        safe_definition = {
            key: definition[key]
            for key in ("id", "label", "transport", "enabled", "url", "headers", "command", "args", "env", "env_templates")
            if key in definition
        }
        if agent_config._mcp_contains_secret_literal(safe_definition):
            raise ValueError(f"MCP {server_id} contains a secret literal")
        result.append(
            {
                "id": server_id,
                "name": str(definition.get("label") or server_id),
                "transport": str(definition.get("transport") or "unknown"),
                "source_agent": str(definition.get("source") or ((definition.get("agents") or [""])[0] if isinstance(definition.get("agents"), list) else "")),
                "source_agents": [str(item) for item in definition.get("source_agents", definition.get("agents", [])) if isinstance(item, str)],
                "target_agents": targets,
                "definition": safe_definition,
                "definition_sha256": agent_config._mcp_definition_digest(definition),
            }
        )
    return result


def workflow_grant_snapshot(workflow: dict[str, Any]) -> list[dict[str, Any]]:
    """Record every output grant used by top-level or nested containers."""
    records: list[dict[str, Any]] = []

    def visit(document: dict[str, Any], prefix: tuple[str, ...] = ()) -> None:
        for node in document.get("nodes", []):
            node_id = str(node.get("id", ""))
            node_path = "/".join((*prefix, node_id))
            if node.get("type") == "container":
                grant_id = str((node.get("data") or {}).get("grant_id") or "").strip()
                if grant_id:
                    with connect_db() as db:
                        row = db.execute(
                            "SELECT id, canonical_path, revoked_at, created_at FROM grants WHERE id=?",
                            (grant_id,),
                        ).fetchone()
                    records.append(
                        {
                            "id": grant_id,
                            "node_path": node_path,
                            "canonical_path": row["canonical_path"] if row is not None else None,
                            "revoked_at": row["revoked_at"] if row is not None else None,
                            "created_at": row["created_at"] if row is not None else None,
                            "status": "active" if row is not None and row["revoked_at"] is None else "unavailable",
                        }
                    )
            elif node.get("type") == "blackbox":
                nested = _nested_workflow(node)
                if nested:
                    visit(nested, (*prefix, node_id))

    visit(workflow)
    return records


def snapshot_workflow_bindings(workflow: dict[str, Any]) -> dict[str, Any]:
    """Freeze optional global agent/model bindings without embedding credentials."""
    result = normalize_workflow(workflow)
    if snapshot_model_binding is None:
        return result

    def visit(document: dict[str, Any]) -> None:
        for node in document.get("nodes", []):
            if node.get("type") == "analyzer":
                data = dict(node.get("data") or {})
                binding = snapshot_model_binding(data)
                data["resolved_model"] = redact_object(binding)
                node["data"] = data
            elif node.get("type") == "blackbox":
                nested = _nested_workflow(node)
                if nested:
                    data = node.setdefault("data", {})
                    data["workflow"] = normalize_workflow(nested)
                    visit(data["workflow"])

    visit(result)
    return result


def flat_component_id(path: tuple[str, ...]) -> str:
    """Create a bounded collision-free LFX id while retaining the original path separately."""
    parts = [re.sub(r"[^A-Za-z0-9_-]", "_", part) for part in path]
    candidate = "__".join(parts)
    if len(candidate) <= 96:
        return candidate
    digest = hashlib.sha256("/".join(path).encode("utf-8")).hexdigest()[:20]
    return f"{candidate[:75]}__{digest}"


def flatten_workflow(workflow: dict[str, Any]) -> dict[str, Any]:
    """Compile nested blackboxes into one flat graph for the single LFX scheduler."""
    records: list[dict[str, Any]] = []
    compiled_edges: list[dict[str, Any]] = []
    blackboxes: dict[str, dict[str, Any]] = {}
    path_by_flat: dict[str, str] = {}

    def visit(document: dict[str, Any], prefix: tuple[str, ...]) -> dict[str, str]:
        boundary: dict[str, tuple[str, str]] = {}
        marker_flats: dict[str, str] = {}
        child_records_start = len(records)
        for node in document.get("nodes", []):
            node_id = str(node["id"])
            node_path = (*prefix, node_id)
            path_label = "/".join(node_path)
            if node.get("type") == "blackbox":
                nested = _nested_workflow(node)
                if nested is None:
                    raise ValueError(f"Blackbox {path_label} is missing data.workflow")
                if loop_enabled(node):
                    # V9 loops are one scheduler vertex.  Their child graph is
                    # compiled afresh by KxyLoopComponent for each round, so
                    # child vertices must not leak into the outer DAG.
                    flat_id = flat_component_id(node_path)
                    if flat_id in path_by_flat:
                        raise ValueError(f"Compiled node id collision at {path_label}")
                    path_by_flat[flat_id] = path_label
                    records.append({"flat_id": flat_id, "path": path_label, "node": node, "prefix": prefix, "loop": True})
                    boundary[node_id] = (flat_id, flat_id)
                    blackboxes[path_label] = {
                        "input": flat_id,
                        "output": flat_id,
                        "inner": [],
                        "depth": len(node_path) - 1,
                        "loop": True,
                    }
                    continue
                child = visit(normalize_workflow(nested), node_path)
                boundary[node_id] = (child["input"], child["output"])
                inner_flats = [item["flat_id"] for item in records if item["path"].startswith(path_label + "/")]
                blackboxes[path_label] = {
                    "input": child["input"],
                    "output": child["output"],
                    "inner": inner_flats,
                    "depth": len(node_path) - 1,
                }
                continue
            flat_id = flat_component_id(node_path)
            if flat_id in path_by_flat:
                raise ValueError(f"Compiled node id collision at {path_label}")
            path_by_flat[flat_id] = path_label
            record = {"flat_id": flat_id, "path": path_label, "node": node, "prefix": prefix}
            records.append(record)
            boundary[node_id] = (flat_id, flat_id)
            if node.get("type") in SUBFLOW_MARKERS:
                marker_flats[node["type"]] = flat_id

        for edge in document.get("edges", []):
            source = str(edge["source"])
            target = str(edge["target"])
            source_flat = boundary[source][1]
            target_flat = boundary[target][0]
            compiled_edges.append(
                {
                    "source": source_flat,
                    "target": target_flat,
                    "sourceHandle": str(edge.get("sourceHandle") or "result"),
                    "targetHandle": str(edge.get("targetHandle") or "items"),
                    "path": "/".join((*prefix, source, target)),
                }
            )

        # A blackbox may intentionally contain a fixed file/text source. Those
        # nodes have no visible input handle, so validation keeps their public
        # schema intact and the compiled graph adds an internal activation gate.
        if prefix and "subflow_input" in marker_flats:
            incoming_ids = {
                str(edge.get("target"))
                for edge in document.get("edges", [])
                if isinstance(edge, dict)
            }
            for node in document.get("nodes", []):
                node_type = node.get("type")
                node_id = str(node.get("id"))
                if node_type not in {"file", "text", "input", "blackbox"} or node_id in incoming_ids:
                    continue
                target_flat = boundary[node_id][0]
                compiled_edges.append(
                    {
                        "source": marker_flats["subflow_input"],
                        "target": target_flat,
                        "sourceHandle": "result",
                        "targetHandle": "gate" if node_type != "blackbox" else "items",
                        "hidden": True,
                        "path": "/".join((*prefix, "<gate>", node_id)),
                    }
                )

        if prefix and marker_flats:
            if "subflow_input" not in marker_flats or "subflow_output" not in marker_flats:
                raise ValueError(f"Blackbox {('/'.join(prefix))} has incomplete boundary markers")
            return {"input": marker_flats["subflow_input"], "output": marker_flats["subflow_output"]}
        if prefix:
            raise ValueError(f"Blackbox {('/'.join(prefix))} has no boundary markers")
        return {}

    visit(workflow, ())
    return {
        "records": records,
        "edges": compiled_edges,
        "blackboxes": blackboxes,
        "path_by_flat": path_by_flat,
    }


def build_lfx_graph(
    workflow: dict[str, Any],
    state: RunState,
    *,
    component_prefix: str = "",
    display_path_prefix: str = "",
    loop_context: dict[str, Any] | None = None,
) -> tuple[Any, dict[str, Any], list[str]]:
    if LFX_IMPORT_ERROR:
        raise RuntimeError(f"LFX 1.12.0 is unavailable: {LFX_IMPORT_ERROR}")
    components: dict[str, Any] = {}
    compiled = flatten_workflow(workflow)
    id_map = {
        str(record["flat_id"]): scoped_component_id(component_prefix, str(record["flat_id"]))
        for record in compiled["records"]
    }
    path_prefix = display_path_prefix.strip("/")
    scoped_paths = {
        id_map[str(flat_id)]: "/".join(item for item in (path_prefix, path) if item)
        for flat_id, path in compiled["path_by_flat"].items()
    }
    state.node_paths.update(scoped_paths)
    for path, info in compiled["blackboxes"].items():
        scoped_path = "/".join(item for item in (path_prefix, path) if item)
        state.blackbox_paths[scoped_path] = {
            **info,
            "input": id_map.get(str(info.get("input")), scoped_component_id(component_prefix, str(info.get("input") or ""))),
            "output": id_map.get(str(info.get("output")), scoped_component_id(component_prefix, str(info.get("output") or ""))),
            "inner": [id_map.get(str(item), scoped_component_id(component_prefix, str(item))) for item in info.get("inner", [])],
        }
    for record in compiled["records"]:
        local_id = str(record["flat_id"])
        node_id = id_map[local_id]
        original_path = scoped_paths[node_id]
        node = record["node"]
        node_type = node["type"]
        data = dict(node.get("data") or {})
        if node_type == "file":
            component = KxySourceComponent(_id=node_id, payload=Data(data=source_payload(str(data.get("file_id", "")), state.workspace)))
        elif node_type in {"text", "input"}:
            payload = (
                input_source_payload(data, state.workspace)
                if node_type == "input"
                else {"status": "source", "active": True, "text": str(data.get("text", "")), "name": data.get("label", "Text input")}
            )
            component = KxySourceComponent(_id=node_id, payload=Data(data=payload))
        elif node_type == "filter":
            component = KxyFilterComponent(_id=node_id)
            component.node_config = normalize_filter_config(data)
        elif node_type == "subflow_input":
            component = KxySubflowInputComponent(_id=node_id)
        elif node_type == "subflow_output":
            component = KxySubflowOutputComponent(_id=node_id)
        elif node_type == "blackbox" and loop_enabled(node):
            component = KxyLoopComponent(_id=node_id)
            component.node_config = data
        elif node_type == "analyzer":
            component = KxyAnalyzerComponent(_id=node_id, prompt=str(data.get("prompt", "")))
            component.runner = lambda prompt, payloads, comp, st=state: execute_cli(st, prompt, payloads, comp)
            component.node_config = data
        elif node_type == "human":
            component = KxyHumanComponent(
                _id=node_id,
                content=str(data.get("content", ""))[:MAX_TEXT_BYTES],
                confirm_label=str(data.get("confirm_label") or "确认继续")[:120],
            )
            component.allow_return = bool(data.get("review_gate") is True and loop_context is not None)
            component.node_config = data
        elif node_type == "condition":
            component = KxyConditionComponent(
                _id=node_id,
                field=str(data.get("field", "status")),
                operator=str(data.get("operator", "equals")),
                expected=str(data.get("expected", "succeeded")),
                missing_strategy=str(data.get("missing_strategy", "error") or "error"),
            )
        elif node_type == "container":
            component = KxyContainerComponent(_id=node_id)
            component.node_config = data
        else:
            raise ValueError(f"Unsupported node type: {node_type}")
        component.run_state = state
        component.original_path = original_path
        components[node_id] = component
    graph = Graph(
        flow_id=f"kxy-{state.run_id}-{component_prefix or 'root'}",
        flow_name=str(workflow.get("name", "kxy workflow")),
    )
    for node_id, component in components.items():
        graph.add_component(component, node_id)
    for edge in compiled["edges"]:
        source_handle = str(edge.get("sourceHandle") or "result")
        target_handle = str(edge.get("targetHandle") or "items")
        graph.add_component_edge(
            id_map[str(edge["source"])],
            (source_handle, target_handle),
            id_map[str(edge["target"])],
        )
    graph.prepare()
    terminal_nodes = [
        node_id
        for node_id in components
        if not any(id_map[str(edge["source"])] == node_id for edge in compiled["edges"])
    ]
    return graph, components, terminal_nodes


def update_blackbox_statuses(run_id: str, status: str, message: str | None = None) -> None:
    """Mirror inner LFX state to durable outer cards without executing them."""
    with connect_db() as db:
        rows = db.execute("SELECT node_id FROM run_nodes WHERE run_id=? AND instr(node_id, '/') > 0", (run_id,)).fetchall()
        for row in rows:
            db.execute(
                "UPDATE run_nodes SET status=?, message=COALESCE(?, message), finished_at=CASE WHEN ? IN ('succeeded','failed','rejected','cancelled','interrupted') THEN COALESCE(finished_at, ?) ELSE finished_at END WHERE run_id=? AND node_id=?",
                (status, message, status, utc_now(), run_id, row["node_id"]),
            )
        # Top-level blackbox ids do not contain a slash; callers update those
        # explicitly from the compiled map when needed.


def mirror_blackbox_results(run_id: str, state: RunState, outputs: dict[str, Any], components: dict[str, Any]) -> None:
    """Map compiled inner results and terminal status back to every blackbox path."""
    if not state.blackbox_paths:
        return
    with connect_db() as db:
        for path, info in state.blackbox_paths.items():
            if info.get("loop"):
                # The loop vertex owns its terminal result directly; its child
                # rounds are durable evidence, not another outer blackbox.
                continue
            inner_ids = info.get("inner", [])
            statuses = [
                row["status"]
                for inner_id in inner_ids
                for row in [db.execute("SELECT status FROM run_nodes WHERE run_id=? AND node_id=?", (run_id, inner_id)).fetchone()]
                if row is not None
            ]
            if any(item == "failed" for item in statuses):
                box_status = "failed"
            elif any(item == "rejected" for item in statuses):
                box_status = "rejected"
            elif any(item == "cancelled" for item in statuses):
                box_status = "cancelled"
            elif any(item == "waiting" for item in statuses):
                box_status = "waiting"
            elif statuses and all(item in {"succeeded", "skipped", "interrupted"} for item in statuses):
                box_status = "succeeded" if "succeeded" in statuses else "skipped"
            else:
                box_status = "running"
            db.execute(
                "UPDATE run_nodes SET status=?, output=COALESCE(?, output), message=COALESCE(?, message), "
                "finished_at=CASE WHEN ? IN ('succeeded','failed','rejected','cancelled','interrupted','skipped') THEN COALESCE(finished_at, ?) ELSE finished_at END "
                "WHERE run_id=? AND node_id=?",
                (
                    box_status,
                    json_text(redact_object(outputs.get(path))) if outputs.get(path) is not None else None,
                    None,
                    box_status,
                    utc_now(),
                    run_id,
                    path,
                ),
            )


def finish_run_attempt(run_id: str, attempt_no: int, status: str, error: str | None = None) -> None:
    with connect_db() as db:
        db.execute(
            "UPDATE run_attempts SET status=?, error=?, finished_at=? "
            "WHERE run_id=? AND attempt_no=? AND status IN ('queued','running')",
            (status, redact(error) if error else None, utc_now(), run_id, attempt_no),
        )


def run_workflow(run_id: str, attempt_no: int = 1):
    with connect_db() as db:
        row = db.execute('SELECT * FROM runs WHERE id=?', (run_id,)).fetchone()
    if row is None or row["status"] == "interrupted":
        return
    try:
        snapshot = json.loads(row['snapshot'])
        workflow = snapshot['workflow']
    except (TypeError, ValueError, KeyError, json.JSONDecodeError):
        snapshot = {}
        workflow = {}
    workspace = Path(row['run_dir']).resolve() / 'workspace'
    with RUN_LOCK:
        cancel_event = RUN_CANCEL_EVENTS.setdefault(run_id, threading.Event())
    state = RunState(
        run_id,
        workspace,
        cancel_event,
        snapshot.get('settings', {}) if isinstance(snapshot, dict) else {},
        attempt_no=attempt_no,
    )
    try:
        if state.cancel_event.is_set():
            raise RuntimeError('Run cancelled by user')
        state.resume_outputs = load_success_checkpoints(run_id, attempt_no, snapshot) if attempt_no > 1 else {}
        state.node_workspace_root.mkdir(parents=True, exist_ok=True)
        (workspace/'manifest.json').write_text(json_text(snapshot), encoding='utf-8')
        with connect_db() as db:
            db.execute(
                "UPDATE runs SET status='running',started_at=COALESCE(started_at,?),error=NULL WHERE id=? AND status='pending'",
                (utc_now(), run_id),
            )
            db.execute(
                "UPDATE run_attempts SET status='running',started_at=COALESCE(started_at,?) WHERE run_id=? AND attempt_no=? AND status='queued'",
                (utc_now(), run_id, attempt_no),
            )
        append_event(run_id, 'run_started', {'run_id': run_id, 'attempt': attempt_no, 'engine': 'lfx', 'lfx_version': '1.12.0', 'reused_nodes': list(state.resume_outputs)})
        graph, components, terminal_nodes = build_lfx_graph(workflow, state)

        async def consume_graph():
            async for result in graph.async_start(inputs=[{}]):
                if hasattr(result, 'vertex'):
                    append_event(run_id, 'lfx_vertex', {'node_id': result.vertex.id, 'valid': getattr(result, 'valid', True), 'attempt': attempt_no})

        asyncio.run(consume_graph())
        if state.cancel_event.is_set():
            raise RuntimeError('Run cancelled by user')
        for excluded in getattr(graph, 'conditionally_excluded_vertices', set()):
            if excluded in components and excluded not in state.node_finished:
                state.finish_node(excluded, 'skipped', message='条件选中了另一条分支')
        outputs = {key: payload_from_data(comp.last_result) for key, comp in components.items() if hasattr(comp, 'last_result')}
        for path, info in state.blackbox_paths.items():
            output_id = info.get("output")
            if output_id and output_id in outputs:
                outputs[path] = outputs[output_id]
        mirror_blackbox_results(run_id, state, outputs, components)
        grants = persist_granted_outputs(run_id, workflow, outputs)
        artifacts = []
        from urllib.parse import quote
        for path in sorted(workspace.rglob('*')):
            if path.is_symlink() or not path.is_file():
                continue
            relative = path.relative_to(workspace)
            parts = relative.parts
            output_index = next((index for index, part in enumerate(parts) if part == "outputs"), -1)
            node_root_index = (
                0
                if parts and parts[0] == "nodes"
                else 2
                if len(parts) >= 3 and parts[0] == "attempts" and parts[2] == "nodes"
                else -1
            )
            is_node_output = node_root_index >= 0 and output_index >= node_root_index + 2
            if parts[0] != 'exports' and not is_node_output:
                continue
            content = path.read_bytes()
            artifacts.append({'name': path.name, 'path': str(relative), 'size': len(content), 'sha256': hashlib.sha256(content).hexdigest(), 'url': f'/api/runs/{run_id}/artifact/{quote(str(relative))}'})
        payload = {'run_id': run_id, 'attempt': attempt_no, 'engine': 'lfx', 'outputs': outputs, 'artifacts': artifacts, 'granted_outputs': grants, 'created_at': utc_now()}
        manifest = write_run_manifest(run_id, workspace, payload)
        with connect_db() as db:
            db.execute(
                "UPDATE runs SET status='succeeded',output_manifest=?,finished_at=? WHERE id=? AND status='running'",
                (json_text({**payload, 'path': str(manifest)}), utc_now(), run_id),
            )
        finish_run_attempt(run_id, attempt_no, 'succeeded')
        append_event(run_id, 'run_finished', {'status': 'succeeded', 'attempt': attempt_no, 'output_nodes': terminal_nodes, 'artifacts': len(artifacts)})
    except RunRejected as exc:
        status = 'rejected'
        message = redact(str(exc))
        with connect_db() as db:
            current = db.execute("SELECT status FROM runs WHERE id=?", (run_id,)).fetchone()
        if current and current["status"] == "interrupted":
            return
        with connect_db() as db:
            db.execute('UPDATE runs SET status=?,error=?,finished_at=? WHERE id=?', (status, message, utc_now(), run_id))
            db.execute("UPDATE run_nodes SET status=?,message=? WHERE run_id=? AND status IN ('pending','running','waiting')", ('skipped', message, run_id))
        finish_run_attempt(run_id, attempt_no, status, message)
        mirror_blackbox_results(run_id, state, {}, {})
        append_event(run_id, 'run_finished', {'status': status, 'attempt': attempt_no, 'error': message})
    except RunCancelled as exc:
        status = 'cancelled'
        message = redact(str(exc))
        with connect_db() as db:
            current = db.execute("SELECT status FROM runs WHERE id=?", (run_id,)).fetchone()
        if current and current["status"] == "interrupted":
            return
        with connect_db() as db:
            db.execute('UPDATE runs SET status=?,error=?,finished_at=? WHERE id=?', (status, message, utc_now(), run_id))
            db.execute("UPDATE run_nodes SET status=?,message=? WHERE run_id=? AND status IN ('pending','running','waiting')", ('cancelled', message, run_id))
        finish_run_attempt(run_id, attempt_no, status, message)
        mirror_blackbox_results(run_id, state, {}, {})
        append_event(run_id, 'run_finished', {'status': status, 'attempt': attempt_no, 'error': message})
    except Exception as exc:
        with connect_db() as db:
            rejected = db.execute("SELECT 1 FROM run_approvals WHERE run_id=? AND status='rejected' LIMIT 1", (run_id,)).fetchone() is not None
            current = db.execute("SELECT status FROM runs WHERE id=?", (run_id,)).fetchone()
        if current and current["status"] == "interrupted":
            return
        status = 'rejected' if rejected else ('cancelled' if state.cancel_event.is_set() else 'failed')
        message = redact(str(exc))
        if rejected:
            with connect_db() as db:
                approval = db.execute("SELECT note FROM run_approvals WHERE run_id=? AND status='rejected' ORDER BY decided_at DESC LIMIT 1", (run_id,)).fetchone()
            message = redact((approval["note"] if approval and approval["note"] else "人工确认已拒绝"))
        with connect_db() as db:
            db.execute('UPDATE runs SET status=?,error=?,finished_at=? WHERE id=?', (status, message, utc_now(), run_id))
            db.execute("UPDATE run_nodes SET status=?,message=? WHERE run_id=? AND status IN ('pending','running','waiting')", ('cancelled' if status == 'cancelled' else 'skipped', message, run_id))
        finish_run_attempt(run_id, attempt_no, status, message)
        mirror_blackbox_results(run_id, state, {}, {})
        append_event(run_id, 'run_finished', {'status': status, 'attempt': attempt_no, 'error': message})
    finally:
        with RUN_LOCK:
            if RUN_CANCEL_EVENTS.get(run_id) is cancel_event:
                RUN_CANCEL_EVENTS.pop(run_id, None)
        # A cancelled worker may still be unwinding when a guarded resume is
        # queued.  Only clean this worker's node workspace; never remove a
        # newer attempt's state and never pop its process handle.
        node_workspace_root = state.node_workspace_root
        run_root = (RUNS_ROOT / str(run_id)).resolve()
        if (
            resolved_inside(run_root, RUNS_ROOT)
            and node_workspace_root.is_dir()
            and resolved_inside(node_workspace_root, run_root)
        ):
            for state_root in node_workspace_root.rglob('.cli-state'):
                if (
                    not state_root.is_symlink()
                    and state_root.is_dir()
                    and resolved_inside(state_root, node_workspace_root)
                ):
                    shutil.rmtree(state_root, ignore_errors=True)


def start_background_run(run_id: str, attempt_no: int = 1) -> None:
    cancel_event = threading.Event()
    with RUN_LOCK:
        RUN_CANCEL_EVENTS[run_id] = cancel_event
    threading.Thread(target=run_workflow, args=(run_id, attempt_no), name=f"kxy-run-{run_id[:8]}-{attempt_no}", daemon=True).start()


def get_run_payload(run_id: str) -> dict[str, Any]:
    with connect_db() as db:
        run = db.execute("SELECT * FROM runs WHERE id=?", (run_id,)).fetchone()
        if run is None:
            raise HTTPException(status_code=404, detail="Run not found")
        nodes = db.execute("SELECT * FROM run_nodes WHERE run_id=? ORDER BY rowid", (run_id,)).fetchall()
        approvals = db.execute("SELECT * FROM run_approvals WHERE run_id=? ORDER BY created_at, node_id", (run_id,)).fetchall()
        attempts = db.execute("SELECT * FROM run_attempts WHERE run_id=? ORDER BY attempt_no", (run_id,)).fetchall()
        loop_rounds = db.execute(
            "SELECT * FROM loop_rounds WHERE run_id=? ORDER BY loop_path, attempt_no, round_no",
            (run_id,),
        ).fetchall()
        events = db.execute("SELECT id, event_type, payload, created_at FROM run_events WHERE run_id=? ORDER BY id DESC LIMIT 80", (run_id,)).fetchall()
    result = dict(run)
    result["snapshot"] = json.loads(result.pop("snapshot"))
    result["output_manifest"] = json.loads(result["output_manifest"]) if result.get("output_manifest") else None
    result["nodes"] = []
    compiled_paths = {
        item.get("flat_id"): item.get("path")
        for item in (result["snapshot"].get("compiled", {}).get("nodes", []) if isinstance(result["snapshot"], dict) else [])
        if isinstance(item, dict)
    }
    for node in nodes:
        item = dict(node)
        item["output"] = json.loads(item["output"]) if item.get("output") else None
        item["node_path"] = compiled_paths.get(item["node_id"], item["node_id"])
        result["nodes"].append(item)
    result["approvals"] = [approval_payload(row) for row in approvals]
    result["loop_rounds"] = [loop_round_payload(row) for row in loop_rounds]
    result["attempts"] = []
    for attempt in attempts:
        item = dict(attempt)
        for key in ("reused_nodes", "rerun_nodes"):
            try:
                item[key] = json.loads(item[key] or "[]")
            except (TypeError, ValueError, json.JSONDecodeError):
                item[key] = []
        result["attempts"].append(item)
    result["events"] = []
    for event in reversed(events):
        item = dict(event)
        item["payload"] = json.loads(item["payload"])
        result["events"].append(item)
    return result


class WorkflowSaveRequest(BaseModel):
    model_config = ConfigDict(extra="allow")
    id: str | None = None
    name: str = "未命名研究流程"
    workflow: dict[str, Any]


class PortableExportRequest(BaseModel):
    model_config = ConfigDict(extra="ignore")
    workflow_id: str | None = None
    workflow: dict[str, Any] | None = None
    include_file_ids: list[str] = Field(default_factory=list)
    include_skill_ids: list[str] = Field(default_factory=list)


class PortableImportCommitRequest(BaseModel):
    model_config = ConfigDict(extra="ignore")
    preview_id: str
    bindings: dict[str, Any] = Field(default_factory=dict)


class SettingsRequest(BaseModel):
    # Optional fields make PUT genuinely backward-compatible: a legacy
    # client can update one appearance key without resetting V3 keys, while
    # the response always contains the complete typed settings object.
    model_config = ConfigDict(extra="ignore")
    palette: str | None = None
    accent: str | None = None
    canvas: str | None = None
    font: str | None = None
    custom_font: str | None = None
    text_font: str | None = None
    code_font: str | None = None
    font_size: int | None = None
    code_font_size: int | None = None
    background_image: str | None = None
    motion_enabled: bool | None = None
    motion_intensity: int | None = None
    undo_limit: int | None = None
    agent_check_on_settings_open: bool | None = None
    auto_check_mounts: bool | None = None
    default_analyzer: dict[str, Any] | None = None
    output_defaults: dict[str, Any] | None = None


class RunRequest(BaseModel):
    model_config = ConfigDict(extra="allow")
    workflow_id: str | None = None
    workflow: dict[str, Any] | None = None
    inputs: dict[str, Any] = Field(default_factory=dict)


class AnalyzerRequest(BaseModel):
    model_config = ConfigDict(extra="allow")
    name: str
    config: dict[str, Any] = Field(default_factory=dict)


class ComponentPresetRequest(BaseModel):
    model_config = ConfigDict(extra="allow")
    kind: str
    name: str
    config: dict[str, Any] = Field(default_factory=dict)


class GrantRequest(BaseModel):
    path: str


class SkillPathRequest(BaseModel):
    path: str


class CredentialLinkRequest(BaseModel):
    provider: str
    credential_ref: str
    endpoint: str = ""
    env_name: str = "OPENAI_API_KEY"
    name: str = ""
    api_format: str | None = None
    models: list[Any] = Field(default_factory=list)


class CredentialTestRequest(BaseModel):
    model_config = ConfigDict(extra="ignore")
    provider: str | None = None
    endpoint: str | None = None
    api_key: str | None = None
    credential_id: str | None = None
    api_format: str | None = None


class ApprovalRequest(BaseModel):
    decision: str
    note: str | None = None
    revision_text: str | None = None


class FilterPreviewRequest(BaseModel):
    model_config = ConfigDict(extra="ignore")
    config: dict[str, Any] = Field(default_factory=dict)
    items: list[Any] = Field(default_factory=list)


class LoopControlRequest(BaseModel):
    model_config = ConfigDict(extra="ignore")
    additional_rounds: int = 1
    additional_seconds: float = 300


@asynccontextmanager
async def lifespan(_app: FastAPI):
    init_db()
    if ensure_agent_config_schema is not None:
        ensure_agent_config_schema()
    yield


app = FastAPI(title="kxy", version="0.1.0", lifespan=lifespan)
if agent_config_router is not None:
    app.include_router(agent_config_router)
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://127.0.0.1:5178", "http://localhost:5178"],
    allow_methods=["GET", "POST", "PUT", "DELETE", "OPTIONS"],
    allow_headers=["Content-Type"],
)


@app.middleware("http")
async def local_origin_guard(request: Request, call_next: Callable):
    from urllib.parse import urlsplit
    host = request.url.hostname
    if host not in {"localhost", "127.0.0.1", "::1", "testserver"}:
        return JSONResponse(status_code=403, content={"detail": "Only local host access is permitted"})
    if request.url.path.startswith("/api/") and request.method in {"POST", "PUT", "PATCH", "DELETE"}:
        origin = request.headers.get("origin")
        if origin:
            try:
                parsed = urlsplit(origin)
                allowed = (parsed.scheme in {"http", "https"} and parsed.hostname in {"localhost", "127.0.0.1", "::1"}
                           and parsed.port in {5178, 8710, request.url.port} and not parsed.username and not parsed.password)
            except ValueError:
                allowed = False
            if not allowed:
                return JSONResponse(status_code=403, content={"detail": "Cross-origin state changes are blocked"})
        if request.headers.get("sec-fetch-site") == "cross-site":
            return JSONResponse(status_code=403, content={"detail": "Cross-site state changes are blocked"})
    return await call_next(request)


@app.get("/api/health")
def health() -> dict[str, Any]:
    return {"ok": True, "app": "kxy", "lfx": "1.12.0" if not LFX_IMPORT_ERROR else "unavailable", "lfx_error": LFX_IMPORT_ERROR}


def _setting_storage(value: Any) -> str:
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (dict, list)):
        return json_text(value)
    return str(value)


def _coerce_setting(key: str, value: Any) -> Any:
    """Read legacy string settings without returning stringly-typed V3 values."""

    default = DEFAULT_SETTINGS[key]
    if key in {"default_analyzer", "output_defaults"}:
        if isinstance(value, str):
            try:
                value = json.loads(value)
            except (TypeError, ValueError, json.JSONDecodeError):
                value = None
        if isinstance(value, dict):
            return value
        return json.loads(json.dumps(default, ensure_ascii=False))
    if key in {"font_size", "code_font_size", "motion_intensity", "undo_limit"}:
        try:
            parsed = int(value)
        except (TypeError, ValueError):
            return default
        if key == "font_size" and 12 <= parsed <= 22:
            return parsed
        if key == "code_font_size" and 10 <= parsed <= 22:
            return parsed
        if key == "motion_intensity" and 0 <= parsed <= MOTION_INTENSITY_MAX:
            return parsed
        if key == "undo_limit" and 1 <= parsed <= 50:
            return parsed
        return default
    if key in {"motion_enabled", "agent_check_on_settings_open", "auto_check_mounts"}:
        if isinstance(value, bool):
            return value
        if str(value).strip().lower() in {"true", "1", "yes", "on"}:
            return True
        if str(value).strip().lower() in {"false", "0", "no", "off"}:
            return False
        return default
    if isinstance(default, str):
        text = str(value) if value is not None else ""
        if key in {"accent", "canvas"} and not HEX_COLOR_RE.fullmatch(text):
            return default
        if key == "palette" and text not in ALLOWED_PALETTES:
            return default
        if key == "font" and text not in ALLOWED_FONTS:
            return default
        if key in {"custom_font", "text_font", "code_font"}:
            if not text or FONT_FAMILY_RE.fullmatch(text.strip()):
                return text.strip()
            return default
        if key == "background_image":
            return text
        return text
    return value


def _default_analyzer_fallback() -> dict[str, Any]:
    return {"mode": "cli", "agent_id": "codex", "prompt": DEFAULT_ANALYZER_PROMPT}


def _normalize_default_analyzer(value: Any) -> dict[str, Any]:
    """Validate and persist only an Agent/model reference, never credentials."""

    try:
        from . import agent_config
    except ImportError as exc:  # pragma: no cover - only legacy import mode
        raise HTTPException(status_code=503, detail="Agent 配置模块不可用") from exc
    if value in (None, "", {}):
        return _default_analyzer_fallback()
    if not isinstance(value, dict):
        raise HTTPException(status_code=422, detail="default_analyzer 必须是对象")
    prompt = value.get("prompt", DEFAULT_ANALYZER_PROMPT)
    if not isinstance(prompt, str):
        raise HTTPException(status_code=422, detail="default_analyzer.prompt 必须是字符串")
    if len(prompt) > MAX_DEFAULT_ANALYZER_PROMPT_CHARS:
        raise HTTPException(status_code=422, detail="default_analyzer.prompt 过长")
    mode = str(value.get("mode") or "cli").strip().lower()
    if mode in {"agent", "plain", "default"}:
        mode = "cli"
    if mode not in {"cli", "model"}:
        raise HTTPException(status_code=422, detail="default_analyzer.mode 必须是 cli 或 model")
    agent_id = str(value.get("agent_id") or value.get("cli") or "").strip().lower()
    if agent_id not in agent_config.AGENT_PRESETS:
        raise HTTPException(status_code=422, detail="default_analyzer.agent_id 不是已知 Agent")
    if mode == "cli":
        # A plain CLI preference remains useful while a local installation is
        # temporarily unavailable; the next node/run will show the precise
        # missing CLI state instead of silently selecting another Agent.
        return {"mode": "cli", "agent_id": agent_id, "prompt": prompt}
    model_ref = str(value.get("model_ref") or "").strip()
    if not model_ref:
        raise HTTPException(status_code=422, detail="model 默认必须提供 model_ref")
    effort = str(value.get("effort") or "").strip()
    try:
        binding = agent_config.resolve_model_binding(
            {"agent_id": agent_id, "model_ref": model_ref, "effort": effort}
        )
    except (ValueError, RuntimeError) as exc:
        raise HTTPException(status_code=422, detail=redact(str(exc), 512)) from exc
    return {
        "mode": "model",
        "agent_id": agent_id,
        "model_ref": str(binding["model_ref"]),
        "effort": str(binding.get("selected_effort") or effort),
        "prompt": prompt,
    }


def _public_default_analyzer(value: Any) -> dict[str, Any]:
    """Decorate stored references with current, secret-free availability."""

    raw = value if isinstance(value, dict) else _default_analyzer_fallback()
    mode = str(raw.get("mode") or "cli").strip().lower()
    agent_id = str(raw.get("agent_id") or raw.get("cli") or "codex").strip().lower()
    prompt = raw.get("prompt", DEFAULT_ANALYZER_PROMPT)
    if not isinstance(prompt, str):
        prompt = DEFAULT_ANALYZER_PROMPT
    result: dict[str, Any] = {"mode": mode, "agent_id": agent_id, "prompt": prompt}
    try:
        from . import agent_config

        if agent_id not in agent_config.AGENT_PRESETS:
            raise ValueError("默认 Agent 不再存在")
        if mode == "model":
            model_ref = str(raw.get("model_ref") or "").strip()
            effort = str(raw.get("effort") or "").strip()
            result.update({"model_ref": model_ref, "effort": effort})
            try:
                binding = agent_config.resolve_model_binding(
                    {"agent_id": agent_id, "model_ref": model_ref, "effort": effort}
                )
            except (ValueError, RuntimeError) as exc:
                result.update(
                    {
                        "status": "missing",
                        "reason": redact(str(exc), 512),
                    }
                )
                return result
            result.update(
                {
                    "model": binding.get("model"),
                    "alias": binding.get("alias"),
                    "source": binding.get("source"),
                    "effort": binding.get("selected_effort") or effort,
                    "status": "available",
                    "reason": "固定模型绑定可用；新建分析器将继承此 Agent、模型和 effort。",
                }
            )
            return result
        profile = agent_config.get_agent_profile(agent_id)
        available = bool(profile.get("available")) and bool(profile.get("resolved_executable"))
        result.update(
            {
                "mode": "cli",
                "status": "available" if available else "missing",
                "reason": (
                    "使用该 Agent 的默认模型；新建分析器不会固定模型。"
                    if available
                    else str(profile.get("message") or "该 Agent CLI 当前不可用；运行时会保留明确失败状态。")
                ),
            }
        )
    except Exception as exc:
        result.update({"status": "missing", "reason": redact(str(exc), 512)})
    return result


def _normalize_output_defaults(value: Any) -> dict[str, Any]:
    """Validate settings for only newly created output containers."""

    default = DEFAULT_SETTINGS["output_defaults"]
    if value in (None, "", {}):
        value = default
    if not isinstance(value, dict):
        raise HTTPException(status_code=422, detail="output_defaults 必须是对象")
    try:
        export_formats = normalize_container_formats(value.get("export_formats", default["export_formats"]))
        allowed = normalize_allowed_file_extensions(
            value.get("allowed_file_extensions", default["allowed_file_extensions"])
        )
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    # Settings always use an explicit allowlist.  None remains reserved for
    # legacy imported containers whose old behavior must be preserved.
    if allowed is None:
        allowed = []
    json_mode = value.get("json_mode", default["json_mode"])
    if json_mode not in JSON_MODES:
        raise HTTPException(status_code=422, detail="output_defaults.json_mode 必须是 content 或 full")
    return {
        "export_formats": export_formats,
        "allowed_file_extensions": allowed,
        "json_mode": json_mode,
    }


def _read_settings() -> dict[str, Any]:
    # Do not call init_db here: it intentionally marks unfinished runs as
    # interrupted and is reserved for application startup/restart recovery.
    ensure_roots()
    with connect_db() as db:
        rows = db.execute("SELECT key, value FROM settings").fetchall()
    result = dict(DEFAULT_SETTINGS)
    for row in rows:
        key = str(row["key"])
        if key in result:
            result[key] = _coerce_setting(key, row["value"])
    result["default_analyzer"] = _public_default_analyzer(result["default_analyzer"])
    result["output_defaults"] = _normalize_output_defaults(result["output_defaults"])
    # A removed background file must not become a dangling URL.  Returning an
    # empty reference is a migration-safe fallback; PUT still rejects any new
    # unknown reference.
    background_id = result.get("background_image", "")
    if background_id:
        with connect_db() as db:
            asset = db.execute("SELECT path FROM appearance_assets WHERE id=?", (background_id,)).fetchone()
        if asset is None or not _managed_background_path(asset["path"]):
            result["background_image"] = ""
    return result


@app.get("/api/settings")
def get_settings() -> dict[str, Any]:
    return _read_settings()


def _validate_font_family(value: Any, field: str) -> str:
    if not isinstance(value, str):
        raise HTTPException(status_code=422, detail=f"{field} must be a local font family name")
    value = value.strip()
    if not value:
        return ""
    if not FONT_FAMILY_RE.fullmatch(value):
        raise HTTPException(status_code=422, detail=f"{field} contains unsupported CSS characters")
    return value


def _managed_background_path(path_value: str | Path) -> bool:
    try:
        path = Path(path_value)
        return path.is_file() and resolved_inside(path, BACKGROUND_ROOT)
    except (OSError, ValueError):
        return False


def _validate_background_reference(value: Any) -> str:
    if not isinstance(value, str):
        raise HTTPException(status_code=422, detail="background_image must be an app-managed image id")
    value = value.strip()
    if not value:
        return ""
    if not re.fullmatch(r"[a-f0-9]{32}", value):
        raise HTTPException(status_code=422, detail="background_image must be an app-managed image id")
    with connect_db() as db:
        row = db.execute("SELECT path FROM appearance_assets WHERE id=?", (value,)).fetchone()
    if row is None or not _managed_background_path(row["path"]):
        raise HTTPException(status_code=422, detail="background_image asset does not exist")
    return value


@app.put("/api/settings")
def put_settings(payload: SettingsRequest) -> dict[str, Any]:
    current = _read_settings()
    incoming = payload.model_dump(exclude_unset=True)
    data = {key: incoming.get(key, current[key]) for key in DEFAULT_SETTINGS}
    # _read_settings decorates the default analyzer with current availability
    # metadata.  When an older/missing binding is merely being preserved while
    # another setting changes, keep the stored reference verbatim instead of
    # rejecting an otherwise unrelated appearance update.
    if "default_analyzer" not in incoming:
        with connect_db() as db:
            stored = db.execute("SELECT value FROM settings WHERE key=?", ("default_analyzer",)).fetchone()
        data["default_analyzer"] = (
            _coerce_setting("default_analyzer", stored["value"])
            if stored is not None
            else _default_analyzer_fallback()
        )
    if not isinstance(data["palette"], str) or data["palette"] not in ALLOWED_PALETTES:
        raise HTTPException(status_code=422, detail="Unsupported palette or font")
    if not isinstance(data["font"], str) or data["font"] not in ALLOWED_FONTS:
        raise HTTPException(status_code=422, detail="Unsupported palette or font")
    if not isinstance(data["accent"], str) or not HEX_COLOR_RE.fullmatch(data["accent"]):
        raise HTTPException(status_code=422, detail="Colors must be six-digit hex values")
    if not isinstance(data["canvas"], str) or not HEX_COLOR_RE.fullmatch(data["canvas"]):
        raise HTTPException(status_code=422, detail="Colors must be six-digit hex values")
    data["custom_font"] = _validate_font_family(data["custom_font"], "custom_font")
    data["text_font"] = _validate_font_family(data["text_font"], "text_font")
    data["code_font"] = _validate_font_family(data["code_font"], "code_font")
    for key, low, high in (("font_size", 12, 22), ("code_font_size", 10, 22), ("motion_intensity", 0, MOTION_INTENSITY_MAX), ("undo_limit", 1, 50)):
        if isinstance(data[key], bool) or not isinstance(data[key], int) or not low <= data[key] <= high:
            raise HTTPException(status_code=422, detail=f"{key} is outside its supported range")
    if not isinstance(data["motion_enabled"], bool):
        raise HTTPException(status_code=422, detail="motion_enabled must be boolean")
    if not isinstance(data["agent_check_on_settings_open"], bool):
        raise HTTPException(status_code=422, detail="agent_check_on_settings_open must be boolean")
    if not isinstance(data["auto_check_mounts"], bool):
        raise HTTPException(status_code=422, detail="auto_check_mounts must be boolean")
    if "default_analyzer" in incoming:
        requested_default = data["default_analyzer"]
        # Keep the current prompt when an older client changes only the Agent
        # or model binding.  A missing prompt in an old database still gets
        # the original built-in prompt through the normalizer fallback.
        if isinstance(requested_default, dict) and "prompt" not in requested_default:
            with connect_db() as db:
                stored = db.execute("SELECT value FROM settings WHERE key=?", ("default_analyzer",)).fetchone()
            stored_default = _coerce_setting("default_analyzer", stored["value"]) if stored is not None else {}
            requested_default = {
                **requested_default,
                "prompt": stored_default.get("prompt", DEFAULT_ANALYZER_PROMPT)
                if isinstance(stored_default, dict)
                else DEFAULT_ANALYZER_PROMPT,
            }
        data["default_analyzer"] = _normalize_default_analyzer(requested_default)
    elif not isinstance(data["default_analyzer"], dict):
        data["default_analyzer"] = _default_analyzer_fallback()
    data["output_defaults"] = _normalize_output_defaults(data["output_defaults"])
    data["background_image"] = _validate_background_reference(data["background_image"])
    with connect_db() as db:
        for key, value in data.items():
            db.execute(
                "INSERT INTO settings(key, value) VALUES (?, ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
                (key, _setting_storage(value)),
            )
    return _read_settings()


def _font_entry(family: str) -> dict[str, Any] | None:
    family = " ".join(str(family).strip().split())
    if not family or not FONT_FAMILY_RE.fullmatch(family):
        return None
    lower = family.casefold()
    monospace = any(token in lower for token in ("mono", "menlo", "monaco", "courier", "consolas", "code", "terminal"))
    return {"family": family, "monospace": monospace}


def _font_families_from_system_profiler(value: Any) -> set[str]:
    found: set[str] = set()

    def walk(item: Any) -> None:
        if isinstance(item, dict):
            for key, child in item.items():
                if str(key).casefold() in {"family", "family_name", "familyname", "font_family", "fontfamily"}:
                    if isinstance(child, str):
                        found.add(child)
                    elif isinstance(child, list):
                        found.update(str(entry) for entry in child if isinstance(entry, str))
                walk(child)
        elif isinstance(item, list):
            for child in item:
                walk(child)

    walk(value)
    return found


def _enumerate_font_families() -> tuple[list[dict[str, Any]], str, str | None]:
    global _FONT_CACHE
    now = time.monotonic()
    if _FONT_CACHE is not None and now - _FONT_CACHE[0] < FONT_CACHE_SECONDS:
        return _FONT_CACHE[1], _FONT_CACHE[2], _FONT_CACHE[3]

    families: set[str] = set()
    source = "unavailable"
    error: str | None = None
    fc_list = shutil.which("fc-list")
    if fc_list:
        try:
            completed = subprocess.run(
                [fc_list, ":", "family"],
                capture_output=True,
                text=True,
                timeout=4,
                check=False,
                env={"PATH": os.environ.get("PATH", os.defpath), "LANG": os.environ.get("LANG", "C")},
            )
            if completed.returncode == 0:
                for line in (completed.stdout or "")[:500_000].splitlines():
                    families.update(part.strip() for part in line.split(",") if part.strip())
                if families:
                    source = "fontconfig"
            else:
                error = (completed.stderr or "fontconfig query failed").strip()[:240]
        except (OSError, subprocess.TimeoutExpired) as exc:
            error = f"fontconfig: {type(exc).__name__}"
    if not families and sys.platform == "darwin":
        try:
            completed = subprocess.run(
                ["/usr/sbin/system_profiler", "SPFontsDataType", "-json"],
                capture_output=True,
                text=True,
                timeout=4,
                check=False,
                env={"PATH": os.environ.get("PATH", os.defpath), "LANG": os.environ.get("LANG", "C")},
            )
            if completed.returncode == 0:
                families = _font_families_from_system_profiler(json.loads((completed.stdout or "{}")[:2_000_000]))
                if families:
                    source = "macos-system-profiler"
            if not families and not error:
                error = (completed.stderr or "macOS font query returned no families").strip()[:240]
        except (OSError, ValueError, json.JSONDecodeError, subprocess.TimeoutExpired) as exc:
            error = f"macOS font query: {type(exc).__name__}"
    entries = sorted((item for family in families if (item := _font_entry(family))), key=lambda item: item["family"].casefold())
    if not entries and error is None:
        error = "No supported installed-font query is available"
    _FONT_CACHE = (now, entries, source, error)
    return entries, source, error


@app.get("/api/fonts")
def list_fonts() -> dict[str, Any]:
    fonts, source, error = _enumerate_font_families()
    result: dict[str, Any] = {"fonts": fonts, "source": source}
    if error:
        result["error"] = error
    return result


def _validate_background_content(content: bytes) -> tuple[str, str, int, int]:
    if len(content) > MAX_BACKGROUND_BYTES:
        raise HTTPException(status_code=413, detail="Background image exceeds 10 MB")
    try:
        from PIL import Image

        with Image.open(io.BytesIO(content)) as image:
            image_format = str(image.format or "").upper()
            width, height = image.size
            image.verify()
        if image_format not in {"PNG", "JPEG", "WEBP"}:
            raise HTTPException(status_code=415, detail="Only PNG, JPEG and WebP backgrounds are supported")
        if width <= 0 or height <= 0 or width > 12_000 or height > 12_000 or width * height > MAX_BACKGROUND_PIXELS:
            raise HTTPException(status_code=422, detail="Background image dimensions are unreasonable")
        # Reopen after verify so a truncated/invalid decoder cannot pass only
        # because metadata was readable.
        with Image.open(io.BytesIO(content)) as image:
            image.load()
    except HTTPException:
        raise
    except _PILDecompressionBombError as exc:
        raise HTTPException(status_code=422, detail="Background image dimensions are unreasonable") from exc
    except (ImportError, OSError, ValueError, SyntaxError) as exc:
        raise HTTPException(status_code=422, detail="Background image is not a valid raster image") from exc
    mime = {"PNG": "image/png", "JPEG": "image/jpeg", "WEBP": "image/webp"}[image_format]
    suffix = {"PNG": ".png", "JPEG": ".jpg", "WEBP": ".webp"}[image_format]
    return mime, suffix, width, height


@app.post("/api/appearance/background")
async def upload_background(file: UploadFile = File(...)) -> dict[str, str]:
    content = await file.read(MAX_BACKGROUND_BYTES + 1)
    mime, suffix, width, height = _validate_background_content(content)
    asset_id = uuid.uuid4().hex
    destination = BACKGROUND_ROOT / f"{asset_id}{suffix}"
    temporary = BACKGROUND_ROOT / f".{asset_id}.{uuid.uuid4().hex}.tmp"
    try:
        temporary.write_bytes(content)
        temporary.replace(destination)
        digest = hashlib.sha256(content).hexdigest()
        with connect_db() as db:
            db.execute(
                "INSERT INTO appearance_assets(id,path,mime,sha256,size,width,height,created_at) VALUES (?,?,?,?,?,?,?,?)",
                (asset_id, str(destination), mime, digest, len(content), width, height, utc_now()),
            )
    except OSError as exc:
        temporary.unlink(missing_ok=True)
        destination.unlink(missing_ok=True)
        raise HTTPException(status_code=500, detail="Unable to store background image") from exc
    except Exception:
        temporary.unlink(missing_ok=True)
        destination.unlink(missing_ok=True)
        raise
    return {"id": asset_id, "url": f"/api/appearance/background/{asset_id}"}


@app.get("/api/appearance/background/{asset_id}")
def get_background(asset_id: str) -> FileResponse:
    if not re.fullmatch(r"[a-f0-9]{32}", asset_id):
        raise HTTPException(status_code=404, detail="Background image not found")
    with connect_db() as db:
        row = db.execute("SELECT path,mime,sha256 FROM appearance_assets WHERE id=?", (asset_id,)).fetchone()
    if row is None or not _managed_background_path(row["path"]):
        raise HTTPException(status_code=404, detail="Background image not found")
    path = Path(row["path"])
    try:
        if hashlib.sha256(path.read_bytes()).hexdigest() != row["sha256"]:
            raise HTTPException(status_code=409, detail="Managed background image changed")
    except OSError as exc:
        raise HTTPException(status_code=404, detail="Background image not found") from exc
    return FileResponse(path, media_type=row["mime"], filename=path.name)


@app.get("/api/branding/logo")
def get_branding_logo() -> FileResponse:
    for filename, media_type in BRANDING_LOGO_FILES:
        path = BRANDING_ROOT / filename
        if path.is_symlink() or not path.is_file() or not resolved_inside(path, BRANDING_ROOT):
            continue
        return FileResponse(path, media_type=media_type, headers={"Cache-Control": "no-store"})
    raise HTTPException(status_code=404, detail="Branding logo not found")


@app.get("/api/templates")
def templates() -> list[dict[str, Any]]:
    return [template_workflow("extraction"), template_workflow("expansion"), template_workflow("cross-check")]


def template_workflow(kind: str) -> dict[str, Any]:
    base = {"version": "kxy.workflow.v1", "settings": {}, "nodes": [], "edges": []}
    if kind == "extraction":
        base.update(
            {
                "id": "template-extraction",
                "name": "资料提取",
                "nodes": [
                    {"id": "source", "type": "text", "position": {"x": 100, "y": 220}, "data": {"label": "研究资料", "text": "把资料拖到画布，或替换为文件节点。"}},
                    {"id": "extract", "type": "analyzer", "position": {"x": 440, "y": 220}, "data": {"label": "提取事实", "prompt": "从输入资料中提取可核查事实，保留页码、段落或表格行号。输出 JSON 或清晰的事实列表。", "cli": "codex", "network": True, "timeout": 600}},
                    {"id": "output", "type": "container", "position": {"x": 800, "y": 220}, "data": {"label": "授权输出"}},
                ],
                "edges": [{"id": "e-source-extract", "source": "source", "target": "extract", "sourceHandle": "result", "targetHandle": "items"}, {"id": "e-extract-output", "source": "extract", "target": "output", "sourceHandle": "result", "targetHandle": "items"}],
            }
        )
    elif kind == "expansion":
        base.update(
            {
                "id": "template-expansion",
                "name": "扩展研究",
                "nodes": [
                    {"id": "question", "type": "text", "position": {"x": 80, "y": 280}, "data": {"label": "研究问题", "text": "哪些未知最值得继续花时间验证？"}},
                    {"id": "research", "type": "analyzer", "position": {"x": 360, "y": 280}, "data": {"label": "扩展检索", "prompt": "围绕研究问题拆解信息缺口，给出可验证的下一步公开资料线索。必须只返回一个 JSON 对象，形如 {\"route\":\"close\" 或 \"more\",\"evidence\":[{\"source\":\"\",\"locator\":\"\",\"fact\":\"\"}],\"unknowns\":[]}。只有当输入资料已经足以形成有来源的回答时 route 才能为 close；否则必须为 more。", "cli": "codex", "network": True, "expect_json": True, "timeout": 900}},
                    {"id": "condition", "type": "condition", "position": {"x": 690, "y": 280}, "data": {"label": "是否足够", "field": "structured.route", "operator": "equals", "expected": "close"}},
                    {"id": "close", "type": "analyzer", "position": {"x": 980, "y": 170}, "data": {"label": "形成结论", "prompt": "基于上游结果整理一条带边界条件的研究结论。", "cli": "codex", "network": True, "timeout": 600}},
                    {"id": "more", "type": "analyzer", "position": {"x": 980, "y": 390}, "data": {"label": "列出缺口", "prompt": "如果证据不足，列出下一轮人工或公开资料验证问题。", "cli": "codex", "network": True, "timeout": 600}},
                    {"id": "output", "type": "container", "position": {"x": 1280, "y": 280}, "data": {"label": "研究输出"}},
                ],
                "edges": [
                    {"id": "e-q-r", "source": "question", "target": "research", "sourceHandle": "result", "targetHandle": "items"},
                    {"id": "e-r-c", "source": "research", "target": "condition", "sourceHandle": "result", "targetHandle": "items"},
                    {"id": "e-c-close", "source": "condition", "target": "close", "sourceHandle": "true", "targetHandle": "items"},
                    {"id": "e-c-more", "source": "condition", "target": "more", "sourceHandle": "false", "targetHandle": "items"},
                    {"id": "e-close-out", "source": "close", "target": "output", "sourceHandle": "result", "targetHandle": "items"},
                    {"id": "e-more-out", "source": "more", "target": "output", "sourceHandle": "result", "targetHandle": "items"},
                ],
            }
        )
    else:
        base.update(
            {
                "id": "template-cross-check",
                "name": "交叉核验",
                "nodes": [
                    {"id": "source-a", "type": "text", "position": {"x": 80, "y": 160}, "data": {"label": "来源 A", "text": "把第一份资料拖到这里。"}},
                    {"id": "source-b", "type": "text", "position": {"x": 80, "y": 420}, "data": {"label": "来源 B", "text": "把第二份资料拖到这里。"}},
                    {"id": "check", "type": "analyzer", "position": {"x": 460, "y": 280}, "data": {"label": "交叉核验", "prompt": "对比所有输入，指出一致、冲突、时间差异和无法判断的部分。不要把推断写成事实。", "cli": "codex", "network": True, "timeout": 600}},
                    {"id": "output", "type": "container", "position": {"x": 840, "y": 280}, "data": {"label": "核验记录"}},
                ],
                "edges": [{"id": "e-a-check", "source": "source-a", "target": "check", "sourceHandle": "result", "targetHandle": "items"}, {"id": "e-b-check", "source": "source-b", "target": "check", "sourceHandle": "result", "targetHandle": "items"}, {"id": "e-check-out", "source": "check", "target": "output", "sourceHandle": "result", "targetHandle": "items"}],
            }
        )
    return base


@app.get("/api/workflows")
def list_workflows() -> list[dict[str, Any]]:
    with connect_db() as db:
        rows = db.execute("SELECT id, name, version, document, created_at, updated_at FROM workflows ORDER BY updated_at DESC").fetchall()
    return [{**dict(row), "workflow": json.loads(row["document"])} for row in rows]


@app.post("/api/workflows/portable/export")
def export_portable_workflow(payload: PortableExportRequest) -> Response:
    workflow = payload.workflow
    if workflow is None and payload.workflow_id:
        with connect_db() as db:
            row = db.execute("SELECT document FROM workflows WHERE id=?", (payload.workflow_id,)).fetchone()
        if row is None:
            raise HTTPException(status_code=404, detail="Workflow not found")
        workflow = json.loads(row["document"])
    if not isinstance(workflow, dict):
        raise HTTPException(status_code=422, detail="Provide workflow or workflow_id")
    try:
        normalized = normalize_workflow(workflow)
        validation = validate_workflow(normalized)
        if not validation.get("ok"):
            raise PortablePackageError("workflow failed kxy validation: " + "; ".join(str(item) for item in validation.get("errors", []))[:1000])
        content = build_package(
            normalized,
            include_file_ids=payload.include_file_ids,
            include_skill_ids=payload.include_skill_ids,
        )
    except PortablePackageError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    return Response(
        content=content,
        media_type="application/zip",
        headers={"Content-Disposition": 'attachment; filename="kxy-workflow-portable.zip"'},
    )


@app.post("/api/workflows/portable/import/preview")
async def preview_portable_workflow(file: UploadFile = File(...)) -> dict[str, Any]:
    content = await file.read(PORTABLE_MAX_ARCHIVE_BYTES + 1)
    try:
        return preview_package(content)
    except PortablePackageError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    except zipfile.BadZipFile as exc:
        raise HTTPException(status_code=422, detail="Invalid portable ZIP") from exc


@app.post("/api/workflows/portable/import/commit")
def commit_portable_workflow(payload: PortableImportCommitRequest) -> dict[str, Any]:
    try:
        return commit_package(payload.preview_id, payload.bindings)
    except PortablePreviewExpired as exc:
        raise HTTPException(status_code=410, detail=str(exc)) from exc
    except PortablePackageError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@app.get("/api/workflows/{workflow_id}")
def get_workflow(workflow_id: str) -> dict[str, Any]:
    with connect_db() as db:
        row = db.execute("SELECT * FROM workflows WHERE id=?", (workflow_id,)).fetchone()
    if row is None:
        raise HTTPException(status_code=404, detail="Workflow not found")
    return {**dict(row), "workflow": json.loads(row["document"])}


@app.post("/api/workflows/validate")
def validate_workflow_api(workflow: dict[str, Any]) -> dict[str, Any]:
    return validate_workflow(workflow)


@app.post("/api/workflows")
@app.put("/api/workflows/{workflow_id}")
def save_workflow(payload: WorkflowSaveRequest, workflow_id: str | None = None) -> dict[str, Any]:
    workflow = normalize_workflow(payload.workflow)
    validation = validate_workflow(workflow)
    if not validation["ok"]:
        raise HTTPException(status_code=422, detail=validation)
    identifier = workflow_id or payload.id or workflow.get("id") or uuid.uuid4().hex
    workflow["id"] = identifier
    workflow["name"] = payload.name or workflow.get("name") or "未命名研究流程"
    timestamp = utc_now()
    with connect_db() as db:
        old = db.execute("SELECT created_at FROM workflows WHERE id=?", (identifier,)).fetchone()
        db.execute(
            "INSERT INTO workflows(id, name, version, document, created_at, updated_at) VALUES (?, ?, ?, ?, ?, ?) "
            "ON CONFLICT(id) DO UPDATE SET name=excluded.name, version=excluded.version, document=excluded.document, updated_at=excluded.updated_at",
            (identifier, workflow["name"], workflow["version"], json_text(workflow), old["created_at"] if old else timestamp, timestamp),
        )
    return {"id": identifier, "name": workflow["name"], "workflow": workflow, "validation": validation}


@app.delete("/api/workflows/{workflow_id}")
def delete_workflow(workflow_id: str) -> dict[str, bool]:
    with connect_db() as db:
        db.execute("DELETE FROM workflows WHERE id=?", (workflow_id,))
    return {"ok": True}


@app.post("/api/files")
async def upload_file(file: UploadFile = File(...)) -> dict[str, Any]:
    content = await file.read(MAX_FILE_BYTES + 1)
    return store_file(file.filename or "untitled", content)


@app.get("/api/files")
def list_files() -> list[dict[str, Any]]:
    with connect_db() as db:
        rows = db.execute("SELECT * FROM files ORDER BY created_at DESC").fetchall()
    return [file_row(row) for row in rows]


@app.get("/api/files/{file_id}")
def get_file(file_id: str) -> dict[str, Any]:
    with connect_db() as db:
        row = db.execute("SELECT * FROM files WHERE id=?", (file_id,)).fetchone()
    if row is None:
        raise HTTPException(status_code=404, detail="File not found")
    return file_row(row)


@app.get("/api/files/{file_id}/content")
def file_content(file_id: str) -> FileResponse:
    with connect_db() as db:
        row = db.execute("SELECT * FROM files WHERE id=?", (file_id,)).fetchone()
    if row is None or not Path(row["stored_path"]).is_file():
        raise HTTPException(status_code=404, detail="File not found")
    return FileResponse(row["stored_path"], media_type=row["mime"], filename=row["display_name"])


@app.get("/api/skills")
def list_skills() -> list[dict[str, Any]]:
    with connect_db() as db:
        rows = db.execute("SELECT * FROM skills ORDER BY created_at DESC").fetchall()
    return [skill_row(row) for row in rows]


@app.post("/api/skills/import-path")
def import_skill_path(payload: SkillPathRequest) -> dict[str, Any]:
    try:
        return import_skill_root(Path(payload.path).expanduser())
    except (OSError, ValueError) as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@app.post("/api/skills/import-zip")
async def import_skill_archive(file: UploadFile = File(...)) -> dict[str, Any]:
    try:
        return import_skill_zip(await file.read(MAX_SKILL_BYTES + 1), file.filename or "skill.zip")
    except (OSError, ValueError, zipfile.BadZipFile) as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


def _skill_delete_stage(skill_id: str) -> Path:
    digest = hashlib.sha256(str(skill_id).encode("utf-8")).hexdigest()[:24]
    stage_root = SKILLS_ROOT / ".deleting"
    stage_root.mkdir(parents=True, exist_ok=True)
    stage = stage_root / f"{digest}-{uuid.uuid4().hex}"
    if not resolved_inside(stage, SKILLS_ROOT):
        raise RuntimeError("Skill 删除暂存路径越过 KXY skills 边界")
    return stage


def _delete_skill_record(skill_id: str) -> dict[str, Any]:
    """Delete one local record with a reversible snapshot move before DB commit."""

    skill_id = str(skill_id or "").strip()
    if not skill_id:
        return {"id": skill_id, "ok": False, "reason": "Skill id 不能为空"}
    with connect_db() as db:
        row = db.execute("SELECT * FROM skills WHERE id=?", (skill_id,)).fetchone()
    if row is None:
        return {"id": skill_id, "ok": False, "not_found": True, "reason": "Skill not found"}

    row_data = dict(row)
    root_value = Path(str(row_data.get("root_path") or "")).expanduser()
    root_lexical = Path(os.path.abspath(os.fspath(root_value)))
    try:
        root_resolved = root_lexical.resolve(strict=False)
    except (OSError, RuntimeError) as exc:
        return {"id": skill_id, "ok": False, "reason": f"Skill 快照路径无法解析：{type(exc).__name__}"}
    skills_root_lexical = SKILLS_ROOT.absolute()
    lexical_inside = False
    try:
        relative_parts = root_lexical.relative_to(skills_root_lexical).parts
        lexical_inside = bool(relative_parts)
    except ValueError:
        relative_parts = ()
    if lexical_inside:
        current = skills_root_lexical
        if any((current := current / part).is_symlink() for part in relative_parts):
            return {"id": skill_id, "ok": False, "reason": "Skill 快照路径包含软链接，拒绝删除"}
    owned = lexical_inside and root_resolved != SKILLS_ROOT.resolve() and resolved_inside(root_resolved, SKILLS_ROOT)
    root = root_lexical if owned else root_resolved
    stage: Path | None = None
    marker_created = False
    db_deleted = False
    try:
        from . import agent_config

        if owned and root.exists():
            if root.is_symlink() or not root.is_dir():
                raise RuntimeError("KXY Skill 快照不是安全的普通目录")
            with connect_db() as db:
                other_rows = db.execute("SELECT id, root_path FROM skills WHERE id != ?", (skill_id,)).fetchall()
            for other in other_rows:
                try:
                    other_root = Path(str(other["root_path"] or "")).expanduser().resolve(strict=False)
                except (OSError, RuntimeError):
                    continue
                if other_root == root.resolve(strict=False):
                    raise RuntimeError(f"Skill 快照被其他 KXY Skill 引用：{other['id']}")
            stage = _skill_delete_stage(skill_id)
        agent_config.skillhub_prepare_skill_delete(row_data, staged_path=stage)
        marker_created = True
        if stage is not None:
            root.rename(stage)
            agent_config.skillhub_update_skill_delete(
                skill_id,
                "pending",
                staged_relative=str(stage.resolve().relative_to(DATA_ROOT.resolve())),
            )
        agent_config.skillhub_update_skill_delete(skill_id, "committing")
        with connect_db() as db:
            cursor = db.execute("DELETE FROM skills WHERE id=?", (skill_id,))
            if cursor.rowcount != 1:
                raise RuntimeError("Skill DB 记录未删除")
        db_deleted = True
    except Exception as exc:
        restore_error = ""
        if stage is not None and stage.exists() and not root.exists():
            try:
                stage.rename(root)
            except (OSError, RuntimeError) as restore_exc:
                restore_error = f"；快照恢复失败：{type(restore_exc).__name__}"
        if marker_created and not db_deleted:
            try:
                agent_config.skillhub_cancel_skill_delete(skill_id)
            except Exception as marker_exc:
                restore_error += f"；删除标记恢复失败：{type(marker_exc).__name__}"
        return {
            "id": skill_id,
            "ok": False,
            "deleted": db_deleted,
            "reason": f"{type(exc).__name__}: {exc}{restore_error}"[:512],
        }

    try:
        agent_config.skillhub_update_skill_delete(
            skill_id,
            "deleted",
            root_deleted=bool(stage is not None),
            source_preserved=not owned,
            cleanup_pending=bool(stage is not None),
        )
    except Exception as exc:
        return {
            "id": skill_id,
            "ok": False,
            "deleted": True,
            "reason": f"数据库记录已删除，但 tombstone 未完成：{type(exc).__name__}"[:512],
        }

    if stage is not None:
        try:
            shutil.rmtree(stage)
            agent_config.skillhub_update_skill_delete(skill_id, "deleted", cleanup_pending=False)
        except (OSError, RuntimeError, ValueError) as exc:
            return {
                "id": skill_id,
                "ok": False,
                "deleted": True,
                "reason": f"数据库记录已删除，快照暂存清理失败：{type(exc).__name__}"[:512],
            }
    return {
        "id": skill_id,
        "ok": True,
        "deleted": True,
        "root_deleted": bool(stage is not None),
        "source_preserved": not owned,
    }


@app.delete("/api/skills/{skill_id}")
def delete_skill(skill_id: str) -> dict[str, Any]:
    result = _delete_skill_record(skill_id)
    if result.get("not_found"):
        raise HTTPException(status_code=404, detail=result["reason"])
    if not result.get("ok"):
        raise HTTPException(status_code=409, detail=result["reason"])
    return {"ok": True, **result}


@app.post("/api/skills/bulk-delete")
def bulk_delete_skills(payload: dict[str, Any]) -> dict[str, Any]:
    if not isinstance(payload, dict) or not isinstance(payload.get("skill_ids"), list):
        raise HTTPException(status_code=422, detail="skill_ids 必须是数组")
    raw_ids = payload["skill_ids"]
    if len(raw_ids) > 100:
        raise HTTPException(status_code=422, detail="一次最多删除 100 个 Skill")
    skill_ids = list(dict.fromkeys(str(value or "").strip() for value in raw_ids))
    results = [_delete_skill_record(skill_id) for skill_id in skill_ids]
    deleted = [item["id"] for item in results if item.get("ok")]
    failed = [
        {"id": item["id"], "reason": item.get("reason", "删除失败"), "deleted": bool(item.get("deleted"))}
        for item in results
        if not item.get("ok")
    ]
    status = "ok" if not failed else "partial" if deleted else "blocked"
    return {"status": status, "deleted": deleted, "failed": failed, "results": results}


@app.get("/api/analyzers")
def list_analyzers() -> list[dict[str, Any]]:
    with connect_db() as db:
        rows = db.execute("SELECT * FROM analyzers ORDER BY updated_at DESC").fetchall()
    return [analyzer_row(row) for row in rows]


@app.post("/api/analyzers")
def save_analyzer(payload: AnalyzerRequest) -> dict[str, Any]:
    clean, removed = sanitize_config(payload.config)
    if removed:
        raise HTTPException(status_code=422, detail={"message": "Credentials cannot be stored in analyzer presets", "fields": removed})
    forbidden_mcp_keys = [key for key in ("mcp", "mcpServers", "mcp_servers", "mcp_config") if key in clean]
    if forbidden_mcp_keys:
        raise HTTPException(status_code=422, detail={"message": "Analyzer stores selected mcp_ids only", "fields": forbidden_mcp_keys})
    if "mcp_ids" in clean:
        values = clean.get("mcp_ids")
        if isinstance(values, str) or not isinstance(values, list):
            raise HTTPException(status_code=422, detail="mcp_ids must be an array of imported server ids")
        normalized_ids = list(dict.fromkeys(str(value or "").strip() for value in values))
        if any(not value for value in normalized_ids):
            raise HTTPException(status_code=422, detail="mcp_ids cannot contain empty ids")
        try:
            from . import agent_config

            target_agent = str(clean.get("agent_id") or clean.get("cli") or "codex").strip().lower()
            supported, reason = agent_config._mcp_support(target_agent)
            if normalized_ids and not supported:
                raise HTTPException(status_code=422, detail=f"Agent {target_agent} does not support selected MCP: {reason}")
            imported = agent_config._mcp_load_index()
            missing = [value for value in normalized_ids if value not in imported]
            if missing:
                raise HTTPException(status_code=422, detail=f"MCP server not imported: {', '.join(missing)}")
            clean["mcp_ids"] = normalized_ids
        except HTTPException:
            raise
        except (AttributeError, ValueError) as exc:
            raise HTTPException(status_code=422, detail=str(exc)) from exc
    # Each POST is a new saved snapshot.  A node config may carry an old
    # ``id`` field, but it must never become the database UPSERT identity.
    identifier = uuid.uuid4().hex
    timestamp = utc_now()
    with connect_db() as db:
        db.execute(
            "INSERT INTO analyzers(id, name, config, created_at, updated_at) VALUES (?, ?, ?, ?, ?) "
            "ON CONFLICT(id) DO UPDATE SET name=excluded.name, config=excluded.config, updated_at=excluded.updated_at",
            (identifier, payload.name[:160], json_text(clean), timestamp, timestamp),
        )
        row = db.execute("SELECT * FROM analyzers WHERE id=?", (identifier,)).fetchone()
    assert row is not None
    return analyzer_row(row)


@app.delete("/api/analyzers/{analyzer_id}")
def delete_analyzer(analyzer_id: str) -> dict[str, bool]:
    with connect_db() as db:
        db.execute("DELETE FROM analyzers WHERE id=?", (analyzer_id,))
    return {"ok": True}


@app.get("/api/presets")
def list_component_presets() -> list[dict[str, Any]]:
    """Return legacy analyzer rows and V8 component rows once each."""

    with connect_db() as db:
        analyzer_rows = db.execute("SELECT * FROM analyzers").fetchall()
        component_rows = db.execute(
            "SELECT * FROM component_presets ORDER BY updated_at DESC"
        ).fetchall()
    result: list[dict[str, Any]] = []
    for row in analyzer_rows:
        item = analyzer_row(row)
        item["kind"] = "analyzer"
        item["source"] = "analyzer"
        result.append(item)
    result.extend(component_preset_row(row) for row in component_rows)
    result.sort(
        key=lambda item: str(item.get("updated_at") or item.get("created_at") or ""),
        reverse=True,
    )
    return result


@app.post("/api/presets")
def save_component_preset(payload: ComponentPresetRequest) -> dict[str, Any]:
    kind = str(payload.kind or "").strip().lower()
    if kind not in PRESET_NODE_TYPES:
        raise HTTPException(
            status_code=422,
            detail="只支持 file、text、input、analyzer、container、condition、human、blackbox 组件预设",
        )
    if kind == "analyzer":
        # Keep the existing analyzer table and API as the single source of
        # truth.  The generic list merely presents that same row with kind and
        # source metadata, so saving here cannot create a duplicate record.
        result = save_analyzer(AnalyzerRequest(name=payload.name, config=payload.config))
        result["kind"] = "analyzer"
        result["source"] = "analyzer"
        return result

    clean, removed = sanitize_component_preset_config(payload.config)
    if removed:
        raise HTTPException(
            status_code=422,
            detail={"message": "Credentials cannot be stored in component presets", "fields": removed},
        )
    if kind == "blackbox":
        nested = clean.get("workflow") if isinstance(clean, dict) else None
        if not isinstance(nested, dict):
            raise HTTPException(status_code=422, detail="Blackbox preset requires data.workflow")
        try:
            # The public validator rejects boundary marker nodes at the root,
            # while a blackbox's nested workflow is required to contain them.
            # Validate through a temporary legal blackbox root so the same
            # depth, reachability and edge rules are applied.
            validation = validate_workflow(
                {
                    "version": "kxy.workflow.v1",
                    "name": "V8 preset validation",
                    "nodes": [
                        {
                            "id": "preset-box",
                            "type": "blackbox",
                            "position": {"x": 0, "y": 0},
                            "data": {"workflow": normalize_workflow(nested)},
                        }
                    ],
                    "edges": [],
                }
            )
        except (TypeError, ValueError) as exc:
            raise HTTPException(status_code=422, detail=str(exc)) from exc
        if not validation.get("ok"):
            raise HTTPException(status_code=422, detail=validation)
    identifier = uuid.uuid4().hex
    timestamp = utc_now()
    with connect_db() as db:
        db.execute(
            "INSERT INTO component_presets(id, kind, name, config, created_at, updated_at) VALUES (?, ?, ?, ?, ?, ?)",
            (identifier, kind, str(payload.name or "未命名预设")[:160], json_text(clean), timestamp, timestamp),
        )
        row = db.execute("SELECT * FROM component_presets WHERE id=?", (identifier,)).fetchone()
    assert row is not None
    return component_preset_row(row)


@app.delete("/api/presets/{preset_id}")
def delete_component_preset(preset_id: str) -> dict[str, bool]:
    with connect_db() as db:
        row = db.execute("SELECT 1 FROM component_presets WHERE id=?", (preset_id,)).fetchone()
        if row is None:
            raise HTTPException(status_code=404, detail="Component preset not found")
        db.execute("DELETE FROM component_presets WHERE id=?", (preset_id,))
    return {"ok": True}


@app.get("/api/grants")
def list_grants() -> list[dict[str, Any]]:
    with connect_db() as db:
        rows = db.execute("SELECT id, canonical_path, revoked_at, created_at FROM grants ORDER BY created_at DESC").fetchall()
    return [dict(row) for row in rows]


@app.post("/api/grants")
def create_grant(payload: GrantRequest) -> dict[str, Any]:
    path = Path(payload.path).expanduser().resolve()
    if not path.is_dir():
        raise HTTPException(status_code=422, detail="Output folder does not exist")
    identifier = uuid.uuid4().hex
    with connect_db() as db:
        db.execute("INSERT INTO grants(id, canonical_path, created_at) VALUES (?, ?, ?)", (identifier, str(path), utc_now()))
    return {"id": identifier, "canonical_path": str(path), "revoked_at": None}


@app.delete("/api/grants/{grant_id}")
def revoke_grant(grant_id: str) -> dict[str, bool]:
    with connect_db() as db:
        db.execute("UPDATE grants SET revoked_at=? WHERE id=?", (utc_now(), grant_id))
    return {"ok": True}


@app.get("/api/cli/status")
def cli_status() -> dict[str, Any]:
    return {"lfx": {"available": not bool(LFX_IMPORT_ERROR), "version": "1.12.0" if not LFX_IMPORT_ERROR else None}, "clis": [probe_cli(name) for name in ("codex", "opencode", "claude", "pi")]}


@app.get("/api/credentials/status")
def credential_status() -> dict[str, Any]:
    with connect_db() as db:
        linked = db.execute("SELECT id, provider, credential_ref, endpoint, env_name, created_at FROM credentials ORDER BY created_at DESC").fetchall()
    try:
        references = [credential_public(row, configured=keychain_has(row["credential_ref"])) for row in linked]
    except RuntimeError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    return {
        "codex": {"mode": "existing-cli-login", "configured": None, "note": "Use `codex login status` outside kxy; no auth files are read."},
        "opencode": {"mode": "existing-cli-login", "configured": None, "note": "Use `opencode auth` outside kxy; no auth files are read."},
        "keychain_references": references,
        "api_key_storage": "Keychain references only; kxy never persists or returns secret values",
    }


@app.get("/api/credentials")
def list_credentials() -> list[dict[str, Any]]:
    with connect_db() as db:
        rows = db.execute("SELECT id, provider, credential_ref, endpoint, env_name, created_at FROM credentials ORDER BY created_at DESC").fetchall()
        try:
            return [credential_public(row, configured=keychain_has(row["credential_ref"]), db=db) for row in rows]
        except RuntimeError as exc:
            raise HTTPException(status_code=503, detail=str(exc)) from exc


def validate_endpoint(endpoint: str) -> str:
    from urllib.parse import urlsplit
    if not endpoint: return ''
    url = urlsplit(endpoint)
    if url.username or url.password or url.query or url.fragment or not url.hostname:
        raise HTTPException(422, 'Endpoint 必须是无凭据、查询参数或片段的 API 基础地址')
    if url.scheme != 'https' and not (url.scheme == 'http' and url.hostname in {'localhost','127.0.0.1','::1'}):
        raise HTTPException(422, '远程 Endpoint 需要 HTTPS，本机模型可用 HTTP')
    return endpoint.rstrip('/')


def _models_endpoint(endpoint: str) -> str:
    from urllib.parse import urlsplit

    base = endpoint.rstrip("/")
    path = urlsplit(base).path.rstrip("/")
    # The default providers already include /v1.  A bare host gets the
    # provider convention; an explicit custom base (/api/v3, /gateway, ...)
    # is respected verbatim and only receives /models.
    return f"{base}/v1/models" if not path else f"{base}/models"


def _credential_name(value: Any, fallback: str) -> str:
    if value in (None, ""):
        value = fallback
    return re.sub(r"[\x00-\x1f\x7f]", " ", str(value)).strip()[:160] or fallback[:160]


def _credential_provider(value: Any) -> str:
    provider = re.sub(r"[\x00-\x1f\x7f]", " ", str(value or "custom")).strip().lower()[:80]
    if provider not in {"openai", "anthropic", "google", "custom"}:
        raise HTTPException(status_code=422, detail="provider 必须是 openai、anthropic、google 或 custom")
    return provider


def _secret_in_public_credential_fields(secret: str, *values: Any) -> bool:
    return bool(secret) and any(secret in json_text(value) for value in values)


@app.post("/api/credentials/test")
async def test_credential(request: Request) -> dict[str, Any]:
    """Probe a provider's model catalog without persisting or logging a secret."""
    try:
        raw = await request.json()
    except (TypeError, ValueError):
        raise HTTPException(status_code=422, detail="请求需要 JSON 对象")
    if not isinstance(raw, dict):
        raise HTTPException(status_code=422, detail="请求需要 JSON 对象")
    # Parse the small public contract manually. FastAPI's default validation
    # response includes an `input` field, which could echo a malformed secret.
    provider_raw = raw.get("provider")
    endpoint_raw = raw.get("endpoint")
    api_key_raw = raw.get("api_key")
    credential_id_raw = raw.get("credential_id")
    api_format_raw = raw.get("api_format")
    env_name_raw = raw.get("env_name")
    if (
        (provider_raw is not None and not isinstance(provider_raw, str))
        or (endpoint_raw is not None and not isinstance(endpoint_raw, str))
        or (api_key_raw is not None and not isinstance(api_key_raw, str))
        or (credential_id_raw is not None and not isinstance(credential_id_raw, str))
        or (api_format_raw is not None and not isinstance(api_format_raw, str))
        or (env_name_raw is not None and not isinstance(env_name_raw, str))
    ):
        raise HTTPException(status_code=422, detail="provider、endpoint、api_key、credential_id、api_format 或 env_name 类型无效")
    provider = str(provider_raw or "").strip().lower()
    api_key = str(api_key_raw or "")
    credential_id = str(credential_id_raw or "").strip()
    has_memory_key = bool(api_key.strip())
    has_saved_ref = bool(credential_id)
    if has_memory_key == has_saved_ref:
        raise HTTPException(status_code=422, detail="请提供未保存 API key 或已保存 credential_id（二选一）")

    secret: str | None = None
    credential: sqlite3.Row | None = None
    metadata: dict[str, Any] = {}
    endpoint_value = ""
    api_format = ""
    try:
        if has_saved_ref:
            if len(credential_id) > 160:
                raise HTTPException(status_code=422, detail="credential_id 无效")
            with connect_db() as db:
                credential = db.execute("SELECT * FROM credentials WHERE id=?", (credential_id,)).fetchone()
                if credential is not None:
                    metadata = credential_metadata_row(credential_id, db) or {}
            if credential is None:
                raise HTTPException(status_code=422, detail="所选凭据不存在")
            provider = str(credential["provider"] or "custom").strip().lower()
            env_name = str(credential["env_name"] or "")
            try:
                api_format = normalize_api_format(metadata.get("api_format"), provider, env_name)
            except ValueError as exc:
                raise HTTPException(status_code=422, detail=str(exc)) from exc
            try:
                secret = keychain_secret(str(credential["credential_ref"]))
            except RuntimeError as exc:
                if is_keychain_permission_error(exc):
                    return {
                        "ok": False,
                        "status": "environment_restricted",
                        "message": str(exc),
                        "latency_ms": 0,
                        "provider": provider,
                        "endpoint": str(credential["endpoint"] or ""),
                        "api_format": api_format,
                        "models": [],
                    }
                return {
                    "ok": False,
                    "status": "unsupported",
                    "message": "无法读取已保存的 Keychain 凭据",
                    "latency_ms": 0,
                    "provider": provider,
                    "endpoint": str(credential["endpoint"] or ""),
                    "api_format": api_format,
                    "models": [],
                }
            except Exception:
                return {
                    "ok": False,
                    "status": "unsupported",
                    "message": "无法读取已保存的 Keychain 凭据",
                    "latency_ms": 0,
                    "provider": provider,
                    "endpoint": str(credential["endpoint"] or ""),
                    "api_format": api_format,
                    "models": [],
                }
            endpoint_value = str(credential["endpoint"] or "").strip()
        else:
            if provider not in {"openai", "anthropic", "google", "custom"}:
                raise HTTPException(status_code=422, detail="provider 必须是 openai、anthropic、google 或 custom")
            secret = api_key
            if not 8 <= len(secret) <= 8192 or "\x00" in secret or "\r" in secret or "\n" in secret:
                raise HTTPException(status_code=422, detail="API key 长度需要在8至8192字符之间且不能含控制换行")
            env_name = str(env_name_raw or "").strip().upper()
            if env_name and env_name not in ALLOWED_CREDENTIAL_ENVS:
                raise HTTPException(status_code=422, detail="不支持的模型密钥类型")
            try:
                api_format = normalize_api_format(api_format_raw, provider, env_name)
            except ValueError as exc:
                raise HTTPException(status_code=422, detail=str(exc)) from exc
            try:
                env_name = normalize_service_env_name(env_name, api_format)
            except ValueError as exc:
                raise HTTPException(status_code=422, detail=str(exc)) from exc
            endpoint_value = str(endpoint_raw or "").strip()
        if not endpoint_value:
            endpoint_value = DEFAULT_API_ENDPOINTS.get(provider, "")
        if not endpoint_value:
            raise HTTPException(status_code=422, detail="自定义服务必须提供 Endpoint")
        endpoint_value = validate_endpoint(endpoint_value)
        if secret and secret in endpoint_value:
            raise HTTPException(status_code=422, detail="Endpoint 不能包含 API key")
    finally:
        # Do not retain a request-body reference to a submitted key after the
        # endpoint and credential selection have been resolved.
        raw.pop("api_key", None)
        api_key = ""

    started = time.perf_counter()
    result: dict[str, Any] = {
        "ok": False,
        "status": "protocol_error",
        "message": "模型目录测试失败",
        "latency_ms": 0,
        "provider": provider,
        "endpoint": endpoint_value,
        "api_format": api_format,
        "models": [],
    }
    try:
        import httpx
    except Exception:
        result["status"] = "unsupported"
        result["message"] = "当前运行环境不支持 API 连接测试"
        result["latency_ms"] = int((time.perf_counter() - started) * 1000)
        secret = None
        return result

    if api_format == "anthropic-messages":
        headers = {"x-api-key": str(secret), "anthropic-version": "2023-06-01"}
    elif api_format == "google-generative-ai":
        headers = {"x-goog-api-key": str(secret)}
    else:
        headers = {"Authorization": f"Bearer {secret}"}
    try:
        deadline = time.perf_counter() + 10.0
        timeout = httpx.Timeout(connect=3.0, read=6.0, write=3.0, pool=3.0)
        async with httpx.AsyncClient(timeout=timeout, follow_redirects=False, trust_env=False) as client:
            async with client.stream("GET", _models_endpoint(endpoint_value), headers=headers) as response:
                if 300 <= response.status_code < 400:
                    result["status"] = "redirect"
                    result["message"] = "Endpoint 返回重定向；为避免泄露凭据，未跟随重定向"
                elif response.status_code in {401, 403}:
                    result["status"] = "auth_error"
                    result["message"] = "API key 未被 Endpoint 接受"
                elif response.status_code == 429:
                    result["status"] = "rate_limited"
                    result["message"] = "Endpoint 返回限流"
                elif response.status_code < 200 or response.status_code >= 300:
                    result["status"] = "protocol_error"
                    result["message"] = f"Endpoint 返回 HTTP {response.status_code}"
                else:
                    content_length = response.headers.get("content-length")
                    try:
                        declared_size = int(content_length) if content_length else 0
                    except ValueError:
                        declared_size = 0
                    if declared_size > MAX_MODELS_RESPONSE_BYTES:
                        result["status"] = "response_too_large"
                        result["message"] = "模型目录响应超过大小限制"
                    else:
                        chunks: list[bytes] = []
                        total = 0
                        async for chunk in response.aiter_bytes():
                            if time.perf_counter() > deadline:
                                result["status"] = "timeout"
                                result["message"] = "Endpoint 响应超时"
                                break
                            total += len(chunk)
                            if total > MAX_MODELS_RESPONSE_BYTES:
                                result["status"] = "response_too_large"
                                result["message"] = "模型目录响应超过大小限制"
                                break
                            chunks.append(chunk)
                        if result["status"] not in {"response_too_large", "timeout"} and time.perf_counter() > deadline:
                            result["status"] = "timeout"
                            result["message"] = "Endpoint 响应超时"
                        if result["status"] not in {"response_too_large", "timeout"}:
                            try:
                                document = json.loads(b"".join(chunks).decode("utf-8"))
                            except (UnicodeDecodeError, ValueError, json.JSONDecodeError):
                                result["status"] = "protocol_error"
                                result["message"] = "Endpoint 未返回合法 JSON 模型目录"
                            else:
                                models = document.get("data") if isinstance(document, dict) else None
                                if models is None and isinstance(document, dict):
                                    models = document.get("models")
                                if not isinstance(models, list) or any(not isinstance(item, dict) for item in models):
                                    result["status"] = "protocol_error"
                                    result["message"] = "Endpoint 返回的模型目录格式不受支持"
                                else:
                                    try:
                                        normalized_models = normalize_service_models(models, secret=secret)
                                    except ValueError:
                                        normalized_models = []
                                    if not normalized_models:
                                        result["status"] = "protocol_error"
                                        result["message"] = "Endpoint 未返回可用的模型目录"
                                    else:
                                        result["ok"] = True
                                        result["status"] = "ok"
                                        result["message"] = "模型目录连接成功；这不代表已完成模型推理"
                                        result["models"] = normalized_models
                                        result["model_count"] = len(normalized_models)
    except httpx.TimeoutException:
        result["status"] = "timeout"
        result["message"] = "Endpoint 连接超时"
    except httpx.RequestError:
        result["status"] = "unreachable"
        result["message"] = "Endpoint 无法连接"
    except Exception:
        result["status"] = "protocol_error"
        result["message"] = "Endpoint 响应无法处理"
    finally:
        result["latency_ms"] = int((time.perf_counter() - started) * 1000)
        secret = None
        headers.clear()
    return result


def store_keychain_secret(ref: str, secret: str) -> None:
    if sys.platform != 'darwin': raise RuntimeError('当前安全密钥保存仅支持 macOS Keychain')
    import ctypes
    security = ctypes.CDLL('/System/Library/Frameworks/Security.framework/Security')
    add = security.SecKeychainAddGenericPassword
    add.argtypes = [ctypes.c_void_p,ctypes.c_uint32,ctypes.c_char_p,ctypes.c_uint32,ctypes.c_char_p,ctypes.c_uint32,ctypes.c_void_p,ctypes.c_void_p]
    add.restype = ctypes.c_int32
    service = ('kxy/'+ref).encode(); account = b'kxy'; password = secret.encode()
    status = add(None,len(service),service,len(account),account,len(password),password,None)
    if status != 0: raise RuntimeError(keychain_error_message(status, '保存'))


@app.post('/api/credentials/store')
async def store_credential(request: Request):
    try:
        payload = await request.json()
    except (TypeError, ValueError):
        raise HTTPException(status_code=422, detail="请求需要 JSON 对象")
    if not isinstance(payload, dict):
        raise HTTPException(status_code=422, detail="请求需要 JSON 对象")
    secret = payload.get("api_key")
    if not isinstance(secret, str) or not 8 <= len(secret) <= 8192 or "\x00" in secret or "\r" in secret or "\n" in secret:
        raise HTTPException(status_code=422, detail="API key 长度需要在8至8192字符之间且不能含控制换行")
    provider = _credential_provider(payload.get("provider") or "custom")
    env_name = str(payload.get("env_name") or "").strip().upper()
    if env_name and env_name not in ALLOWED_CREDENTIAL_ENVS:
        raise HTTPException(status_code=422, detail="不支持的模型密钥类型")
    try:
        api_format = normalize_api_format(payload.get("api_format"), provider, env_name)
        env_name = normalize_service_env_name(env_name, api_format)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    endpoint_raw = payload.get("endpoint")
    if endpoint_raw is not None and not isinstance(endpoint_raw, str):
        raise HTTPException(status_code=422, detail="endpoint 类型无效")
    endpoint = str(endpoint_raw or "").strip() or DEFAULT_API_ENDPOINTS.get(provider, "")
    if not endpoint:
        raise HTTPException(status_code=422, detail="自定义服务必须提供 Endpoint")
    endpoint = validate_endpoint(endpoint)
    if secret in endpoint:
        raise HTTPException(status_code=422, detail="Endpoint 不能包含 API key")
    raw_models = payload.get("models", [])
    try:
        models = normalize_service_models(raw_models, secret=secret)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    name = _credential_name(payload.get("name"), provider)
    if _secret_in_public_credential_fields(secret, provider, endpoint, env_name, name, api_format, raw_models):
        raise HTTPException(status_code=422, detail="API key 不能出现在服务公开字段中")
    ref = "managed-" + uuid.uuid4().hex
    identifier = uuid.uuid4().hex
    try:
        store_keychain_secret(ref, secret)
    except RuntimeError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    finally:
        payload.pop("api_key", None)
        secret = None
    with connect_db() as db:
        db.execute(
            "INSERT INTO credentials(id,provider,credential_ref,endpoint,env_name,created_at) VALUES (?,?,?,?,?,?)",
            (identifier, provider, ref, endpoint, env_name, utc_now()),
        )
        save_credential_metadata(db, identifier, name=name, api_format=api_format, models=models)
        row = db.execute("SELECT id, provider, credential_ref, endpoint, env_name, created_at FROM credentials WHERE id=?", (identifier,)).fetchone()
        assert row is not None
        return credential_public(row, configured=True, db=db)


@app.post("/api/credentials")
def link_credential(payload: CredentialLinkRequest) -> dict[str, Any]:
    ref = payload.credential_ref.strip()
    if not re.fullmatch(r"[A-Za-z0-9._-]{1,80}", ref):
        raise HTTPException(status_code=422, detail="credential_ref must be a short Keychain label")
    if payload.env_name not in ALLOWED_CREDENTIAL_ENVS:
        raise HTTPException(status_code=422, detail="Only supported provider API-key environment names are accepted")
    provider = _credential_provider(payload.provider)
    payload.endpoint = validate_endpoint(payload.endpoint)
    try:
        api_format = normalize_api_format(payload.api_format, provider, payload.env_name)
        env_name = normalize_service_env_name(payload.env_name, api_format, default=False)
        models = normalize_service_models(payload.models)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    if not keychain_has(ref):
        raise HTTPException(status_code=422, detail="The Keychain reference is not configured; kxy did not receive a secret")
    identifier = uuid.uuid4().hex
    with connect_db() as db:
        db.execute(
            "INSERT INTO credentials(id, provider, credential_ref, endpoint, env_name, created_at) VALUES (?, ?, ?, ?, ?, ?)",
            (identifier, provider, ref, payload.endpoint, env_name, utc_now()),
        )
        save_credential_metadata(
            db,
            identifier,
            name=_credential_name(payload.name, provider),
            api_format=api_format,
            models=models,
        )
        row = db.execute("SELECT id, provider, credential_ref, endpoint, env_name, created_at FROM credentials WHERE id=?", (identifier,)).fetchone()
        assert row is not None
        return credential_public(row, configured=True, db=db)


@app.put("/api/credentials/{credential_id}")
async def update_credential(credential_id: str, request: Request) -> dict[str, Any]:
    """Edit public service metadata while keeping the secret Keychain-only."""
    if not re.fullmatch(r"[A-Za-z0-9._-]{1,160}", credential_id):
        raise HTTPException(status_code=422, detail="credential_id 无效")
    try:
        payload = await request.json()
    except (TypeError, ValueError):
        raise HTTPException(status_code=422, detail="请求需要 JSON 对象")
    if not isinstance(payload, dict):
        raise HTTPException(status_code=422, detail="请求需要 JSON 对象")
    api_key = payload.get("api_key")
    if api_key is not None and not isinstance(api_key, str):
        raise HTTPException(status_code=422, detail="api_key 类型无效")
    if isinstance(api_key, str) and api_key and (not 8 <= len(api_key) <= 8192 or "\x00" in api_key or "\r" in api_key or "\n" in api_key):
        raise HTTPException(status_code=422, detail="API key 长度需要在8至8192字符之间且不能含控制换行")
    with connect_db() as db:
        current = db.execute("SELECT id, provider, credential_ref, endpoint, env_name, created_at FROM credentials WHERE id=?", (credential_id,)).fetchone()
        if current is None:
            raise HTTPException(status_code=404, detail="所选凭据不存在")
        metadata = credential_metadata_row(credential_id, db) or {}
        current_provider = str(current["provider"] or "custom").strip().lower()
        current_env = str(current["env_name"] or "")
        current_endpoint = str(current["endpoint"] or "")
        try:
            current_format = normalize_api_format(metadata.get("api_format"), current_provider, current_env)
        except ValueError as exc:
            raise HTTPException(status_code=422, detail=str(exc)) from exc
        provider = _credential_provider(payload.get("provider", current_provider))
        env_was_supplied = "env_name" in payload
        env_name = str(payload.get("env_name", current_env) or "").strip().upper()
        if env_name and env_name not in ALLOWED_CREDENTIAL_ENVS:
            raise HTTPException(status_code=422, detail="不支持的模型密钥类型")
        endpoint_raw = payload.get("endpoint", current_endpoint)
        if endpoint_raw is not None and not isinstance(endpoint_raw, str):
            raise HTTPException(status_code=422, detail="endpoint 类型无效")
        endpoint = str(endpoint_raw or "").strip()
        if not endpoint:
            endpoint = DEFAULT_API_ENDPOINTS.get(provider, "")
        if not endpoint:
            raise HTTPException(status_code=422, detail="自定义服务必须提供 Endpoint")
        endpoint = validate_endpoint(endpoint)
        try:
            api_format = normalize_api_format(payload.get("api_format", current_format), provider, env_name)
            if env_was_supplied:
                env_name = normalize_service_env_name(env_name, api_format)
            elif "api_format" in payload and api_format != current_format:
                env_name = normalize_service_env_name("", api_format)
        except ValueError as exc:
            raise HTTPException(status_code=422, detail=str(exc)) from exc
        raw_models = payload.get("models") if "models" in payload else None
        try:
            if "models" in payload:
                models = normalize_service_models(payload.get("models"))
            else:
                models = normalize_service_models(json.loads(metadata.get("models") or "[]"))
        except (TypeError, ValueError, json.JSONDecodeError) as exc:
            raise HTTPException(status_code=422, detail="models 格式无效") from exc
        name = _credential_name(payload.get("name", metadata.get("name")), provider)
        destination_changed = (
            endpoint != current_endpoint
            or provider != current_provider
            or env_name != current_env
            or api_format != current_format
        )
        replacement = api_key.strip() if isinstance(api_key, str) else ""
        if replacement and _secret_in_public_credential_fields(
            replacement,
            provider,
            endpoint,
            env_name,
            name,
            api_format,
            raw_models if raw_models is not None else models,
        ):
            raise HTTPException(status_code=422, detail="API key 不能出现在服务公开字段中")
        if destination_changed and not replacement:
            raise HTTPException(status_code=422, detail="修改 Endpoint、密钥类型或 API 格式前必须重新输入 API key")
        new_ref = str(current["credential_ref"])
        if replacement:
            new_ref = "managed-" + uuid.uuid4().hex
            try:
                store_keychain_secret(new_ref, replacement)
            except RuntimeError as exc:
                raise HTTPException(status_code=422, detail=str(exc)) from exc
        db.execute(
            "UPDATE credentials SET provider=?, credential_ref=?, endpoint=?, env_name=? WHERE id=?",
            (provider, new_ref, endpoint, env_name, credential_id),
        )
        save_credential_metadata(db, credential_id, name=name, api_format=api_format, models=models)
        row = db.execute("SELECT id, provider, credential_ref, endpoint, env_name, created_at FROM credentials WHERE id=?", (credential_id,)).fetchone()
        assert row is not None
        payload.pop("api_key", None)
        return credential_public(row, configured=True, db=db)


@app.delete("/api/credentials/{credential_id}")
def unlink_credential(credential_id: str) -> dict[str, bool]:
    with connect_db() as db:
        db.execute("DELETE FROM credentials WHERE id=?", (credential_id,))
        db.execute("DELETE FROM credential_metadata WHERE credential_id=?", (credential_id,))
    return {"ok": True}


@app.post("/api/filters/preview")
def preview_filter(payload: FilterPreviewRequest) -> dict[str, Any]:
    try:
        return filter_payload_items(payload.items, payload.config)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@app.post("/api/runs")
def create_run(payload: RunRequest) -> dict[str, Any]:
    workflow = payload.workflow; workflow_id = payload.workflow_id
    if workflow is None and workflow_id:
        with connect_db() as db: row=db.execute('SELECT document FROM workflows WHERE id=?',(workflow_id,)).fetchone()
        if row is None: raise HTTPException(404,'流程不存在')
        workflow=json.loads(row['document'])
    if workflow is None: raise HTTPException(422,'需要 workflow 或 workflow_id')
    workflow=normalize_workflow(workflow); validation=validate_workflow(workflow)
    if not validation['ok']: raise HTTPException(422,validation)
    run_id=uuid.uuid4().hex; run_dir=RUNS_ROOT/run_id; workspace=run_dir/'workspace'
    workspace.mkdir(parents=True)
    try:
        compiled = flatten_workflow(workflow)
        # Resolve model/API bindings once, before a run can reach a human pause.
        # The optional configuration module is intentionally called through a
        # direct helper; legacy V1 analyzer fields remain valid when no binding
        # reference is present.
        workflow = snapshot_workflow_bindings(workflow)
        selected=workflow_skill_snapshot(workflow)
        selected_mcps = workflow_mcp_snapshot(workflow)
        for skill in selected:
            with connect_db() as db: row=db.execute('SELECT root_path FROM skills WHERE id=?',(skill['id'],)).fetchone()
            root=Path(row['root_path'])
            if not resolved_inside(root,SKILLS_ROOT) or snapshot_skill(root)[0]!=skill['snapshot_hash']: raise ValueError('Skill 已变化，请重新导入')
            target=workspace/'skill-snapshots'/skill['id']; target.parent.mkdir(exist_ok=True)
            shutil.copytree(root,target)
            skill['path']=str(target.relative_to(workspace))
        sources={}
        for source_path, node in workflow_file_nodes(workflow):
            file_id=str((node.get('data') or {}).get('file_id',''))
            if file_id in sources:
                continue
            with connect_db() as db: row=db.execute('SELECT * FROM files WHERE id=?',(file_id,)).fetchone()
            if row is None: raise ValueError(f'资料节点 {source_path} 未绑定有效资料')
            source=file_row(row); original=Path(row['stored_path'])
            if not resolved_inside(original,FILES_ROOT): raise ValueError('资料存储路径无效')
            content=original.read_bytes()
            if hashlib.sha256(content).hexdigest()!=source['sha256']: raise ValueError('原始资料校验失败')
            target=workspace/'inputs'/safe_run_input_name(row['display_name'], source['sha256']); target.parent.mkdir(exist_ok=True)
            if not target.exists(): target.write_bytes(content)
            sources[file_id]={'status':'source','file_id':file_id,'name':source['display_name'],'sha256':source['sha256'],'mime':source['mime'],'text':source.get('preview') or '', 'file_path':str(target.relative_to(workspace)),'source_status':source['status'],'metadata':source['metadata']}
        snapshot={'version':'kxy.run.v1','checkpoint_version':CHECKPOINT_VERSION,'workflow_id':workflow_id or workflow.get('id'),'workflow':workflow,'settings':get_settings(),'skills':selected,'mcps':selected_mcps,'skill_mappings':{},'mcp_mappings':{},'grants':workflow_grant_snapshot(workflow),'sources':sources,'compiled':{'nodes':[{'flat_id':item['flat_id'],'path':item['path'],'type':item['node']['type']} for item in compiled['records']], 'edges':compiled['edges'], 'blackboxes':compiled['blackboxes']},'created_at':utc_now()}
    except (ValueError,OSError) as exc:
        shutil.rmtree(run_dir,ignore_errors=True)
        raise HTTPException(422,str(exc)) from exc
    with connect_db() as db:
        db.execute("INSERT INTO runs(id,workflow_id,status,snapshot,run_dir,created_at) VALUES (?,?,'pending',?,?,?)",(run_id,workflow_id or workflow.get('id'),json_text(snapshot),str(run_dir),utc_now()))
        compiled = flatten_workflow(workflow)
        for record in compiled['records']:
            db.execute("INSERT INTO run_nodes(run_id,node_id,status) VALUES (?,?,'pending')",(run_id,record['flat_id']))
        # Outer blackbox cards do not become a second scheduler vertex, but they
        # do get a durable row for visible status/error mapping.
        for path in compiled['blackboxes']:
            db.execute("INSERT OR IGNORE INTO run_nodes(run_id,node_id,status) VALUES (?,?,'pending')",(run_id,path))
        db.execute(
            "INSERT INTO run_attempts(run_id,attempt_no,status,reused_nodes,rerun_nodes,created_at) VALUES (?,?,?, ?, ?, ?)",
            (
                run_id,
                1,
                "queued",
                json_text([]),
                json_text([record["flat_id"] for record in compiled["records"]]),
                utc_now(),
            ),
        )
    append_event(run_id,'run_queued',{'run_id':run_id,'attempt':1}); start_background_run(run_id, 1)
    return {'id':run_id,'status':'pending','snapshot':snapshot,'validation':validation}


@app.get("/api/runs")
def list_runs() -> list[dict[str, Any]]:
    with connect_db() as db:
        rows = db.execute("SELECT id, workflow_id, status, error, created_at, started_at, finished_at FROM runs ORDER BY created_at DESC LIMIT 40").fetchall()
    return [dict(row) for row in rows]


@app.get("/api/runs/{run_id}")
def get_run(run_id: str) -> dict[str, Any]:
    return get_run_payload(run_id)


@app.post("/api/runs/{run_id}/resume")
def resume_run(run_id: str) -> dict[str, Any]:
    """Queue one explicit same-run recovery attempt from durable checkpoints."""

    with RUN_LOCK:
        # The database attempt row can already be terminal while the previous
        # worker is still unwinding.  Keep the in-memory ownership marker as a
        # second mutex so a new attempt cannot overlap its cleanup.
        if run_id in RUN_CANCEL_EVENTS:
            raise HTTPException(status_code=409, detail="该运行仍有 worker 正在退出，请稍后重试")
        try:
            with connect_db() as db:
                row = db.execute("SELECT * FROM runs WHERE id=?", (run_id,)).fetchone()
                if row is None:
                    raise HTTPException(status_code=404, detail="Run not found")
                if row["status"] not in {"failed", "interrupted"}:
                    raise HTTPException(
                        status_code=409,
                        detail=f"Only failed or interrupted runs can be resumed; current status is {row['status']}",
                    )
                active = db.execute(
                    "SELECT 1 FROM run_attempts WHERE run_id=? AND status IN ('queued','running') LIMIT 1",
                    (run_id,),
                ).fetchone()
                if active is not None:
                    raise HTTPException(status_code=409, detail="该运行已有恢复尝试正在进行")
                snapshot, reused, rerun = resume_checkpoint_plan(run_id, row, db)
                attempt_row = db.execute(
                    "SELECT COALESCE(MAX(attempt_no), 0) AS max_attempt FROM run_attempts WHERE run_id=?",
                    (run_id,),
                ).fetchone()
                attempt_no = int(attempt_row["max_attempt"] or 0) + 1
                # A restart leaves an unfinished approval as interrupted.  It
                # has no decision, so presenting the same approval again is
                # safe; rejected/approved rows are never rewritten here.
                db.execute(
                    "UPDATE run_approvals SET status='pending', decided_at=NULL "
                    "WHERE run_id=? AND status='interrupted' AND decision IS NULL",
                    (run_id,),
                )
                db.execute(
                    "UPDATE run_nodes SET status='pending', message=NULL, output=NULL, "
                    "started_at=NULL, finished_at=NULL WHERE run_id=? AND status <> 'succeeded'",
                    (run_id,),
                )
                db.execute(
                    "INSERT INTO run_attempts(run_id,attempt_no,status,reason,reused_nodes,rerun_nodes,created_at) "
                    "VALUES (?,?,?,?,?,?,?)",
                    (
                        run_id,
                        attempt_no,
                        "queued",
                        "explicit_resume",
                        json_text(sorted(reused)),
                        json_text(rerun),
                        utc_now(),
                    ),
                )
                db.execute(
                    "UPDATE runs SET status='pending', error=NULL, output_manifest=NULL, "
                    "cancel_requested=0, finished_at=NULL WHERE id=?",
                    (run_id,),
                )
        except HTTPException:
            raise
        except (ValueError, OSError, sqlite3.IntegrityError) as exc:
            raise HTTPException(status_code=409, detail=str(exc)) from exc
    append_event(
        run_id,
        "run_resume_queued",
        {"run_id": run_id, "attempt": attempt_no, "reused_nodes": sorted(reused), "rerun_nodes": rerun},
    )
    start_background_run(run_id, attempt_no)
    return {
        "id": run_id,
        "status": "pending",
        "attempt": attempt_no,
        "reused_nodes": sorted(reused),
        "rerun_nodes": rerun,
        "snapshot": snapshot,
    }


def _waiting_loop_round(run_id: str, loop_path: str) -> sqlite3.Row:
    if not loop_path or len(loop_path) > MAX_TEXT_BYTES:
        raise HTTPException(status_code=422, detail="loop_path 无效")
    with connect_db() as db:
        if db.execute("SELECT 1 FROM runs WHERE id=?", (run_id,)).fetchone() is None:
            raise HTTPException(status_code=404, detail="Run not found")
        row = db.execute(
            "SELECT * FROM loop_rounds WHERE run_id=? AND loop_path=? "
            "AND status IN ('budget_waiting','round_limit_waiting') "
            "ORDER BY attempt_no DESC, round_no DESC LIMIT 1",
            (run_id, loop_path),
        ).fetchone()
    if row is None:
        raise HTTPException(status_code=409, detail="该循环当前没有等待中的预算或轮数控制")
    if row["control"]:
        raise HTTPException(status_code=409, detail="该循环控制请求已在处理中")
    return row


def _queue_loop_control(run_id: str, loop_path: str, control: str, payload: dict[str, Any]) -> dict[str, Any]:
    row = _waiting_loop_round(run_id, loop_path)
    with connect_db() as db:
        updated = db.execute(
            "UPDATE loop_rounds SET control=?, control_payload=? "
            "WHERE run_id=? AND attempt_no=? AND loop_path=? AND round_no=? AND control IS NULL",
            (
                control,
                json_text(redact_object(payload)),
                run_id,
                row["attempt_no"],
                loop_path,
                row["round_no"],
            ),
        ).rowcount
    if updated != 1:
        raise HTTPException(status_code=409, detail="该循环控制请求与另一请求冲突")
    append_event(
        run_id,
        "loop_control_queued",
        {
            "loop_path": loop_path,
            "round": row["round_no"],
            "control": control,
            "payload": payload,
        },
    )
    return get_run_payload(run_id)


@app.post("/api/runs/{run_id}/loops/{loop_path:path}/continue")
def continue_loop(run_id: str, loop_path: str, payload: LoopControlRequest | None = None) -> dict[str, Any]:
    """Continue one paused loop with explicitly bounded extra budget."""

    row = _waiting_loop_round(run_id, loop_path)
    request = payload or LoopControlRequest()
    try:
        additional_rounds = int(request.additional_rounds)
        additional_seconds = float(request.additional_seconds)
    except (TypeError, ValueError, OverflowError) as exc:
        raise HTTPException(status_code=422, detail="循环继续参数必须是有限的非负数") from exc
    if (
        isinstance(request.additional_rounds, bool)
        or additional_rounds < 0
        or additional_rounds > LOOP_MAX_ROUNDS
        or not (0 <= additional_seconds <= LOOP_MAX_ACTIVE_BUDGET_SECONDS)
        or additional_rounds == 0 and additional_seconds == 0
    ):
        raise HTTPException(status_code=422, detail="循环继续参数超出有界范围")
    current_rounds = int(row["effective_max_rounds"] or LOOP_DEFAULT_MAX_ROUNDS)
    current_budget = float(row["budget_seconds"] or LOOP_DEFAULT_ACTIVE_BUDGET_SECONDS)
    if current_rounds + additional_rounds > LOOP_MAX_ROUNDS or current_budget + additional_seconds > LOOP_MAX_ACTIVE_BUDGET_SECONDS:
        raise HTTPException(status_code=422, detail="循环继续后的总上限超过 V9 安全边界")
    return _queue_loop_control(
        run_id,
        loop_path,
        "continue",
        {
            "additional_rounds": additional_rounds,
            "additional_seconds": additional_seconds,
        },
    )


@app.post("/api/runs/{run_id}/loops/{loop_path:path}/stop")
def stop_loop(run_id: str, loop_path: str) -> dict[str, Any]:
    """Stop a paused loop; no non-passing executor output is accepted."""

    return _queue_loop_control(run_id, loop_path, "stop", {"reason": "user_stop"})


@app.post("/api/runs/{run_id}/approvals/{node_id:path}")
def decide_approval(run_id: str, node_id: str, payload: ApprovalRequest) -> dict[str, Any]:
    decision = payload.decision.strip().lower()
    if decision in {"revise", "return_for_revision"}:
        decision = "return"
    if decision not in {"approve", "reject", "return"}:
        raise HTTPException(status_code=422, detail="decision must be approve, reject, or return")
    note = redact(str(payload.note or ""), max_length=MAX_EVENT_TEXT)
    revision_present = payload.revision_text is not None
    revision_text = redact(str(payload.revision_text), max_length=MAX_TEXT_BYTES) if revision_present else None
    if decision == "return" and not note and not revision_present:
        raise HTTPException(status_code=422, detail="return requires feedback note or a revision copy")
    with connect_db() as db:
        row = db.execute(
            "SELECT * FROM run_approvals WHERE run_id=? AND (node_id=? OR node_path=?) ORDER BY CASE WHEN node_id=? THEN 0 ELSE 1 END LIMIT 1",
            (run_id, node_id, node_id, node_id),
        ).fetchone()
        if row is None:
            if db.execute("SELECT 1 FROM runs WHERE id=?", (run_id,)).fetchone() is None:
                raise HTTPException(status_code=404, detail="Run not found")
            raise HTTPException(status_code=404, detail="Approval not found")
        if row["status"] != "pending":
            raise HTTPException(status_code=409, detail=f"Approval is already {row['status']}")
        if decision == "return" and not bool(row["allow_return"] if "allow_return" in row.keys() else False):
            raise HTTPException(status_code=409, detail="This human node is not the supported loop review return gate")
        status = "approved" if decision == "approve" else "returned" if decision == "return" else "rejected"
        db.execute(
            "UPDATE run_approvals SET status=?, decision=?, note=?, revision_text=?, revision_note=?, revision_at=?, decided_at=? WHERE run_id=? AND node_id=? AND status='pending'",
            (status, decision, note or None, revision_text, note or None, utc_now() if revision_present else None, utc_now(), run_id, row["node_id"]),
        )
        updated = db.execute("SELECT * FROM run_approvals WHERE run_id=? AND node_id=?", (run_id, row["node_id"])).fetchone()
    if updated is None:
        raise HTTPException(status_code=409, detail="Approval changed while deciding")
    append_event(
        run_id,
        "approval_decided",
        {"node_id": row["node_id"], "node_path": row["node_path"], "decision": decision, "note": note or None, "revised": revision_present},
    )
    with RUN_LOCK:
        event = RUN_APPROVAL_EVENTS.get((run_id, row["node_id"]))
    if event:
        event.set()
    return approval_payload(updated)


@app.post("/api/runs/{run_id}/cancel")
def cancel_run(run_id: str) -> dict[str, Any]:
    with connect_db() as db:
        row = db.execute("SELECT status FROM runs WHERE id=?", (run_id,)).fetchone()
        if row is None:
            raise HTTPException(status_code=404, detail="Run not found")
        db.execute("UPDATE runs SET cancel_requested=1 WHERE id=?", (run_id,))
        approval_rows = db.execute("SELECT node_id FROM run_approvals WHERE run_id=? AND status='pending'", (run_id,)).fetchall()
        if row["status"] == "waiting":
            db.execute(
                "UPDATE run_approvals SET status='cancelled', decision='cancel', note='Run cancelled by user', decided_at=? WHERE run_id=? AND status='pending'",
                (utc_now(), run_id),
            )
    with RUN_LOCK:
        event = RUN_CANCEL_EVENTS.get(run_id)
        process = RUN_PROCESSES.get(run_id)
        approval_events = [RUN_APPROVAL_EVENTS.get((run_id, item["node_id"])) for item in approval_rows]
    if event:
        event.set()
    for approval_event_handle in approval_events:
        if approval_event_handle:
            approval_event_handle.set()
    if process:
        kill_process_group(process)
    append_event(run_id, "cancel_requested", {"run_id": run_id})
    return {"ok": True, "status": row["status"]}


@app.get("/api/runs/{run_id}/events")
async def run_events(run_id: str, after: int = Query(default=0, ge=0)) -> StreamingResponse:
    with connect_db() as db:
        if db.execute("SELECT 1 FROM runs WHERE id=?", (run_id,)).fetchone() is None:
            raise HTTPException(status_code=404, detail="Run not found")

    async def stream():
        cursor = after
        idle = 0
        while idle < 300:
            with connect_db() as db:
                rows = db.execute("SELECT id, event_type, payload, created_at FROM run_events WHERE run_id=? AND id>? ORDER BY id", (run_id, cursor)).fetchall()
                run = db.execute("SELECT status FROM runs WHERE id=?", (run_id,)).fetchone()
            if rows:
                idle = 0
                for row in rows:
                    cursor = row["id"]
                    yield f"id: {cursor}\ndata: {json_text({'id': cursor, 'type': row['event_type'], 'payload': json.loads(row['payload']), 'created_at': row['created_at']})}\n\n"
            else:
                idle += 1
                yield ": heartbeat\n\n"
            if run and run["status"] in {"succeeded", "failed", "rejected", "cancelled", "interrupted"} and not rows:
                break
            await asyncio.sleep(0.4)

    return StreamingResponse(stream(), media_type="text/event-stream", headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


@app.get("/api/runs/{run_id}/artifact/{artifact_path:path}")
def download_artifact(run_id: str, artifact_path: str):
    run = get_run_payload(run_id)
    artifacts = (run.get("output_manifest") or {}).get("artifacts", [])
    if artifact_path not in {item["path"] for item in artifacts}:
        raise HTTPException(404, "产物不存在或尚未完成")
    root = Path(run["run_dir"]) / "workspace"
    path = root / artifact_path
    if path.is_symlink() or not resolved_inside(path, root) or not path.is_file():
        raise HTTPException(404, "产物路径无效")
    return FileResponse(path, filename=path.name, media_type="application/octet-stream")


def _approval_file_candidates(run_id: str, row: sqlite3.Row) -> list[dict[str, Any]]:
    run_dir = RUNS_ROOT / run_id
    workspace = (run_dir / "workspace").resolve()
    try:
        inputs = json.loads(row["inputs"] or "[]")
    except (TypeError, ValueError, json.JSONDecodeError):
        return []
    manifest: dict[str, Any] = {}
    manifest_path = workspace / "manifest.json"
    if manifest_path.is_file() and not manifest_path.is_symlink():
        try:
            parsed = json.loads(manifest_path.read_text(encoding="utf-8"))
            manifest = parsed if isinstance(parsed, dict) else {}
        except (OSError, ValueError, json.JSONDecodeError):
            manifest = {}
    sources = manifest.get("sources") if isinstance(manifest.get("sources"), dict) else {}
    candidates: dict[str, dict[str, Any]] = {}

    def add(raw_path: Any, expected_sha: Any, expected_size: Any, name: Any) -> None:
        relative = str(raw_path or "").replace("\\", "/")
        candidate = Path(relative)
        if not relative or candidate.is_absolute() or ".." in candidate.parts:
            return
        digest = str(expected_sha or "")
        if not re.fullmatch(r"[0-9a-fA-F]{64}", digest):
            return
        candidates[relative] = {
            "path": relative,
            "sha256": digest.lower(),
            "size": expected_size,
            "name": safe_display_name(str(name or candidate.name)),
        }

    def visit(value: Any) -> None:
        if isinstance(value, list):
            for child in value:
                visit(child)
            return
        if not isinstance(value, dict):
            return
        file_path = value.get("file_path")
        source = sources.get(str(value.get("file_id") or "")) if value.get("file_id") else None
        add(
            file_path,
            value.get("sha256") or (source.get("sha256") if isinstance(source, dict) else None),
            value.get("size") or (source.get("size") if isinstance(source, dict) else None),
            value.get("display_name") or value.get("name") or (source.get("name") if isinstance(source, dict) else None),
        )
        artifacts = value.get("artifacts")
        if isinstance(artifacts, list):
            for artifact in artifacts:
                if isinstance(artifact, dict):
                    add(artifact.get("path"), artifact.get("sha256"), artifact.get("size"), artifact.get("name"))
        for key in ("items", "upstream", "inputs"):
            if key in value:
                visit(value[key])

    visit(inputs)
    return list(candidates.values())


@app.get("/api/runs/{run_id}/approval-file/{node_id:path}")
def download_approval_file(run_id: str, node_id: str, path: str = Query(...)):
    with connect_db() as db:
        row = db.execute(
            "SELECT * FROM run_approvals WHERE run_id=? AND (node_id=? OR node_path=?) "
            "ORDER BY CASE WHEN node_id=? THEN 0 ELSE 1 END LIMIT 1",
            (run_id, node_id, node_id, node_id),
        ).fetchone()
        run = db.execute("SELECT run_dir FROM runs WHERE id=?", (run_id,)).fetchone()
    if row is None or run is None:
        raise HTTPException(404, "Approval not found")
    relative = str(path or "").replace("\\", "/")
    allowed = next((item for item in _approval_file_candidates(run_id, row) if item["path"] == relative), None)
    if allowed is None:
        raise HTTPException(404, "文件不在当前人工节点的已验证输入白名单中")
    workspace = (Path(run["run_dir"]) / "workspace").resolve()
    target = workspace / relative
    if target.is_symlink() or not resolved_inside(target, workspace) or not target.is_file():
        raise HTTPException(404, "审批附件路径无效")
    if hashlib.sha256(target.read_bytes()).hexdigest().lower() != allowed["sha256"]:
        raise HTTPException(409, "审批附件哈希校验失败")
    if allowed.get("size") not in (None, ""):
        try:
            if target.stat().st_size != int(allowed["size"]):
                raise HTTPException(409, "审批附件大小校验失败")
        except (TypeError, ValueError):
            raise HTTPException(409, "审批附件大小元数据无效")
    return FileResponse(target, filename=allowed["name"], media_type="application/octet-stream")


STATIC_DIR = ROOT / "frontend" / "dist"
if STATIC_DIR.is_dir():
    app.mount("/", StaticFiles(directory=STATIC_DIR, html=True), name="frontend")
